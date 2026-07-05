# encoding: utf-8
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "compliance-wechatmp-investment-flow.md"
OUT_DIR = ROOT / "docs" / "generated"
OUT_DOCX = OUT_DIR / "公众号投研服务业务流程说明.docx"


FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simhei.ttf"),
    Path("C:/Windows/Fonts/simsun.ttc"),
]


def font_path() -> str:
    for candidate in FONT_CANDIDATES:
        if candidate.exists():
            return str(candidate)
    return "arial.ttf"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, *, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=10, bold=bold, color=color)


def set_table_borders(table, color: str = "B8C2CC", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("公众号投研服务业务流程说明")
    set_run_font(footer_run, size=9)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)


def wrap_text(text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for ch in text:
        candidate = current + ch
        bbox = font.getbbox(candidate)
        if bbox[2] - bbox[0] <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_box(draw: ImageDraw.ImageDraw, xy, text: str, font, fill: str, outline: str) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    lines = wrap_text(text, font, x2 - x1 - 34)
    total_h = len(lines) * 24
    y = y1 + ((y2 - y1) - total_h) / 2 - 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=font, fill="#1F2933")
        y += 24


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], label: str = "", font=None) -> None:
    draw.line([start, end], fill="#5B6775", width=3)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) < abs(ey - sy):
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 8, ey - 14 * direction), (ex + 8, ey - 14 * direction)]
    else:
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 14 * direction, ey - 8), (ex - 14 * direction, ey + 8)]
    draw.polygon(points, fill="#5B6775")
    if label and font:
        mx = (sx + ex) / 2
        my = (sy + ey) / 2
        bbox = draw.textbbox((0, 0), label, font=font)
        pad = 6
        draw.rounded_rectangle(
            (mx - (bbox[2] - bbox[0]) / 2 - pad, my - 16, mx + (bbox[2] - bbox[0]) / 2 + pad, my + 12),
            radius=8,
            fill="#FFFFFF",
            outline="#D0D7DE",
        )
        draw.text((mx - (bbox[2] - bbox[0]) / 2, my - 12), label, font=font, fill="#44546A")


def render_overall(path: Path) -> None:
    img = Image.new("RGB", (1700, 2200), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    f = ImageFont.truetype(font_path(), 25)
    small = ImageFont.truetype(font_path(), 22)
    boxes = [
        ("A", "客户在微信公众号输入内容", 650, 70),
        ("B", "微信服务器转发消息到云端服务", 650, 240),
        ("C", "云端校验消息、识别客户身份", 650, 410),
        ("D", "权限校验：是否已开通服务", 650, 580),
        ("E", "返回未开通或无权限提示", 190, 760),
        ("F", "识别业务类型：技术分析、利率、转债或其他", 650, 760),
        ("G", "执行业务处理", 650, 930),
        ("H", "生成或读取结果图片", 650, 1100),
        ("I", "结果暂存并绑定客户本次请求", 650, 1270),
        ("J", "公众号提示客户回复 1 获取结果，或直接返回图片", 650, 1440),
        ("K", "客户领取结果", 650, 1610),
        ("L", "云端记录请求、结果、领取状态和异常信息", 650, 1780),
    ]
    coords = {}
    for key, text, x, y in boxes:
        coords[key] = (x, y, x + 400, y + 95)
        draw_box(draw, coords[key], text, f, "#F5F8FB", "#7FA6C9")
    for a, b in [("A", "B"), ("B", "C"), ("C", "D"), ("F", "G"), ("G", "H"), ("H", "I"), ("I", "J"), ("J", "K"), ("K", "L")]:
        ax = (coords[a][0] + coords[a][2]) // 2
        ay = coords[a][3]
        bx = (coords[b][0] + coords[b][2]) // 2
        by = coords[b][1]
        arrow(draw, (ax, ay), (bx, by))
    arrow(draw, (650, 628), (590, 807), "已授权", small)
    arrow(draw, (650, 628), (590, 807), "", small)
    arrow(draw, (650, 628), (390, 807), "未授权", small)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, quality=95)


def render_technical(path: Path) -> None:
    img = Image.new("RGB", (1800, 1700), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    f = ImageFont.truetype(font_path(), 25)
    small = ImageFont.truetype(font_path(), 22)
    coords = {
        "A": (690, 70, 1110, 165),
        "B": (690, 240, 1110, 335),
        "C": (690, 410, 1110, 505),
        "D": (690, 580, 1110, 675),
        "E": (160, 760, 580, 855),
        "F": (690, 760, 1110, 855),
        "G": (160, 960, 580, 1055),
        "H": (690, 960, 1110, 1055),
        "I": (690, 1130, 1110, 1225),
        "J": (690, 1300, 1110, 1395),
        "K": (690, 1470, 1110, 1565),
        "L": (1220, 1470, 1640, 1565),
        "M": (1220, 1300, 1640, 1395),
    }
    labels = {
        "A": "客户输入股票代码/名称/#标的",
        "B": "云端识别为技术分析请求",
        "C": "校验客户权限",
        "D": "解析并标准化标的",
        "E": "返回标的不存在/多标的选择提示",
        "F": "查询是否已有有效缓存或已发布产品",
        "G": "直接使用缓存图片",
        "H": "调用技术分析程序生成报告和图表",
        "I": "AI 将技术分析报告整理为标准展示文本",
        "J": "渲染为信号卡/图表图片",
        "K": "保存结果文件并记录请求",
        "L": "提示客户回复 1 领取",
        "M": "客户回复 1 后返回图片",
    }
    for key, box in coords.items():
        draw_box(draw, box, labels[key], f, "#F5F8FB", "#7FA6C9")
    for a, b in [("A", "B"), ("B", "C"), ("C", "D"), ("H", "I"), ("I", "J"), ("J", "K"), ("L", "M")]:
        arrow(draw, ((coords[a][0] + coords[a][2]) // 2, coords[a][3]), ((coords[b][0] + coords[b][2]) // 2, coords[b][1]))
    arrow(draw, (690, 628), (580, 807), "标的不明确或不存在", small)
    arrow(draw, (900, 675), (900, 760), "标的有效", small)
    arrow(draw, (690, 807), (580, 1007), "命中有效结果", small)
    arrow(draw, (900, 855), (900, 960), "未命中", small)
    arrow(draw, (580, 1007), (690, 1518))
    arrow(draw, (1110, 1518), (1220, 1518))
    img.save(path, quality=95)


def render_daily(path: Path) -> None:
    img = Image.new("RGB", (1700, 1450), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    f = ImageFont.truetype(font_path(), 25)
    small = ImageFont.truetype(font_path(), 22)
    coords = {
        "A": (650, 70, 1050, 165),
        "B": (650, 240, 1050, 335),
        "C": (650, 410, 1050, 505),
        "D": (650, 580, 1050, 675),
        "E": (650, 750, 1050, 845),
        "F": (650, 920, 1050, 1015),
        "G": (170, 1100, 570, 1195),
        "H": (650, 1100, 1050, 1195),
        "I": (650, 1270, 1050, 1365),
        "J": (1130, 1270, 1530, 1365),
    }
    labels = {
        "A": "运营/管理员准备利率或转债素材",
        "B": "云端生成标准投研内容",
        "C": "渲染为图片",
        "D": "设置生效日期、失效时间并发布",
        "E": "客户在公众号输入“利率”或“转债”",
        "F": "云端校验客户权限",
        "G": "返回无权限提示",
        "H": "读取当前生效内容",
        "I": "返回对应图片",
        "J": "返回暂无内容提示",
    }
    for key, box in coords.items():
        draw_box(draw, box, labels[key], f, "#F5F8FB", "#7FA6C9")
    for a, b in [("A", "B"), ("B", "C"), ("C", "D"), ("D", "E"), ("E", "F"), ("H", "I")]:
        arrow(draw, ((coords[a][0] + coords[a][2]) // 2, coords[a][3]), ((coords[b][0] + coords[b][2]) // 2, coords[b][1]))
    arrow(draw, (650, 968), (570, 1148), "未授权", small)
    arrow(draw, (850, 1015), (850, 1100), "已授权", small)
    arrow(draw, (1050, 1148), (1130, 1318), "无有效内容", small)
    arrow(draw, (850, 1195), (850, 1270), "存在有效内容", small)
    img.save(path, quality=95)


def add_formatted_text(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)
    for run in paragraph.runs:
        set_run_font(run)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows) - 1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    header = rows[0]
    body = rows[2:] if len(rows) > 1 and all(set(c) <= {"-", " "} for c in rows[1]) else rows[1:]
    widths_map = {
        2: [3.2, 12.9],
        4: [2.6, 4.0, 6.0, 3.5],
    }
    widths = widths_map.get(len(header), [16.1 / len(header)] * len(header))
    set_table_widths(table, widths)
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "F2F4F7")
        set_cell_text(cell, text, bold=True, color="1F4D78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for r_idx, row in enumerate(body, start=1):
        cells = table.rows[r_idx].cells
        for c_idx, text in enumerate(row):
            set_cell_text(cells[c_idx], text)
            cells[c_idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    images = [
        OUT_DIR / "flow-overall.png",
        OUT_DIR / "flow-technical-analysis.png",
        OUT_DIR / "flow-rate-bond.png",
    ]
    render_overall(images[0])
    render_technical(images[1])
    render_daily(images[2])

    doc = Document()
    configure_doc(doc)

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    diagram_index = 0
    table_rows: list[list[str]] = []
    in_mermaid = False

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_markdown_table(doc, table_rows)
            table_rows = []

    for line in lines:
        raw = line.rstrip()
        if raw.startswith("```mermaid"):
            flush_table()
            in_mermaid = True
            continue
        if in_mermaid:
            if raw.startswith("```"):
                doc.add_picture(str(images[diagram_index]), width=Cm(15.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = False
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(f"图 {diagram_index + 1}：流程图")
                set_run_font(run, size=10, bold=True)
                run.font.color.rgb = RGBColor(75, 85, 99)
                diagram_index += 1
                in_mermaid = False
            continue
        if raw.startswith("|"):
            table_rows.append([cell.strip() for cell in raw.strip("|").split("|")])
            continue
        flush_table()
        if not raw:
            continue
        if raw.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(raw[2:])
            set_run_font(run, size=22, bold=True, color="0B2545")
            continue
        if raw.startswith("## "):
            doc.add_paragraph(raw[3:], style="Heading 1")
            continue
        if raw.startswith("### "):
            doc.add_paragraph(raw[4:], style="Heading 2")
            continue
        stripped = raw.lstrip()
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_text(p, stripped[2:])
            continue
        match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.30)
            p.paragraph_format.first_line_indent = Inches(-0.30)
            p.paragraph_format.space_after = Pt(8)
            marker = p.add_run(f"{match.group(1)}.  ")
            set_run_font(marker)
            add_formatted_text(p, match.group(2))
            continue
        p = doc.add_paragraph()
        add_formatted_text(p, raw)

    flush_table()
    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build()
