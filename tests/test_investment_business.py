# encoding:utf-8
import base64
from datetime import UTC, datetime, timedelta
from io import BytesIO
from collections import OrderedDict
import json
import os
from pathlib import Path
import subprocess
from types import SimpleNamespace
import sys
from uuid import uuid4
from zipfile import ZipFile

import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url


def test_investment_schema_declares_all_tables():
    from business.investment.schema import metadata

    assert {
        "investment_users",
        "investment_admin_users",
        "investment_admin_sessions",
        "investment_request_records",
        "investment_daily_contents",
        "investment_output_files",
        "investment_cache_entries",
        "investment_configs",
        "investment_stock_symbols",
        "investment_operation_audits",
    }.issubset(metadata.tables)

    assert {
        "created_by_admin_id",
        "created_by_username",
        "updated_by_admin_id",
        "updated_by_username",
        "deleted_at",
        "deleted_by_admin_id",
        "deleted_by_username",
        "delete_reason",
    }.issubset(metadata.tables["investment_users"].columns.keys())
    assert {
        "created_by_admin_id",
        "created_by_username",
        "created_by_role",
        "updated_by_admin_id",
        "updated_by_username",
        "updated_by_role",
        "published_by_admin_id",
        "published_by_username",
        "published_by_role",
    }.issubset(metadata.tables["investment_daily_contents"].columns.keys())
    assert {
        "operator_admin_id",
        "operator_username",
        "operator_role",
        "operation_category",
        "result_status",
        "error_code",
        "error_message",
        "elapsed_ms",
        "before_state",
        "after_state",
        "request_ip",
        "user_agent",
    }.issubset(metadata.tables["investment_operation_audits"].columns.keys())
    assert {
        "updated_by_admin_id",
        "updated_by_username",
        "updated_by_role",
    }.issubset(metadata.tables["investment_configs"].columns.keys())
    assert "owner_type" in metadata.tables["investment_output_files"].columns


def test_investment_auth_service_hashes_passwords_and_checks_role_permissions(investment_env):
    from business.investment.auth_service import (
        authenticate_admin,
        create_admin_session,
        create_admin_user,
        get_admin_session,
        hash_password,
        require_permission,
        verify_password,
    )

    password_hash = hash_password("secret-pass")

    assert password_hash.startswith("pbkdf2_sha256$")
    assert "secret-pass" not in password_hash
    assert verify_password("secret-pass", password_hash) is True
    assert verify_password("wrong-pass", password_hash) is False

    user_id = create_admin_user("operator-a", "secret-pass", role="content_operator")
    assert authenticate_admin("operator-a", "wrong-pass") is None
    admin = authenticate_admin("operator-a", "secret-pass")
    assert admin is not None
    assert admin.id == user_id
    assert admin.role == "content_operator"

    token = create_admin_session(admin)
    session = get_admin_session(token)
    assert session is not None
    assert session.username == "operator-a"
    assert session.role == "content_operator"
    assert require_permission(session, "content.upload").allowed is True
    assert require_permission(session, "content.publish").allowed is True
    assert require_permission(session, "content.write").allowed is True
    assert require_permission(session, "content.effective").allowed is True
    assert require_permission(session, "customers.write").allowed is False
    assert require_permission(session, "admin_users.write").allowed is False
    assert require_permission(session, "audits.read").allowed is False
    assert require_permission(session, "config.write").allowed is False

    with pytest.raises(ValueError, match="unsupported admin role"):
        create_admin_user("legacy-poster-a", "poster-pass", role="poster")
    with pytest.raises(ValueError, match="unsupported admin role"):
        create_admin_user("legacy-tech-a", "tech-pass", role="technical_admin")


def test_technical_operator_only_has_config_permissions(investment_env):
    from business.investment.auth_service import authenticate_admin, create_admin_user, require_permission

    create_admin_user("tech-ops-a", "tech-pass", role="technical_operator")

    admin = authenticate_admin("tech-ops-a", "tech-pass")
    assert admin is not None
    assert admin.role == "technical_operator"
    assert require_permission(admin, "config.read").allowed is True
    assert require_permission(admin, "config.write").allowed is True
    assert require_permission(admin, "customers.read").allowed is False
    assert require_permission(admin, "admin_users.read").allowed is False
    assert require_permission(admin, "content.read").allowed is False
    assert require_permission(admin, "skills.read").allowed is False
    assert require_permission(admin, "stocks.read").allowed is False
    assert require_permission(admin, "health.read").allowed is False


def test_investment_migrations_seed_default_admin_and_posters(investment_env):
    from business.investment.auth_service import authenticate_admin

    admin = authenticate_admin("admin", "password")
    assert admin is not None
    assert admin.role == "admin"

    for username in ("poster1", "poster2", "poster3"):
        poster = authenticate_admin(username, "password")
        assert poster is not None
        assert poster.role == "content_operator"


def test_admin_user_service_lists_updates_and_resets_password(investment_env):
    from business.investment.auth_service import (
        authenticate_admin,
        create_admin_user,
        list_admin_users,
        reset_admin_password,
        update_admin_user,
    )

    create_admin_user("ops-a", "old-pass", role="content_operator")

    users = {user.username: user for user in list_admin_users()}
    assert users["ops-a"].role == "content_operator"
    assert users["ops-a"].enabled is True
    assert "password_hash" not in users["ops-a"].__dict__

    update_admin_user("ops-a", role="admin", enabled=False)
    assert authenticate_admin("ops-a", "old-pass") is None
    disabled = {user.username: user for user in list_admin_users()}["ops-a"]
    assert disabled.role == "admin"
    assert disabled.enabled is False

    update_admin_user("ops-a", enabled=True)
    reset_admin_password("ops-a", "new-pass")
    assert authenticate_admin("ops-a", "old-pass") is None
    reset_user = authenticate_admin("ops-a", "new-pass")
    assert reset_user is not None
    assert reset_user.role == "admin"

    with pytest.raises(ValueError, match="not found"):
        update_admin_user("missing-admin", enabled=False)
    with pytest.raises(ValueError, match="not found"):
        reset_admin_password("missing-admin", "new-pass")

    from business.investment.db import connect
    from business.investment.schema import investment_admin_users

    update_admin_user("ops-a", role="content_operator")
    create_admin_user("solo-admin", "admin-pass", role="admin")
    with connect() as conn:
        conn.execute(
            investment_admin_users.update()
            .where(investment_admin_users.c.username != "solo-admin")
            .values(enabled=0)
        )
    with pytest.raises(ValueError, match="last enabled admin"):
        update_admin_user("solo-admin", role="content_operator")
    with pytest.raises(ValueError, match="last enabled admin"):
        update_admin_user("solo-admin", enabled=False)


def test_web_investment_api_enforces_admin_roles(investment_env, monkeypatch):
    from business.investment.audit_service import list_operation_audits
    from business.investment.auth_service import authenticate_admin, create_admin_session, create_admin_user
    from business.investment.records import get_content_record
    from channel.web import web_channel
    from channel.web.web_channel import (
        InvestmentAdminUserPasswordHandler,
        InvestmentAdminUserStatusHandler,
        InvestmentAdminUsersHandler,
        InvestmentConfigHandler,
        InvestmentDailyContentHandler,
        InvestmentHealthHandler,
        InvestmentUsersHandler,
    )

    create_admin_user("admin-a", "admin-pass", role="admin")
    create_admin_user("operator-a", "operator-pass", role="content_operator")
    create_admin_user("tech-a", "tech-pass", role="technical_operator")
    admin_token = create_admin_session(authenticate_admin("admin-a", "admin-pass"))
    operator_token = create_admin_session(authenticate_admin("operator-a", "operator-pass"))
    tech_token = create_admin_session(authenticate_admin("tech-a", "tech-pass"))

    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web, "input", lambda **kwargs: SimpleNamespace(openid="", enabled=""))

    def use_token(token):
        monkeypatch.setattr(web_channel.web, "cookies", lambda: {"cow_investment_session": token})

    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"openid": "openid-a", "name": "Alice", "allowed_services": ["全部"]}).encode("utf-8"),
    )
    use_token(operator_token)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)
    with pytest.raises(web_channel.web.HTTPError) as denied_error:
        InvestmentUsersHandler().POST()
    denied = json.loads(denied_error.value.data)
    assert denied["status"] == "error"
    assert denied["code"] == "permission_denied"

    use_token(admin_token)
    created = json.loads(InvestmentUsersHandler().POST())
    assert created["status"] == "success"
    assert created["action"] == "created"

    use_token(tech_token)
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"configs": {"router.enable_agent_fallback": True}}).encode("utf-8"),
    )
    config = json.loads(InvestmentConfigHandler().POST())
    assert config["status"] == "success"
    with pytest.raises(web_channel.web.HTTPError) as tech_health_error:
        InvestmentHealthHandler().GET()
    tech_health_denied = json.loads(tech_health_error.value.data)
    assert tech_health_denied["status"] == "error"
    assert tech_health_denied["permission"] == "health.read"

    use_token(operator_token)
    with pytest.raises(web_channel.web.HTTPError) as poster_config_error:
        InvestmentConfigHandler().GET()
    poster_denied = json.loads(poster_config_error.value.data)
    assert poster_denied["status"] == "error"
    assert poster_denied["code"] == "permission_denied"
    assert poster_denied["permission"] == "config.read"

    with pytest.raises(web_channel.web.HTTPError) as poster_admin_error:
        InvestmentAdminUsersHandler().GET()
    poster_admin_denied = json.loads(poster_admin_error.value.data)
    assert poster_admin_denied["status"] == "error"
    assert poster_admin_denied["code"] == "permission_denied"

    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"service_type": "利率", "source_text": "rate source", "operator": "forged-admin"}).encode("utf-8"),
    )
    content = json.loads(InvestmentDailyContentHandler().POST())
    assert content["status"] == "success"
    assert get_content_record(content["content_id"]).operator == "operator-a"

    use_token(admin_token)
    monkeypatch.setattr(web_channel.web, "input", lambda **kwargs: SimpleNamespace())
    admin_users = json.loads(InvestmentAdminUsersHandler().GET())
    assert admin_users["status"] == "success"
    assert any(item["username"] == "admin-a" and item["role"] == "admin" for item in admin_users["users"])
    assert all("password" not in item for item in admin_users["users"])

    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"username": "ops-web", "password": "ops-pass", "role": "technical_operator"}).encode("utf-8"),
    )
    created_admin = json.loads(InvestmentAdminUsersHandler().POST())
    assert created_admin["status"] == "success"
    assert created_admin["action"] == "created"
    assert authenticate_admin("ops-web", "ops-pass").role == "technical_operator"

    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"username": "ops-web", "role": "admin"}).encode("utf-8"),
    )
    updated_admin = json.loads(InvestmentAdminUsersHandler().POST())
    assert updated_admin["status"] == "success"
    assert updated_admin["action"] == "updated"
    assert authenticate_admin("ops-web", "ops-pass").role == "admin"

    missing_admin = json.loads(InvestmentAdminUserStatusHandler().POST("missing-web", "disable"))
    assert missing_admin["status"] == "error"

    disabled_admin = json.loads(InvestmentAdminUserStatusHandler().POST("ops-web", "disable"))
    assert disabled_admin["status"] == "success"
    assert authenticate_admin("ops-web", "ops-pass") is None

    enabled_admin = json.loads(InvestmentAdminUserStatusHandler().POST("ops-web", "enable"))
    assert enabled_admin["status"] == "success"
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"password": "reset-pass"}).encode("utf-8"),
    )
    reset_admin = json.loads(InvestmentAdminUserPasswordHandler().POST("ops-web"))
    assert reset_admin["status"] == "success"
    assert authenticate_admin("ops-web", "reset-pass").role == "admin"

    audits = list_operation_audits(limit=20, target_type="admin_user", target_id="ops-web")
    actions = [audit.action for audit in audits]
    assert "admin_user.create" in actions
    assert "admin_user.update" in actions
    assert "admin_user.disable" in actions
    assert "admin_user.reset_password" in actions


def test_investment_auth_me_allows_content_operator_without_customer_or_audit_permission(investment_env, monkeypatch):
    from business.investment.auth_service import authenticate_admin, create_admin_session, create_admin_user
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentAuthMeHandler

    create_admin_user("operator-a", "operator-pass", role="content_operator")
    token = create_admin_session(authenticate_admin("operator-a", "operator-pass"))

    monkeypatch.setattr(web_channel.web, "cookies", lambda: {"cow_investment_session": token})
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)

    payload = json.loads(InvestmentAuthMeHandler().GET())

    assert payload["status"] == "success"
    assert payload["admin"]["role"] == "content_operator"
    assert "content.upload" in payload["admin"]["permissions"]
    assert "content.publish" in payload["admin"]["permissions"]
    assert "content.write" in payload["admin"]["permissions"]
    assert "customers.write" not in payload["admin"]["permissions"]
    assert "audits.read" not in payload["admin"]["permissions"]


def test_web_investment_auth_falls_back_to_web_password_until_admin_exists(investment_env, monkeypatch):
    from business.investment.db import connect
    from business.investment.schema import investment_admin_sessions, investment_admin_users
    from business.investment.auth_service import create_admin_user
    from channel.web import web_channel

    with connect() as conn:
        conn.execute(investment_admin_sessions.delete())
        conn.execute(investment_admin_users.delete())

    monkeypatch.setattr(web_channel, "_check_auth", lambda: True)
    monkeypatch.setattr(web_channel.web, "cookies", lambda: {})

    fallback = web_channel._require_investment_permission("customers.write")
    assert fallback.role == "admin"
    assert fallback.bootstrap is True

    create_admin_user("admin-a", "admin-pass", role="admin")
    denied = web_channel._investment_permission_error("customers.write")
    assert denied["status"] == "error"
    assert denied["code"] == "unauthorized"


def test_admin_login_with_web_password_enabled_also_authenticates_console(investment_env, monkeypatch):
    from business.investment.auth_service import create_admin_user
    from channel.web import web_channel
    from channel.web.web_channel import AuthCheckHandler, AuthLoginHandler

    create_admin_user("admin-a", "admin-pass", role="admin")
    cookies = {}
    headers = []

    monkeypatch.setattr(web_channel, "conf", lambda: {"web_password": "legacy-pass", "web_session_expire_days": 30})
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: headers.append(args))
    monkeypatch.setattr(
        web_channel.web,
        "setcookie",
        lambda name, value, **_kwargs: cookies.__setitem__(name, value),
    )
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({"username": "admin-a", "password": "admin-pass"}).encode("utf-8"),
    )

    login_payload = json.loads(AuthLoginHandler().POST())

    assert login_payload["status"] == "success"
    assert cookies["cow_investment_session"]
    assert cookies["cow_auth_token"]

    monkeypatch.setattr(web_channel.web, "cookies", lambda: cookies)
    check_payload = json.loads(AuthCheckHandler().GET())

    assert check_payload["status"] == "success"
    assert check_payload["authenticated"] is True
    web_channel._require_auth()


def test_chat_page_redirects_to_login_when_console_session_missing(monkeypatch):
    from channel.web import web_channel
    from channel.web.web_channel import ChatHandler

    redirects = []

    def fake_seeother(target):
        redirects.append(target)
        raise RuntimeError(target)

    monkeypatch.setattr(web_channel, "_check_console_auth", lambda: False)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web, "seeother", fake_seeother)
    monkeypatch.setattr(web_channel.web.ctx, "fullpath", "/chat?view=invest-records", raising=False)

    with pytest.raises(RuntimeError):
        ChatHandler().GET()

    assert redirects == ["/login?next=%2Fchat%3Fview%3Dinvest-records"]


def test_web_message_api_requires_console_admin_login(monkeypatch):
    from channel.web import web_channel
    from channel.web.web_channel import MessageHandler

    class FakeWebChannel:
        def post_message(self):
            return "unexpected"

    monkeypatch.setattr(web_channel, "_check_console_auth", lambda: False)
    monkeypatch.setattr(web_channel, "_require_auth", lambda: None)
    monkeypatch.setattr(web_channel, "WebChannel", lambda: FakeWebChannel())
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    with pytest.raises(web_channel.web.HTTPError) as exc_info:
        MessageHandler().POST()

    payload = json.loads(exc_info.value.data)
    assert payload["status"] == "error"
    assert payload["message"] == "Unauthorized"


def test_login_page_redirects_authenticated_user_to_next_path(monkeypatch):
    from channel.web import web_channel
    from channel.web.web_channel import LoginPageHandler

    redirects = []

    def fake_seeother(target):
        redirects.append(target)
        raise RuntimeError(target)

    monkeypatch.setattr(web_channel, "_check_console_auth", lambda: True)
    monkeypatch.setattr(web_channel.web, "input", lambda **kwargs: SimpleNamespace(next="/chat"))
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web, "seeother", fake_seeother)

    with pytest.raises(RuntimeError):
        LoginPageHandler().GET()

    assert redirects == ["/chat"]


def test_content_operator_session_denies_records_and_cache(investment_env, monkeypatch):
    from business.investment.auth_service import authenticate_admin, create_admin_session, create_admin_user
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentCacheHandler, InvestmentRequestRecordsHandler

    create_admin_user("operator-a", "operator-pass", role="content_operator")
    token = create_admin_session(authenticate_admin("operator-a", "operator-pass"))

    monkeypatch.setattr(web_channel.web, "cookies", lambda: {"cow_investment_session": token})
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(limit="20"))
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    with pytest.raises(web_channel.web.HTTPError) as records_denied_error:
        InvestmentRequestRecordsHandler().GET()
    records_denied = json.loads(records_denied_error.value.data)
    assert records_denied["status"] == "error"
    assert records_denied["code"] == "permission_denied"
    assert records_denied["permission"] == "records.read"

    with pytest.raises(web_channel.web.HTTPError) as denied_error:
        InvestmentCacheHandler().GET()
    denied = json.loads(denied_error.value.data)
    assert denied["status"] == "error"
    assert denied["code"] == "permission_denied"
    assert denied["permission"] == "cache.read"


def _xlsx_bytes(headers, rows):
    values = [list(map(str, headers))]
    values.extend([["" if value is None else str(value) for value in row] for row in rows])
    shared_strings = []
    shared_index = {}

    def shared(value):
        if value not in shared_index:
            shared_index[value] = len(shared_strings)
            shared_strings.append(value)
        return shared_index[value]

    def column_name(index):
        name = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            name = chr(65 + remainder) + name
        return name

    sheet_rows = []
    for row_index, row in enumerate(values, start=1):
        cells = []
        for col_index, value in enumerate(row, start=1):
            cell_ref = f"{column_name(col_index)}{row_index}"
            cells.append(f'<c r="{cell_ref}" t="s"><v>{shared(value)}</v></c>')
        sheet_rows.append(f'<row r="{row_index}">{"".join(cells)}</row>')

    shared_xml = "".join(f"<si><t>{value}</t></si>" for value in shared_strings)
    data = BytesIO()
    with ZipFile(data, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/workbook.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Users" sheetId="1" r:id="rId1"/></sheets>
</workbook>""",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData>{"".join(sheet_rows)}</sheetData>
</worksheet>""",
        )
        archive.writestr(
            "xl/sharedStrings.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(shared_strings)}" uniqueCount="{len(shared_strings)}">{shared_xml}</sst>""",
        )
    return data.getvalue()


@pytest.fixture()
def investment_env(tmp_path, monkeypatch):
    from business.investment import db, storage

    base_url = os.environ.get("COWAGENT_TEST_POSTGRES_URL") or db.DEFAULT_DATABASE_URL
    schema_name = f"cowagent_test_{uuid4().hex}"
    url = make_url(base_url)
    schema_url = url.set(
        query={
            **dict(url.query),
            "options": f"-csearch_path={schema_name}",
        },
    ).render_as_string(hide_password=False)

    admin_engine = create_engine(base_url, future=True)
    with admin_engine.begin() as conn:
        conn.execute(text(f'create schema "{schema_name}"'))

    monkeypatch.setenv("COWAGENT_INVESTMENT_DATABASE_URL", schema_url)
    monkeypatch.setenv("COWAGENT_INVESTMENT_STORAGE_ROOT", str(tmp_path / "storage"))
    monkeypatch.delenv("COWAGENT_INVESTMENT_DB_PATH", raising=False)

    db.reset_engine_for_tests()
    storage._MIGRATED_DATABASE_URL = None
    storage.initialize_storage()
    try:
        yield tmp_path
    finally:
        db.reset_engine_for_tests()
        storage._MIGRATED_DATABASE_URL = None
        with admin_engine.begin() as conn:
            conn.execute(text(f'drop schema if exists "{schema_name}" cascade'))
        admin_engine.dispose()


def test_investment_database_url_defaults_to_docker_postgres(monkeypatch):
    from business.investment import db

    monkeypatch.delenv("COWAGENT_INVESTMENT_DATABASE_URL", raising=False)
    monkeypatch.delenv("COWAGENT_INVESTMENT_DB_PATH", raising=False)

    url = db.get_database_url()

    assert url == "postgresql+psycopg://cowagent:cowagent@127.0.0.1:55432/cowagent_investment"


def test_investment_database_url_prefers_postgres_env(investment_env, monkeypatch):
    from business.investment import db

    monkeypatch.setenv("COWAGENT_INVESTMENT_DATABASE_URL", "postgresql+psycopg://u:p@localhost:5432/cowagent")

    assert db.get_database_url() == "postgresql+psycopg://u:p@localhost:5432/cowagent"


def test_investment_database_url_rejects_sqlite_env(monkeypatch):
    from business.investment import db

    monkeypatch.setenv("COWAGENT_INVESTMENT_DATABASE_URL", "sqlite:///tmp/investment.db")

    with pytest.raises(ValueError, match="PostgreSQL"):
        db.get_database_url()


def test_investment_database_url_rejects_sqlite_config(monkeypatch):
    import config

    from business.investment import db

    monkeypatch.delenv("COWAGENT_INVESTMENT_DATABASE_URL", raising=False)
    monkeypatch.setattr(config, "conf", lambda: {"investment_database_url": "sqlite:///tmp/investment.db"})

    with pytest.raises(ValueError, match="PostgreSQL"):
        db.get_database_url()


def test_storage_initializes_schema_with_alembic_upgrade(tmp_path, monkeypatch):
    from business.investment import migrations, schema, storage

    monkeypatch.setenv("COWAGENT_INVESTMENT_STORAGE_ROOT", str(tmp_path / "storage"))
    calls = []

    monkeypatch.setattr(migrations, "upgrade", lambda revision="head": calls.append(revision))
    monkeypatch.setattr(schema.metadata, "create_all", lambda bind: pytest.fail("initialize_storage must use Alembic"))

    storage.initialize_storage()

    assert calls == ["head"]
    assert (tmp_path / "storage" / "uploads").is_dir()
    assert (tmp_path / "storage" / "generated").is_dir()
    assert (tmp_path / "storage" / "technical-analysis").is_dir()


def test_investment_migration_smoke_creates_schema(investment_env):
    from business.investment import storage
    from business.investment.db import get_engine
    from sqlalchemy import inspect

    storage.initialize_storage()
    inspector = inspect(get_engine())

    assert inspector.has_table("investment_users")
    assert inspector.has_table("investment_stock_symbols")


def test_investment_alembic_runner_exposes_upgrade():
    from business.investment import migrations

    assert callable(migrations.upgrade)
    assert migrations.alembic_config_path().name == "alembic.ini"


def test_config_service_uses_investment_db_connection_helpers():
    from business.investment import config_service

    assert not hasattr(config_service, "get_connection")


def test_user_service_uses_investment_db_connection_helpers():
    from business.investment import user_service

    assert not hasattr(user_service, "get_connection")


def test_records_service_uses_investment_db_connection_helpers():
    from business.investment import records

    assert not hasattr(records, "get_connection")


def test_daily_content_service_uses_investment_db_connection_helpers():
    from business.investment import daily_content

    assert not hasattr(daily_content, "get_connection")


def test_stock_resolver_uses_investment_db_connection_helpers():
    from business.investment import stock_resolver

    assert not hasattr(stock_resolver, "get_connection")


def test_health_uses_investment_db_connection_helpers():
    from business.investment import health

    assert not hasattr(health, "get_connection")


def test_config_service_direct_call_initializes_storage_schema(investment_env):
    from business.investment import db
    from business.investment.config_service import get_config, save_config
    from sqlalchemy import inspect

    db.reset_engine_for_tests()

    save_config("tushare.token", "direct-token", operator_role="admin")

    assert get_config("tushare.token") == "direct-token"
    assert inspect(db.get_engine()).has_table("investment_configs")


def test_config_masks_sensitive_values_and_checks_permissions(investment_env, monkeypatch):
    from business.investment.config_service import (
        can_modify_config,
        get_config,
        mask_sensitive_value,
        save_config,
        safe_log_value,
    )

    monkeypatch.setattr("business.investment.config_service.conf", lambda: {"tushare_token": "fallback-token"})

    assert get_config("tushare.token") == "fallback-token"
    save_config("tushare.token", "ts-1234567890abcdef", operator_role="admin")
    assert get_config("tushare.token") == "ts-1234567890abcdef"
    assert mask_sensitive_value("ts-1234567890abcdef") == "ts-1**********cdef"
    assert "ts-1234567890abcdef" not in safe_log_value("tushare.token", "ts-1234567890abcdef")
    assert mask_sensitive_value("abc") == "***"
    assert can_modify_config("tushare.token", "uploader") is False
    assert can_modify_config("tushare.token", "content_operator") is False
    assert can_modify_config("tushare.token", "technical_operator") is True
    assert can_modify_config("tushare.token", "admin") is True


def test_investment_user_message_uses_reply_config_defaults(investment_env):
    from business.investment.constants import ErrorCode, user_message

    assert user_message(ErrorCode.UNAUTHORIZED) == "您暂未开通该服务，如需开通请联系服务人员。"
    assert user_message(ErrorCode.SYSTEM_ERROR) == "系统暂时繁忙，请稍后重试。"


def test_investment_user_message_can_be_overridden_from_database(investment_env):
    from business.investment.config_service import save_config
    from business.investment.constants import ErrorCode, user_message

    save_config("reply.investment.unauthorized", "请联系客户经理开通权限。", operator_role="admin", operator="pytest")

    assert user_message(ErrorCode.UNAUTHORIZED) == "请联系客户经理开通权限。"
    assert user_message(ErrorCode.SYSTEM_ERROR) == "系统暂时繁忙，请稍后重试。"


def test_web_open_chat_config_is_admin_only(investment_env):
    from business.investment.config_service import can_modify_config, get_config, save_config

    assert get_config("router.enable_web_open_chat", False) is False
    assert can_modify_config("router.enable_web_open_chat", "technical_operator") is False
    assert can_modify_config("router.enable_web_open_chat", "admin") is True

    with pytest.raises(PermissionError):
        save_config("router.enable_web_open_chat", True, operator_role="technical_operator")

    save_config("router.enable_web_open_chat", True, operator_role="admin")
    assert get_config("router.enable_web_open_chat") is True


def test_investment_config_rejects_model_and_wechatmp_keys_without_persisting(investment_env):
    from sqlalchemy import select

    from business.investment import db
    from business.investment.config_service import get_config, save_config, save_configs
    from business.investment.schema import investment_configs

    forbidden = {
        "model.name": "investment-model",
        "model.api_key": "sk-investment-secret",
        "wechatmp.app_id": "wx-investment",
        "wechatmp.token": "wx-token",
    }

    for key, value in forbidden.items():
        with pytest.raises(ValueError, match="Investment config"):
            save_config(key, value, operator_role="admin")

    with pytest.raises(ValueError, match="Investment config"):
        save_configs(forbidden, operator_role="admin")

    with db.connect() as conn:
        rows = conn.execute(select(investment_configs.c.config_key)).fetchall()
    assert {row[0] for row in rows}.isdisjoint(forbidden)
    assert get_config("tushare.token", "") == ""


def test_web_console_config_save_preserves_masked_investment_sensitive_values(investment_env):
    from business.investment.config_service import get_config, get_configs, save_configs

    token = "ts-console-secret-1234567890"
    save_configs({"tushare.token": token, "router.enable_agent_fallback": False}, operator_role="admin")

    masked = get_configs(["tushare.token", "router.enable_agent_fallback"], masked=True)
    save_configs(
        {
            "tushare.token": masked["tushare.token"],
            "router.enable_agent_fallback": True,
        },
        operator_role="admin",
    )

    assert get_config("tushare.token") == token
    assert get_config("router.enable_agent_fallback") is True


def test_investment_skill_loader_reads_builtin_packages_and_excludes_cowagent_skills(investment_env):
    from business.investment.skill_registry import list_investment_skills

    keys = {item["skill_key"] for item in list_investment_skills()}

    assert {"technical-analysis", "rate", "convertible-bond", "signal-card-renderer"} <= keys
    assert "image-generation" not in keys
    assert "knowledge-wiki" not in keys


def test_investment_skill_loader_applies_web_trigger_override(investment_env):
    from business.investment.config_service import save_config
    from business.investment.constants import ServiceType
    from business.investment.skill_registry import match_investment_skill

    save_config("skill.rate.triggers", ["今日利率"], operator_role="admin", operator="pytest")

    assert match_investment_skill("利率") is None
    matched = match_investment_skill("今日利率")
    assert matched is not None
    assert matched.skill_key == "rate"
    assert matched.service_type == ServiceType.RATE


def test_investment_skill_loader_extracts_suffix_target(investment_env):
    from business.investment.config_service import save_config
    from business.investment.constants import ServiceType
    from business.investment.skill_registry import match_investment_skill

    save_config("skill.technical-analysis.triggers", ["走势分析"], operator_role="admin", operator="pytest")

    matched = match_investment_skill("300502.SZ 走势分析")
    assert matched is not None
    assert matched.skill_key == "technical-analysis"
    assert matched.service_type == ServiceType.TECHNICAL_ANALYSIS
    assert matched.target_text == "300502.SZ"


def test_uploaded_investment_skill_package_appears_in_registry(investment_env, tmp_path):
    from business.investment.constants import ServiceType
    from business.investment.skill_registry import list_investment_skills, match_investment_skill
    from business.investment.skill_versions import save_package_upload

    package = tmp_path / "macro.zip"
    skill_md = """---
name: macro-analysis
description: 宏观分析投研 Skill
investment:
  label: 宏观分析
  enabled: true
  routable: true
  service_type: macro_analysis
  match_type: exact
  triggers:
    - 宏观
  handler_type: script
  entry: scripts/macro_analysis.py
  output_mode: text
---
# Macro Analysis
"""
    script = "import json\nprint(json.dumps({'success': True, 'reply_text': 'macro ok', 'output_files': []}, ensure_ascii=False))\n"
    with ZipFile(package, "w") as archive:
        archive.writestr("SKILL.md", skill_md)
        archive.writestr("scripts/macro_analysis.py", script)

    save_package_upload("macro.zip", package.read_bytes(), operator="pytest")

    keys = {item["skill_key"] for item in list_investment_skills()}
    assert "macro-analysis" in keys
    matched = match_investment_skill("宏观")
    assert matched is not None
    assert matched.skill_key == "macro-analysis"
    assert matched.service_type == ServiceType.UNMATCHED


def test_uploaded_investment_skill_package_rejects_unsafe_skill_key(investment_env, tmp_path):
    import pytest

    from business.investment.skill_versions import save_package_upload

    package = tmp_path / "unsafe.zip"
    skill_md = """---
name: ../escape
description: unsafe
investment:
  label: Unsafe
  triggers: [unsafe]
  handler_type: script
  entry: scripts/run.py
---
# Unsafe
"""
    with ZipFile(package, "w") as archive:
        archive.writestr("SKILL.md", skill_md)
        archive.writestr("scripts/run.py", "print('{}')\n")

    with pytest.raises(ValueError, match="unsafe investment skill key"):
        save_package_upload("unsafe.zip", package.read_bytes(), operator="pytest")


def test_uploaded_script_investment_skill_executes_from_route(investment_env, tmp_path):
    from business.investment.constants import ServiceType
    from business.investment.router import handle_text_message
    from business.investment.skill_versions import save_package_upload
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    package = tmp_path / "macro.zip"
    skill_md = """---
name: macro-script
description: 宏观分析投研 Skill
investment:
  label: 宏观脚本
  enabled: true
  routable: true
  service_type: macro_analysis
  match_type: exact
  triggers:
    - 宏观脚本
  handler_type: script
  entry: scripts/macro.py
  output_mode: text
---
# Macro
"""
    script = "import json\nprint(json.dumps({'success': True, 'reply_text': 'macro script ok', 'output_files': []}, ensure_ascii=False))\n"
    with ZipFile(package, "w") as archive:
        archive.writestr("SKILL.md", skill_md)
        archive.writestr("scripts/macro.py", script)

    save_package_upload("macro.zip", package.read_bytes(), operator="pytest")

    reply = handle_text_message("ok", "宏观脚本")

    assert reply.handled is True
    assert reply.success is True
    assert reply.service_type == ServiceType.UNMATCHED
    assert reply.reply_text == "macro script ok"


def test_investment_skill_upload_python_file_creates_version_and_activates_it(investment_env):
    from pathlib import Path

    from business.investment.config_service import get_config
    from business.investment.skill_versions import list_versions, save_upload

    result = save_upload("signal-card-renderer", "render_card.py", b"print('renderer v1')", operator="tester")

    script_path = Path(result["script_path"])
    assert result["version_id"].startswith("skill-")
    assert result["skill_key"] == "signal-card-renderer"
    assert script_path.name == "render_card.py"
    assert script_path.parent.name == "scripts"
    assert script_path.read_text(encoding="utf-8") == "print('renderer v1')"
    assert (script_path.parents[1] / "assets" / "template_ta.html").is_file()
    assert get_config("render.renderer_path") == str(script_path)

    versions = list_versions("signal-card-renderer")
    uploaded = [item for item in versions if item["version_id"] == result["version_id"]][0]
    assert uploaded["active"] is True
    assert uploaded["source"] == "upload"
    assert any(item["version_id"] == "builtin-default" for item in versions)


def test_investment_skill_upload_zip_rejects_path_escape_and_requires_script(investment_env):
    from io import BytesIO
    from zipfile import ZipFile

    import pytest

    from business.investment.skill_versions import save_upload

    unsafe = BytesIO()
    with ZipFile(unsafe, "w") as archive:
        archive.writestr("../escape.py", "bad")
        archive.writestr("scripts/render_card.py", "print('ok')")
    with pytest.raises(ValueError, match="unsafe"):
        save_upload("signal-card-renderer", "skill.zip", unsafe.getvalue(), operator="tester")

    missing = BytesIO()
    with ZipFile(missing, "w") as archive:
        archive.writestr("README.md", "no script")
    with pytest.raises(ValueError, match="render_card.py"):
        save_upload("signal-card-renderer", "skill.zip", missing.getvalue(), operator="tester")


def test_investment_skill_version_activation_switches_between_upload_and_builtin(investment_env):
    from business.investment.config_service import get_config
    from business.investment.skill_versions import activate_version, list_versions, save_upload

    uploaded = save_upload("signal-card-renderer", "render_card.py", b"print('renderer v2')", operator="tester")

    activate_version("signal-card-renderer", "builtin-default", operator="tester")

    assert get_config("render.renderer_path", "") == ""
    builtin = [item for item in list_versions("signal-card-renderer") if item["version_id"] == "builtin-default"][0]
    assert builtin["active"] is True

    activate_version("signal-card-renderer", uploaded["version_id"], operator="tester")

    assert get_config("render.renderer_path") == uploaded["script_path"]


def test_investment_skill_uploaded_version_can_be_deleted_and_active_delete_falls_back_to_builtin(investment_env):
    import pytest

    from business.investment.config_service import get_config
    from business.investment.skill_versions import delete_version, list_versions, save_upload

    uploaded = save_upload("technical-analysis", "analyze_universal.py", b"print('ta delete')", operator="tester")
    assert get_config("technical_analysis.skill_path") == uploaded["script_path"]

    deleted = delete_version("technical-analysis", uploaded["version_id"], operator="tester")

    assert deleted["version_id"] == uploaded["version_id"]
    assert get_config("technical_analysis.skill_path", "") == ""
    assert uploaded["version_id"] not in [item["version_id"] for item in list_versions("technical-analysis")]

    with pytest.raises(ValueError, match="builtin"):
        delete_version("technical-analysis", "builtin-default", operator="tester")


def test_investment_skill_versions_include_loaded_skills(investment_env):
    from pathlib import Path

    from business.investment.config_service import get_config
    from business.investment.skill_versions import activate_version, list_all_skills, save_upload

    technical = save_upload("technical-analysis", "analyze_universal.py", b"print('ta v1')", operator="tester")
    renderer = save_upload("signal-card-renderer", "render_card.py", b"print('renderer v1')", operator="tester")

    assert Path(technical["script_path"]).name == "analyze_universal.py"
    assert Path(renderer["script_path"]).name == "render_card.py"
    assert get_config("technical_analysis.skill_path") == technical["script_path"]
    assert get_config("render.renderer_path") == renderer["script_path"]

    payload = list_all_skills()
    by_key = {item["skill"]["skill_key"]: item for item in payload}
    assert {"technical-analysis", "signal-card-renderer", "rate", "convertible-bond"} <= set(by_key)
    assert by_key["technical-analysis"]["versions"][0]["version_id"] == "builtin-default"
    assert by_key["signal-card-renderer"]["versions"][0]["version_id"] == "builtin-default"
    assert by_key["rate"]["versions"] == []
    assert by_key["convertible-bond"]["versions"] == []

    activate_version("technical-analysis", "builtin-default", operator="tester")
    assert get_config("technical_analysis.skill_path", "") == ""


def test_web_investment_skill_handlers_list_upload_and_activate_versions(investment_env, monkeypatch):
    from business.investment.config_service import get_config
    from channel.web import web_channel
    from channel.web.web_channel import (
        InvestmentSkillActivateHandler,
        InvestmentSkillDeleteHandler,
        InvestmentSkillUploadHandler,
        InvestmentSkillVersionsHandler,
    )

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)

    get_payload = json.loads(InvestmentSkillVersionsHandler().GET())

    assert get_payload["status"] == "success"
    assert {"technical-analysis", "signal-card-renderer", "rate", "convertible-bond"} <= {
        item["skill"]["skill_key"] for item in get_payload["skills"]
    }

    upload_file = SimpleNamespace(filename="analyze_universal.py", value=b"print('web technical')")
    monkeypatch.setattr(web_channel, "_raw_web_input", lambda: {"file": upload_file, "operator": "tester"})

    upload_payload = json.loads(InvestmentSkillUploadHandler().POST("technical-analysis"))

    assert upload_payload["status"] == "success"
    assert upload_payload["version"]["active"] is True
    assert get_config("technical_analysis.skill_path") == upload_payload["version"]["script_path"]

    monkeypatch.setattr(web_channel.web, "data", lambda: json.dumps({"operator": "tester"}).encode("utf-8"))
    activate_payload = json.loads(InvestmentSkillActivateHandler().POST("technical-analysis", "builtin-default"))

    assert activate_payload["status"] == "success"
    assert activate_payload["version"]["active"] is True
    assert get_config("technical_analysis.skill_path", "") == ""

    renderer_file = SimpleNamespace(filename="render_card.py", value=b"print('web renderer')")
    monkeypatch.setattr(web_channel, "_raw_web_input", lambda: {"file": renderer_file, "operator": "tester"})
    renderer_payload = json.loads(InvestmentSkillUploadHandler().POST("signal-card-renderer"))
    deleted_payload = json.loads(
        InvestmentSkillDeleteHandler().POST("signal-card-renderer", renderer_payload["version"]["version_id"])
    )

    assert deleted_payload["status"] == "success"
    assert deleted_payload["deleted"]["version_id"] == renderer_payload["version"]["version_id"]
    assert get_config("render.renderer_path", "") == ""


def test_web_user_disable_button_updates_permission_path(investment_env, monkeypatch):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user, get_user_by_openid
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUserDisableHandler

    create_user("disable-button-openid", enabled=True, allowed_services=[ServiceType.ALL])
    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr("business.investment.technical_analysis.run_technical_analysis", lambda *_args, **_kwargs: pytest.fail("disabled user must not enter technical analysis"))

    payload = json.loads(InvestmentUserDisableHandler().POST("disable-button-openid"))
    reply = handle_text_message("disable-button-openid", "300502.SZ 技术分析")

    assert payload["status"] == "success"
    assert get_user_by_openid("disable-button-openid").enabled is False
    assert reply.success is False
    assert reply.error_code == ErrorCode.USER_DISABLED
    assert reply.reply_text == "您的服务已停用，如需恢复请联系服务人员。"


def test_web_customer_search_enable_and_audits_use_customer_permissions(investment_env, monkeypatch):
    from business.investment.audit_service import list_operation_audits
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user, get_user_by_openid
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUserStatusHandler, InvestmentUsersHandler

    _login_default_investment_admin(monkeypatch, username="audit-admin")
    create_user(
        "search-openid",
        name="Search Alice",
        institution="North Fund",
        mobile="13800138000",
        enabled=False,
        allowed_services=[ServiceType.RATE],
    )
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(keyword="North", enabled=""))
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    listed = json.loads(InvestmentUsersHandler().GET())
    assert [user["openid"] for user in listed["users"]] == ["search-openid"]

    enabled = json.loads(InvestmentUserStatusHandler().POST("search-openid", "enable"))
    assert enabled["status"] == "success"
    assert get_user_by_openid("search-openid").enabled is True

    audits = list_operation_audits(limit=10, target_type="customer", target_id="search-openid")
    assert [audit.action for audit in audits][:1] == ["customer.enable"]
    assert audits[0].operator == "audit-admin"


def test_web_customer_create_audit_binds_session_admin_not_body_operator(investment_env, monkeypatch):
    from business.investment.audit_service import list_operation_audits
    from business.investment.auth_service import authenticate_admin
    from business.investment.db import connect
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUsersHandler

    _login_default_investment_admin(monkeypatch, username="session-admin")
    admin = authenticate_admin("session-admin", "password")
    assert admin is not None
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({
            "openid": "actor-openid",
            "name": "Actor Customer",
            "allowed_services": ["全部"],
            "operator": "spoofed-operator",
        }).encode("utf-8"),
    )

    payload = json.loads(InvestmentUsersHandler().POST())

    assert payload["status"] == "success"
    audit = list_operation_audits(limit=1, target_type="customer", target_id="actor-openid")[0]
    assert audit.operator == "session-admin"
    assert audit.operator_admin_id == admin.id
    assert audit.operator_username == "session-admin"
    assert audit.operator_role == "admin"
    assert audit.operation_category == "customer"
    assert audit.result_status == "success"
    assert audit.error_code == ""
    assert audit.error_message == ""
    assert audit.before_state == {}
    assert audit.after_state["openid"] == "actor-openid"
    assert audit.after_state["name"] == "Actor Customer"
    with connect() as conn:
        row = conn.execute(
            text(
                "select created_by_admin_id, created_by_username, updated_by_admin_id, updated_by_username "
                "from investment_users where openid = 'actor-openid'"
            )
        ).fetchone()
    assert row is not None
    assert row.created_by_admin_id == admin.id
    assert row.created_by_username == "session-admin"
    assert row.updated_by_admin_id == admin.id
    assert row.updated_by_username == "session-admin"


def test_web_user_management_apis_support_keyword_and_pagination(investment_env, monkeypatch):
    from business.investment.auth_service import create_admin_user
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentAdminUsersHandler, InvestmentUsersHandler

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    for index in range(5):
        create_user(
            f"customer-page-{index}",
            name=f"Paged Customer {index}",
            institution="North Fund" if index < 4 else "South Fund",
            mobile=f"1380000000{index}",
            allowed_services=[ServiceType.ALL],
        )

    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(keyword="North", enabled="", page="2", page_size="2"),
    )
    customers = json.loads(InvestmentUsersHandler().GET())

    assert customers["status"] == "success"
    assert customers["pagination"] == {"page": 2, "page_size": 2, "total": 4, "total_pages": 2}
    assert len(customers["users"]) == 2
    assert all("North" in user["institution"] for user in customers["users"])

    for index in range(5):
        role = "admin" if index == 0 else "content_operator"
        create_admin_user(f"ops-page-{index}", "password", role=role)

    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(keyword="ops-page", page="2", page_size="2"),
    )
    admins = json.loads(InvestmentAdminUsersHandler().GET())

    assert admins["status"] == "success"
    assert admins["pagination"] == {"page": 2, "page_size": 2, "total": 5, "total_pages": 3}
    assert [user["username"] for user in admins["users"]] == ["ops-page-2", "ops-page-3"]


def test_web_customer_keyword_search_keeps_rows_and_total_consistent(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUsersHandler

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    create_user("1o92zVw4BYH2qai", name="A", allowed_services=[ServiceType.ALL])
    create_user("o92zVw4BYH2qaib", name="B", allowed_services=[ServiceType.ALL])
    create_user("runtime-user-62b63e3cca25", name="Runtime", allowed_services=[ServiceType.ALL])
    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(keyword="92", enabled="", page="1", page_size="20"),
    )

    payload = json.loads(InvestmentUsersHandler().GET())

    assert payload["status"] == "success"
    assert payload["pagination"]["total"] == 2
    assert [user["openid"] for user in payload["users"]] == ["o92zVw4BYH2qaib", "1o92zVw4BYH2qai"]


def test_web_customer_keyword_search_supports_field_categories(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUsersHandler

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    create_user("openid-name-hit", name="分类客户", mobile="13800000000", allowed_services=[ServiceType.ALL])
    create_user("分类-openid-only", name="OpenID Only", mobile="13900000000", allowed_services=[ServiceType.ALL])
    create_user("service-rate-hit", name="普通客户", mobile="13700000000", allowed_services=[ServiceType.RATE])
    create_user("name-rate-only", name="利率客户", mobile="13600000000", allowed_services=[ServiceType.ALL])

    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(keyword="分类", keyword_field="name", enabled="", page="1", page_size="20"),
    )
    name_payload = json.loads(InvestmentUsersHandler().GET())

    assert name_payload["status"] == "success"
    assert name_payload["pagination"]["total"] == 1
    assert [user["openid"] for user in name_payload["users"]] == ["openid-name-hit"]

    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(keyword="利率", keyword_field="service", enabled="", page="1", page_size="20"),
    )
    service_payload = json.loads(InvestmentUsersHandler().GET())

    assert service_payload["status"] == "success"
    assert service_payload["pagination"]["total"] == 1
    assert [user["openid"] for user in service_payload["users"]] == ["service-rate-hit"]

    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(keyword="利率", keyword_field="all", enabled="", page="1", page_size="20"),
    )
    all_payload = json.loads(InvestmentUsersHandler().GET())

    assert all_payload["status"] == "success"
    assert all_payload["pagination"]["total"] == 2


def test_router_authenticates_before_parsing_unmatched_input(investment_env, monkeypatch):
    import pytest

    from business.investment.constants import ErrorCode, ServiceType
    import business.investment.router as router
    from business.investment.router import DEFAULT_UNMATCHED_PROMPT, handle_text_message
    from business.investment.user_service import create_user

    monkeypatch.setattr(router, "parse_route", lambda _raw_input: pytest.fail("unauthorized input must not be parsed"))

    denied = handle_text_message("missing-openid", "hello")

    assert denied.success is False
    assert denied.error_code == ErrorCode.UNAUTHORIZED
    assert denied.reply_text == "您暂未开通该服务，如需开通请联系服务人员。"

    create_user("ok-unmatched-openid", enabled=True, allowed_services=[ServiceType.RATE])
    monkeypatch.setattr(
        router,
        "parse_route",
        lambda raw_input: router.RouteResult(False, ServiceType.UNMATCHED, raw_input, error_code=ErrorCode.INPUT_ERROR),
    )

    miss = handle_text_message("ok-unmatched-openid", "hello")

    assert miss.success is False
    assert miss.error_code == ErrorCode.INPUT_ERROR
    assert miss.reply_text == DEFAULT_UNMATCHED_PROMPT


def test_web_user_edit_updates_existing_user_permissions(investment_env, monkeypatch):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.user_service import create_user, verify_permission
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUsersHandler

    create_user("edit-button-openid", enabled=True, allowed_services=[ServiceType.ALL])
    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "env", {"CONTENT_TYPE": "application/json"}, raising=False)
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps(
            {
                "openid": "edit-button-openid",
                "name": "Edited",
                "institution": "Edited Inst",
                "enabled": True,
                "allowed_services": ["利率"],
            },
            ensure_ascii=False,
        ).encode("utf-8"),
    )

    payload = json.loads(InvestmentUsersHandler().POST())

    assert payload["status"] == "success"
    assert payload["action"] == "updated"
    assert verify_permission("edit-button-openid", ServiceType.RATE).allowed is True
    technical = verify_permission("edit-button-openid", ServiceType.TECHNICAL_ANALYSIS)
    assert technical.allowed is False
    assert technical.error_code == ErrorCode.UNAUTHORIZED


def test_investment_skill_settings_post_updates_triggers_and_enabled(investment_env, monkeypatch):
    from business.investment.config_service import get_config
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentSkillSettingsHandler

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps(
            {
                "enabled": False,
                "triggers": ["今日利率", "利率观察"],
                "operator": "pytest",
            },
            ensure_ascii=False,
        ).encode("utf-8"),
    )

    payload = json.loads(InvestmentSkillSettingsHandler().POST("rate"))

    assert payload["status"] == "success"
    assert get_config("skill.rate.enabled") is False
    assert get_config("skill.rate.triggers") == ["今日利率", "利率观察"]


def _login_default_investment_admin(monkeypatch, *, username="admin", password="password"):
    from business.investment.auth_service import authenticate_admin, create_admin_session, create_admin_user
    from channel.web import web_channel

    admin = authenticate_admin(username, password)
    if admin is None:
        create_admin_user(username, password, role="admin")
        admin = authenticate_admin(username, password)
    assert admin is not None
    token = create_admin_session(admin)
    monkeypatch.setattr(web_channel.web, "cookies", lambda: {"cow_investment_session": token})
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)
    return token


def _call_investment_json_handler(monkeypatch, handler, *, params=None, body=None):
    from channel.web import web_channel

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(**(params or {})))
    monkeypatch.setattr(web_channel.web, "data", lambda: json.dumps(body or {}, ensure_ascii=False).encode("utf-8"))
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)

    raw = handler()
    payload = json.loads(raw)
    return payload


def test_web_investment_config_returns_masked_tushare_token(investment_env, monkeypatch):
    from business.investment.config_service import save_config
    from channel.web.web_channel import InvestmentConfigHandler

    save_config("tushare.token", "ts-web-secret-1234567890", operator_role="admin")

    payload = _call_investment_json_handler(monkeypatch, InvestmentConfigHandler().GET)

    assert payload["status"] == "success"
    assert payload["configs"]["tushare.token"] == "ts-w**********7890"
    assert "ts-web-secret-1234567890" not in json.dumps(payload, ensure_ascii=False)


def test_web_investment_config_returns_reply_text_metadata(investment_env, monkeypatch):
    from channel.web.web_channel import InvestmentConfigHandler

    payload = _call_investment_json_handler(monkeypatch, InvestmentConfigHandler().GET)

    assert payload["status"] == "success"
    assert "reply_texts" in payload
    assert any(group["title"] == "公众号处理状态" for group in payload["reply_texts"]["groups"])
    assert payload["reply_texts"]["definitions"]["reply.wechatmp.technical_ack"]["label"] == "技术分析开始生成提示"
    assert "reply.wechatmp.technical_ack" in payload["configs"]
    assert payload["configs"]["reply.wechatmp.technical_ack"].startswith("已收到，正在运行")


def test_web_investment_config_saves_reply_text_values(investment_env, monkeypatch):
    from business.investment.config_service import get_config
    from channel.web.web_channel import InvestmentConfigHandler

    body = {"configs": {"reply.wechatmp.immediate_ack": "已收到，请稍候。"}}
    payload = _call_investment_json_handler(monkeypatch, InvestmentConfigHandler().POST, body=body)

    assert payload["status"] == "success"
    assert get_config("reply.wechatmp.immediate_ack") == "已收到，请稍候。"


def test_web_investment_config_excludes_and_rejects_global_model_and_wechatmp_keys(investment_env, monkeypatch):
    from sqlalchemy import select

    from business.investment import db
    from business.investment.schema import investment_configs
    from channel.web.web_channel import InvestmentConfigHandler

    body = {
        "configs": {
            "model.name": "bad-investment-model",
            "model.api_key": "sk-bad-investment",
            "wechatmp.app_id": "wx-bad",
            "tushare.token": "ts-web-secret-1234567890",
        },
        "operator_role": "admin",
    }

    post_payload = _call_investment_json_handler(monkeypatch, InvestmentConfigHandler().POST, body=body)

    assert post_payload["status"] == "error"
    assert "model.name" in post_payload["message"]
    assert "wechatmp.app_id" in post_payload["message"]
    with db.connect() as conn:
        keys = {
            row[0]
            for row in conn.execute(select(investment_configs.c.config_key)).fetchall()
        }
    assert keys.isdisjoint({"model.name", "model.api_key", "wechatmp.app_id"})
    assert "tushare.token" not in keys

    get_payload = _call_investment_json_handler(monkeypatch, InvestmentConfigHandler().GET)
    assert get_payload["status"] == "success"
    assert "tushare.token" in get_payload["configs"]
    assert not any(key.startswith(("model.", "wechatmp.")) for key in get_payload["configs"])


def test_web_stock_query_returns_matches_and_stats_without_refresh(investment_env, monkeypatch):
    from business.investment import stock_resolver
    from channel.web.web_channel import InvestmentStocksHandler

    stock_resolver.refresh_stock_symbols(
        [
            {"code": "300502.SZ", "name": "新易盛", "market": "SZ", "source": "seed"},
            {"code": "600519.SH", "name": "贵州茅台", "market": "SH", "source": "seed"},
        ],
        source="seed",
    )
    monkeypatch.setattr(stock_resolver, "refresh_from_auto", lambda: pytest.fail("query must not refresh"))
    monkeypatch.setattr(stock_resolver, "refresh_from_akshare", lambda: pytest.fail("query must not refresh"))
    monkeypatch.setattr(stock_resolver, "refresh_from_tushare", lambda: pytest.fail("query must not refresh"))

    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentStocksHandler().GET,
        params={"name": "新易盛", "limit": "20"},
    )

    assert payload["status"] == "success"
    assert [item["code"] for item in payload["stocks"]] == ["300502.SZ"]
    assert payload["stats"]["total"] == 2
    assert payload["stats"]["source_count"] == 1


def test_web_stock_refresh_dispatches_sources_and_reports_failures(investment_env, monkeypatch):
    from business.investment import stock_resolver
    from channel.web.web_channel import InvestmentStocksRefreshHandler

    monkeypatch.setattr(stock_resolver, "refresh_from_akshare", lambda: 3)
    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentStocksRefreshHandler().POST,
        body={"source": "akshare"},
    )
    assert payload["status"] == "success"
    assert payload["result"] == {"akshare": {"count": 3}}
    assert "stats" in payload

    monkeypatch.setattr(stock_resolver, "refresh_from_tushare", lambda: (_ for _ in ()).throw(RuntimeError("ts failed")))
    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentStocksRefreshHandler().POST,
        body={"source": "tushare"},
    )
    assert payload["status"] == "error"
    assert payload["message"] == "ts failed"

    monkeypatch.setattr(stock_resolver, "refresh_from_auto", lambda: {"akshare": {"error": "ak failed"}, "tushare": {"count": 4}})
    payload = _call_investment_json_handler(monkeypatch, InvestmentStocksRefreshHandler().POST, body={})
    assert payload["status"] == "success"
    assert payload["result"] == {"akshare": {"error": "ak failed"}, "tushare": {"count": 4}}


def test_business_record_cleanup_dry_run_and_execute_remove_useless_records(investment_env):
    from business.investment.db import connect
    from business.investment.record_cleanup import cleanup_useless_business_records
    from business.investment.schema import (
        investment_admin_sessions,
        investment_configs,
        investment_request_records,
        investment_stock_symbols,
    )

    with connect() as conn:
        conn.execute(
            investment_configs.insert(),
            {
                "config_key": "runtime.pg.test.cleanup",
                "config_value": "temp",
                "updated_at": "2026-01-01T00:00:00+00:00",
                "updated_by": "pytest",
            },
        )
        conn.execute(
            investment_stock_symbols.insert(),
            {
                "code": "000000.SZ",
                "name": "测试股票",
                "market": "SZ",
                "source": "runtime-test",
                "updated_at": "2026-01-01T00:00:00+00:00",
            },
        )
        conn.execute(
            investment_admin_sessions.insert(),
            {
                "session_id": "expired-session",
                "user_id": 1,
                "token_hash": "expired-token-hash",
                "created_at": "2026-01-01T00:00:00+00:00",
                "expires_at": "2026-01-02T00:00:00+00:00",
            },
        )
        conn.execute(
            investment_request_records.insert(),
            {
                "request_id": "old-unmatched-cleanup",
                "openid": "cleanup-openid",
                "raw_input": "nonsense",
                "service_type": "unmatched",
                "status": "failed",
                "created_at": "2026-01-01T00:00:00+00:00",
                "updated_at": "2026-01-01T00:00:00+00:00",
            },
        )

    dry_run = cleanup_useless_business_records(now="2026-06-02T00:00:00+00:00", dry_run=True)
    assert dry_run["runtime_test_configs"] == 1
    assert dry_run["runtime_test_stocks"] == 1
    assert dry_run["expired_admin_sessions"] == 1
    assert dry_run["old_exception_requests"] == 1

    with connect() as conn:
        assert conn.execute(text("select count(*) from investment_configs where config_key = 'runtime.pg.test.cleanup'")).scalar_one() == 1

    executed = cleanup_useless_business_records(now="2026-06-02T00:00:00+00:00", dry_run=False)
    assert executed == dry_run
    with connect() as conn:
        assert conn.execute(text("select count(*) from investment_configs where config_key = 'runtime.pg.test.cleanup'")).scalar_one() == 0
        assert conn.execute(text("select count(*) from investment_stock_symbols where source = 'runtime-test'")).scalar_one() == 0
        assert conn.execute(text("select count(*) from investment_admin_sessions where session_id = 'expired-session'")).scalar_one() == 0
        assert conn.execute(text("select count(*) from investment_request_records where request_id = 'old-unmatched-cleanup'")).scalar_one() == 0


def test_web_daily_content_generate_marks_generating_before_background_task(investment_env, monkeypatch):
    from business.investment.constants import ServiceType, Status
    from business.investment.daily_content import create_rate_content_draft
    from business.investment.records import get_content_record
    from channel.web.web_channel import InvestmentDailyContentGenerateHandler

    content_id = create_rate_content_draft(source_text="rate source")
    calls = []

    def fake_generate_content(task_content_id):
        calls.append(task_content_id)
        return SimpleNamespace(
            success=True,
            content_id=task_content_id,
            generated_text="",
            output_image="",
            output_files=[],
            error_code=None,
            user_prompt="",
            detail="",
        )

    monkeypatch.setattr("business.investment.daily_content.generate_content", fake_generate_content)

    payload = _call_investment_json_handler(
        monkeypatch,
        lambda: InvestmentDailyContentGenerateHandler().POST(content_id),
    )

    assert payload["status"] == "success"
    assert payload["content_id"] == content_id
    assert payload["generation_status"] == "started"
    assert get_content_record(content_id).service_type == ServiceType.RATE
    assert get_content_record(content_id).status == Status.GENERATING


def test_web_daily_content_create_binds_session_admin_not_body_operator(investment_env, monkeypatch):
    from business.investment.audit_service import list_operation_audits
    from business.investment.auth_service import authenticate_admin, create_admin_session, create_admin_user
    from business.investment.db import connect
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentDailyContentHandler

    create_admin_user("content-session", "content-pass", role="content_operator")
    admin = authenticate_admin("content-session", "content-pass")
    assert admin is not None
    token = create_admin_session(admin)
    monkeypatch.setattr(web_channel.web, "cookies", lambda: {"cow_investment_session": token})
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)
    monkeypatch.setattr(
        web_channel.web,
        "data",
        lambda: json.dumps({
            "service_type": "rate",
            "source_text": "content source",
            "operator": "spoofed-content-operator",
        }).encode("utf-8"),
    )

    payload = json.loads(InvestmentDailyContentHandler().POST())

    assert payload["status"] == "success"
    content_id = payload["content_id"]
    with connect() as conn:
        row = conn.execute(
            text(
                "select operator, created_by_admin_id, created_by_username, created_by_role, "
                "updated_by_admin_id, updated_by_username, updated_by_role "
                "from investment_daily_contents where content_id = :content_id"
            ),
            {"content_id": content_id},
        ).fetchone()
    assert row is not None
    assert row.operator == "content-session"
    assert row.created_by_admin_id == admin.id
    assert row.created_by_username == "content-session"
    assert row.created_by_role == "content_operator"
    assert row.updated_by_admin_id == admin.id
    assert row.updated_by_username == "content-session"
    assert row.updated_by_role == "content_operator"
    audit = list_operation_audits(limit=1, target_type="daily_content", target_id=content_id)[0]
    assert audit.action == "content.create"
    assert audit.operator == "content-session"
    assert audit.operator_admin_id == admin.id
    assert audit.operation_category == "content"


def test_web_daily_content_get_returns_current_effective_content(investment_env, tmp_path, monkeypatch):
    from pathlib import Path

    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.storage import get_storage_dirs
    from channel.web.web_channel import InvestmentDailyContentHandler

    image = tmp_path / "rate-current.png"
    image.write_bytes(b"png")
    older_image = tmp_path / "rate-older.png"
    older_image.write_bytes(b"png")
    older_id = create_content_draft(ServiceType.RATE, source_text="older", operator="operator-old")
    current_id = create_content_draft(ServiceType.RATE, source_text="current", operator="operator-current")
    set_content_effective(older_id, str(older_image), operator="operator-old")
    set_content_effective(current_id, str(image), operator="operator-current")

    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentDailyContentHandler().GET,
        params={"service_type": "rate", "limit": "50"},
    )

    assert payload["status"] == "success"
    assert payload["current_effective"]["content_id"] == current_id
    assert payload["current_effective"]["service_type"] == ServiceType.RATE
    assert payload["current_effective"]["status"] == "effective"
    current_output = Path(payload["current_effective"]["output_image"])
    assert current_output != image
    assert current_output.is_file()
    assert current_output.read_bytes() == b"png"
    assert current_output.resolve().is_relative_to((get_storage_dirs()["generated"] / "archive").resolve())
    assert payload["current_effective"]["operator"] == "operator-current"


def test_web_daily_content_get_returns_history_artifacts(investment_env, tmp_path, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, update_generation_success
    from channel.web.web_channel import InvestmentDailyContentHandler

    image = tmp_path / "rate-history.png"
    image.write_bytes(b"history-png")
    content_id = create_content_draft(
        ServiceType.RATE,
        source_text="history source text",
        operator="operator-history",
        effective_date="2026-06-03",
    )
    update_generation_success(content_id, "history generated text", str(image))

    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentDailyContentHandler().GET,
        params={"service_type": "rate", "limit": "50"},
    )

    history = next(content for content in payload["contents"] if content["content_id"] == content_id)
    assert history["source_text"] == "history source text"
    assert history["generated_text"] == "history generated text"
    assert history["output_artifacts"]
    assert history["output_artifacts"][0]["artifact_role"] == "output_image"
    assert history["output_artifacts"][0]["file_path"] == history["output_image"]


def test_set_content_effective_archives_external_output_image(investment_env, tmp_path):
    from pathlib import Path

    from business.investment.constants import ServiceType, Status
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.records import get_content_record, list_output_files
    from business.investment.storage import get_storage_dirs

    image = tmp_path / "rate_card.png"
    image.write_bytes(b"legacy-rate-card")
    content_id = create_content_draft(
        ServiceType.RATE,
        source_text="legacy image",
        operator="operator-effective",
        effective_date="2026-06-04",
    )

    set_content_effective(content_id, str(image), operator="operator-effective")

    record = get_content_record(content_id)
    artifacts = list_output_files(content_id)
    assert record.status == Status.EFFECTIVE
    assert record.output_image != str(image)
    assert Path(record.output_image).is_file()
    assert Path(record.output_image).read_bytes() == b"legacy-rate-card"
    assert Path(record.output_image).resolve().is_relative_to((get_storage_dirs()["generated"] / "archive").resolve())
    assert Path(record.output_image).name.startswith("output_image_rate_card_")
    assert artifacts
    assert artifacts[0]["artifact_role"] == "output_image"
    assert artifacts[0]["file_path"] == record.output_image


def test_list_content_records_filters_by_effective_date(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft
    from business.investment.records import list_content_records

    create_content_draft(ServiceType.RATE, source_text="old rate", effective_date="2026-05-28")
    expected_id = create_content_draft(ServiceType.RATE, source_text="today rate", effective_date="2026-05-29")
    create_content_draft(ServiceType.CONVERTIBLE_BOND, source_text="today cb", effective_date="2026-05-29")

    records = list_content_records(
        service_type=ServiceType.RATE,
        effective_date="2026-05-29",
    )

    assert [record.content_id for record in records] == [expected_id]
    assert {record.effective_date for record in records} == {"2026-05-29"}


def test_web_content_handlers_accept_effective_date_filter(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft
    from channel.web.web_channel import InvestmentContentRecordsHandler, InvestmentDailyContentHandler

    old_id = create_content_draft(ServiceType.RATE, source_text="old rate", effective_date="2026-05-28")
    expected_id = create_content_draft(ServiceType.RATE, source_text="today rate", effective_date="2026-05-29")

    daily_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentDailyContentHandler().GET,
        params={"service_type": "rate", "limit": "20", "effective_date": "2026-05-29"},
    )
    records_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentContentRecordsHandler().GET,
        params={"service_type": "rate", "limit": "20", "effective_date": "2026-05-29"},
    )

    assert [content["content_id"] for content in daily_payload["contents"]] == [expected_id]
    assert [record["content_id"] for record in records_payload["records"]] == [expected_id]
    assert old_id not in {content["content_id"] for content in daily_payload["contents"]}


def test_web_content_handlers_sanitize_limit_and_reject_unmatched_service_type(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft
    from channel.web.web_channel import InvestmentContentRecordsHandler, InvestmentDailyContentHandler

    rate_id = create_content_draft(ServiceType.RATE, source_text="rate", effective_date="2026-05-29")
    cb_id = create_content_draft(ServiceType.CONVERTIBLE_BOND, source_text="cb", effective_date="2026-05-29")

    invalid_limit_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentDailyContentHandler().GET,
        params={"service_type": "rate", "limit": "bad"},
    )
    negative_limit_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentContentRecordsHandler().GET,
        params={"service_type": "rate", "limit": "-10"},
    )
    oversized_limit_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentDailyContentHandler().GET,
        params={"service_type": "rate", "limit": "9999"},
    )
    unmatched_daily_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentDailyContentHandler().GET,
        params={"service_type": "unknown-service", "limit": "20"},
    )
    unmatched_records_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentContentRecordsHandler().GET,
        params={"service_type": "unknown-service", "limit": "20"},
    )

    assert invalid_limit_payload["status"] == "success"
    assert [content["content_id"] for content in invalid_limit_payload["contents"]] == [rate_id]
    assert negative_limit_payload["status"] == "success"
    assert len(negative_limit_payload["records"]) == 1
    assert oversized_limit_payload["status"] == "success"
    assert [content["content_id"] for content in oversized_limit_payload["contents"]] == [rate_id]
    assert unmatched_daily_payload["contents"] == []
    assert unmatched_daily_payload["current_effective"] is None
    assert unmatched_records_payload["records"] == []
    assert cb_id not in {content["content_id"] for content in oversized_limit_payload["contents"]}


def test_web_records_handlers_sanitize_invalid_limit(investment_env, monkeypatch):
    from business.investment.audit_service import record_operation_audit
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record
    from channel.web.web_channel import InvestmentOperationAuditsHandler, InvestmentRequestRecordsHandler

    request_id = create_request_record("openid", "利率", ServiceType.RATE)
    audit_id = record_operation_audit("created", "request", request_id, operator="tester")

    records_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentRequestRecordsHandler().GET,
        params={"limit": "bad"},
    )
    audits_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentOperationAuditsHandler().GET,
        params={"limit": "bad"},
    )

    assert records_payload["status"] == "success"
    assert [record["request_id"] for record in records_payload["records"]] == [request_id]
    assert audits_payload["status"] == "success"
    assert [audit["audit_id"] for audit in audits_payload["audits"]] == [audit_id]


def test_request_records_page_returns_total_offset_and_api_pagination(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.records import create_request_record, list_request_records_page
    from channel.web.web_channel import InvestmentRequestRecordsHandler

    request_ids = [
        create_request_record(f"openid-page-{index}", f"分页请求 {index}", ServiceType.RATE)
        for index in range(5)
    ]
    with connect() as conn:
        for index, request_id in enumerate(request_ids):
            conn.execute(
                text(
                    "update investment_request_records "
                    "set created_at = :created_at, updated_at = :created_at "
                    "where request_id = :request_id"
                ),
                {"created_at": f"2026-05-29T00:0{index}:00+00:00", "request_id": request_id},
            )

    records, total = list_request_records_page(page=2, page_size=2, service_type=ServiceType.RATE, keyword="分页请求")
    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentRequestRecordsHandler().GET,
        params={"page": "2", "page_size": "2", "service_type": "rate", "keyword": "分页请求"},
    )
    capped_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentRequestRecordsHandler().GET,
        params={"page": "1", "page_size": "9999", "service_type": "rate", "keyword": "分页请求"},
    )

    assert total == 5
    assert [record.request_id for record in records] == [request_ids[2], request_ids[1]]
    assert [record["request_id"] for record in payload["records"]] == [request_ids[2], request_ids[1]]
    assert payload["pagination"] == {"page": 2, "page_size": 2, "total": 5, "total_pages": 3}
    assert capped_payload["pagination"]["page_size"] == 200
    assert capped_payload["pagination"]["total"] == 5


def test_request_records_page_filters_by_openid_or_mobile(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, list_request_records_page, succeed_request_record
    from business.investment.user_service import create_user

    create_user("openid-a", name="Alice", institution="Inst A", mobile="13800000000", enabled=True, allowed_services=[ServiceType.ALL])
    create_user("openid-b", name="Bob", institution="Inst B", mobile="13900000000", enabled=True, allowed_services=[ServiceType.ALL])
    first = create_request_record("openid-a", "新易盛 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    second = create_request_record("openid-b", "利率", ServiceType.RATE)
    succeed_request_record(first, output_files=["/tmp/a.png"], elapsed_ms=1)
    succeed_request_record(second, output_files=["/tmp/b.png"], elapsed_ms=1)

    by_openid, total_openid = list_request_records_page(page=1, page_size=20, customer="openid-a")
    by_mobile, total_mobile = list_request_records_page(page=1, page_size=20, customer="13900000000")
    by_name, total_name = list_request_records_page(page=1, page_size=20, customer="Alice")
    by_institution, total_institution = list_request_records_page(page=1, page_size=20, customer="Inst B")

    assert total_openid == 1
    assert [record.openid for record in by_openid] == ["openid-a"]
    assert total_mobile == 1
    assert [record.openid for record in by_mobile] == ["openid-b"]
    assert total_name == 1
    assert [record.openid for record in by_name] == ["openid-a"]
    assert total_institution == 1
    assert [record.openid for record in by_institution] == ["openid-b"]


def test_request_records_page_unknown_service_returns_empty(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, list_request_records_page, succeed_request_record

    request_id = create_request_record("openid", "利率", ServiceType.RATE)
    succeed_request_record(request_id, output_files=["/tmp/rate.png"], elapsed_ms=1)

    records, total = list_request_records_page(page=1, page_size=20, service_type="unknown-service")

    assert records == []
    assert total == 0


def test_content_records_page_returns_total_and_filter_pagination(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft
    from business.investment.db import connect
    from business.investment.records import list_content_records_page
    from channel.web.web_channel import InvestmentContentRecordsHandler

    content_ids = [
        create_content_draft(ServiceType.RATE, source_text=f"分页内容 {index}", effective_date="2026-05-29")
        for index in range(5)
    ]
    create_content_draft(ServiceType.CONVERTIBLE_BOND, source_text="excluded", effective_date="2026-05-29")
    with connect() as conn:
        for index, content_id in enumerate(content_ids):
            conn.execute(
                text(
                    "update investment_daily_contents "
                    "set created_at = :created_at, updated_at = :created_at "
                    "where content_id = :content_id"
                ),
                {"created_at": f"2026-05-29T01:0{index}:00+00:00", "content_id": content_id},
            )

    records, total = list_content_records_page(
        page=2,
        page_size=2,
        service_type=ServiceType.RATE,
        effective_date="2026-05-29",
    )
    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentContentRecordsHandler().GET,
        params={"page": "2", "page_size": "2", "service_type": "rate", "effective_date": "2026-05-29"},
    )

    assert total == 5
    assert [record.content_id for record in records] == [content_ids[2], content_ids[1]]
    assert [record["content_id"] for record in payload["records"]] == [content_ids[2], content_ids[1]]
    assert payload["pagination"] == {"page": 2, "page_size": 2, "total": 5, "total_pages": 3}


def test_operation_audits_page_returns_total_and_filter_pagination(investment_env, monkeypatch):
    from business.investment.audit_service import list_operation_audits_page, record_operation_audit
    from business.investment.db import connect
    from channel.web.web_channel import InvestmentOperationAuditsHandler

    audit_ids = [
        record_operation_audit("cache.invalidate", "investment_cache_entry", f"cache-{index}", operator="page-ops")
        for index in range(5)
    ]
    record_operation_audit("config.update", "investment_config", "excluded", operator="page-ops")
    with connect() as conn:
        for index, audit_id in enumerate(audit_ids):
            conn.execute(
                text("update investment_operation_audits set created_at = :created_at where audit_id = :audit_id"),
                {"created_at": f"2026-05-29T02:0{index}:00+00:00", "audit_id": audit_id},
            )

    audits, total = list_operation_audits_page(
        page=2,
        page_size=2,
        action="cache.invalidate",
        operator="page-ops",
        keyword="cache-",
    )
    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentOperationAuditsHandler().GET,
        params={"page": "2", "page_size": "2", "action": "cache.invalidate", "operator": "page-ops", "keyword": "cache-"},
    )

    assert total == 5
    assert [audit.audit_id for audit in audits] == [audit_ids[2], audit_ids[1]]
    assert [audit["audit_id"] for audit in payload["audits"]] == [audit_ids[2], audit_ids[1]]
    assert payload["pagination"] == {"page": 2, "page_size": 2, "total": 5, "total_pages": 3}


def test_operation_audit_records_admin_actor_fields(investment_env):
    from business.investment.audit_service import AdminActor, list_operation_audits, record_operation_audit

    audit_id = record_operation_audit(
        "customer.update",
        "customer",
        "openid-actor",
        actor=AdminActor(admin_id=42, username="actor-admin", role="admin"),
        detail={"changed": ["mobile"]},
        result_status="failed",
        error_code="VALIDATION_ERROR",
        error_message="mobile invalid",
        elapsed_ms=123,
        before_state={"mobile": "13800138000"},
        after_state={"mobile": "bad-mobile"},
    )

    audit = list_operation_audits(limit=1, target_type="customer", target_id="openid-actor")[0]
    assert audit.audit_id == audit_id
    assert audit.operator == "actor-admin"
    assert audit.operator_admin_id == 42
    assert audit.operator_username == "actor-admin"
    assert audit.operator_role == "admin"
    assert audit.operation_category == "customer"
    assert audit.result_status == "failed"
    assert audit.error_code == "VALIDATION_ERROR"
    assert audit.error_message == "mobile invalid"
    assert audit.elapsed_ms == 123
    assert audit.before_state == {"mobile": "13800138000"}
    assert audit.after_state == {"mobile": "bad-mobile"}


def test_cache_entries_page_returns_total_and_filter_pagination(investment_env, monkeypatch):
    from business.investment.cache_service import build_cache_key, list_cache_entries_page, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from channel.web.web_channel import InvestmentCacheHandler

    cache_keys = []
    for index in range(5):
        cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, f"30050{index}.SZ", "2026-05-29", "v1")
        cache_keys.append(cache_key)
        write_cache_entry(
            cache_key=cache_key,
            service_type=ServiceType.TECHNICAL_ANALYSIS,
            normalized_target=f"30050{index}.SZ",
            market_date="2026-05-29",
            version_fingerprint="v1",
            output_files=[f"/tmp/card-{index}.png"],
        )
    write_cache_entry(
        cache_key=build_cache_key(ServiceType.RATE, "RATE", "2026-05-29", "v1"),
        service_type=ServiceType.RATE,
        normalized_target="RATE",
        market_date="2026-05-29",
        version_fingerprint="v1",
        output_files=["/tmp/rate.png"],
    )
    with connect() as conn:
        for index, cache_key in enumerate(cache_keys):
            conn.execute(
                text("update investment_cache_entries set updated_at = :updated_at where cache_key = :cache_key"),
                {"updated_at": f"2026-05-29T03:0{index}:00+00:00", "cache_key": cache_key},
            )

    entries, total = list_cache_entries_page(
        page=2,
        page_size=2,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        market_date="2026-05-29",
    )
    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentCacheHandler().GET,
        params={"page": "2", "page_size": "2", "service_type": "technical_analysis", "market_date": "2026-05-29"},
    )

    assert total == 5
    assert [entry.cache_key for entry in entries] == [cache_keys[2], cache_keys[1]]
    assert [entry["cache_key"] for entry in payload["entries"]] == [cache_keys[2], cache_keys[1]]
    assert payload["pagination"] == {"page": 2, "page_size": 2, "total": 5, "total_pages": 3}
    assert payload["market_dates"] == ["2026-05-29"]


def test_cache_handler_without_market_date_returns_history_across_dates(investment_env, monkeypatch):
    from business.investment.cache_service import build_cache_key, write_cache_entry
    from business.investment.constants import ServiceType
    from channel.web.web_channel import InvestmentCacheHandler

    write_cache_entry(
        cache_key=build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300500.SZ", "2026-05-28", "v1"),
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300500.SZ",
        market_date="2026-05-28",
        version_fingerprint="v1",
        output_files=["/tmp/card-old.png"],
    )
    write_cache_entry(
        cache_key=build_cache_key(ServiceType.RATE, "RATE", "2026-06-03", "v1"),
        service_type=ServiceType.RATE,
        normalized_target="RATE",
        market_date="2026-06-03",
        version_fingerprint="v1",
        output_files=["/tmp/rate-new.png"],
    )

    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentCacheHandler().GET,
        params={"page": "1", "page_size": "20"},
    )

    assert {entry["market_date"] for entry in payload["entries"]} == {"2026-05-28", "2026-06-03"}
    assert payload["market_dates"] == ["2026-06-03", "2026-05-28"]
    assert payload["pagination"]["total"] == 2


def test_request_records_api_filters_and_prefers_mobile_customer_display(investment_env, monkeypatch):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.db import connect
    from business.investment.records import create_request_record, fail_request_record, succeed_request_record
    from business.investment.user_service import create_user
    from channel.web.web_channel import InvestmentRequestRecordsHandler

    create_user(
        "openid-mobile",
        name="Mobile Customer",
        institution="Inst Mobile",
        mobile="13900000000",
        allowed_services=[ServiceType.ALL],
    )
    create_user("openid-plain", name="Plain Customer", allowed_services=[ServiceType.ALL])
    included = create_request_record("openid-mobile", "新易盛 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    succeed_request_record(included, output_files=["/tmp/card.png"], elapsed_ms=18, stock_name="新易盛")
    failed = create_request_record("openid-plain", "利率", ServiceType.RATE)
    fail_request_record(failed, ErrorCode.NO_CONTENT, detail="no active content", elapsed_ms=3)

    with connect() as conn:
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :created_at "
                "where request_id = :request_id"
            ),
            {"created_at": "2026-05-29T01:00:00+00:00", "request_id": included},
        )
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :created_at "
                "where request_id = :request_id"
            ),
            {"created_at": "2026-05-30T01:00:00+00:00", "request_id": failed},
        )

    mobile_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentRequestRecordsHandler().GET,
        params={
            "service_type": "technical_analysis",
            "status": "success",
            "keyword": "13900000000",
            "start_date": "2026-05-29",
            "end_date": "2026-05-29",
            "limit": "20",
        },
    )
    openid_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentRequestRecordsHandler().GET,
        params={"keyword": "openid-plain", "limit": "20"},
    )

    assert mobile_payload["status"] == "success"
    assert [record["request_id"] for record in mobile_payload["records"]] == [included]
    assert mobile_payload["records"][0]["customer_mobile"] == "13900000000"
    assert mobile_payload["records"][0]["customer_display"] == "13900000000"
    assert mobile_payload["records"][0]["customer_name"] == "Mobile Customer"
    assert mobile_payload["records"][0]["institution"] == "Inst Mobile"
    assert [record["request_id"] for record in openid_payload["records"]] == [failed]
    assert openid_payload["records"][0]["customer_mobile"] == ""
    assert openid_payload["records"][0]["customer_display"] == "openid-plain"


def test_web_record_endpoints_filter_main_fields_with_realistic_web_input(investment_env, monkeypatch):
    from business.investment.audit_service import record_operation_audit
    from business.investment.cache_service import build_cache_key, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft
    from business.investment.db import connect
    from business.investment.records import create_request_record
    from channel.web import web_channel
    from channel.web.web_channel import (
        InvestmentCacheHandler,
        InvestmentContentRecordsHandler,
        InvestmentOperationAuditsHandler,
        InvestmentRequestRecordsHandler,
    )

    rate_request_id = create_request_record("openid-rate", "利率", ServiceType.RATE)
    ta_request_id = create_request_record("openid-ta", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    rate_content_id = create_content_draft(ServiceType.RATE, source_text="rate", effective_date="2026-05-31")
    cb_content_id = create_content_draft(ServiceType.CONVERTIBLE_BOND, source_text="cb", effective_date="2026-05-31")
    audit_id = record_operation_audit(
        "cache.invalidate",
        "investment_cache_entry",
        target_id="rate-cache",
        operator="rate-operator",
        detail={"reason": "rate filter check"},
    )
    excluded_audit_id = record_operation_audit(
        "config.update",
        "investment_config",
        target_id="router",
        operator="other-operator",
        detail={"reason": "other"},
    )
    rate_cache_key = build_cache_key(ServiceType.RATE, "RATE", "2026-05-31", "v1")
    ta_cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-31", "v1")
    write_cache_entry(
        cache_key=rate_cache_key,
        service_type=ServiceType.RATE,
        normalized_target="RATE",
        market_date="2026-05-31",
        version_fingerprint="v1",
        output_files=["/tmp/rate.png"],
    )
    write_cache_entry(
        cache_key=ta_cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-31",
        version_fingerprint="v1",
        output_files=["/tmp/ta.png"],
    )
    with connect() as conn:
        for request_id in (rate_request_id, ta_request_id):
            conn.execute(
                text(
                    "update investment_request_records "
                    "set created_at = :created_at, updated_at = :created_at "
                    "where request_id = :request_id"
                ),
                {"created_at": "2026-05-31T01:00:00+00:00", "request_id": request_id},
            )
        for current_audit_id in (audit_id, excluded_audit_id):
            conn.execute(
                text("update investment_operation_audits set created_at = :created_at where audit_id = :audit_id"),
                {"created_at": "2026-05-31T02:00:00+00:00", "audit_id": current_audit_id},
            )

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)

    def call(handler, params):
        monkeypatch.setattr(
            web_channel.web,
            "input",
            lambda **defaults: SimpleNamespace(**({**defaults, **params})),
        )
        return json.loads(handler())

    requests_payload = call(
        InvestmentRequestRecordsHandler().GET,
        {
            "service_type": "rate",
            "keyword": "利率",
            "start_date": "2026-05-31",
            "end_date": "2026-05-31",
            "page": "1",
            "page_size": "80",
        },
    )
    contents_payload = call(
        InvestmentContentRecordsHandler().GET,
        {"service_type": "rate", "effective_date": "2026-05-31", "page": "1", "page_size": "80"},
    )
    cache_payload = call(
        InvestmentCacheHandler().GET,
        {"service_type": "rate", "market_date": "2026-05-31", "page": "1", "page_size": "80"},
    )
    audits_payload = call(
        InvestmentOperationAuditsHandler().GET,
        {
            "action": "cache.invalidate",
            "target_type": "investment_cache_entry",
            "target_id": "rate-cache",
            "operator": "rate-operator",
            "keyword": "rate filter",
            "start_date": "2026-05-31",
            "end_date": "2026-05-31",
            "page": "1",
            "page_size": "80",
        },
    )

    assert [record["request_id"] for record in requests_payload["records"]] == [rate_request_id]
    assert ta_request_id not in {record["request_id"] for record in requests_payload["records"]}
    assert requests_payload["pagination"]["total"] == 1
    assert [record["content_id"] for record in contents_payload["records"]] == [rate_content_id]
    assert cb_content_id not in {record["content_id"] for record in contents_payload["records"]}
    assert contents_payload["pagination"]["total"] == 1
    assert [entry["cache_key"] for entry in cache_payload["entries"]] == [rate_cache_key]
    assert ta_cache_key not in {entry["cache_key"] for entry in cache_payload["entries"]}
    assert cache_payload["pagination"]["total"] == 1
    assert [audit["audit_id"] for audit in audits_payload["audits"]] == [audit_id]
    assert excluded_audit_id not in {audit["audit_id"] for audit in audits_payload["audits"]}
    assert audits_payload["pagination"]["total"] == 1


def test_operation_audits_api_filters_by_operator_action_keyword_and_date(investment_env, monkeypatch):
    from business.investment.audit_service import record_operation_audit
    from business.investment.db import connect
    from channel.web.web_channel import InvestmentOperationAuditsHandler

    included = record_operation_audit(
        "cache.invalidate",
        "investment_cache_entry",
        "cache-a",
        operator="ops-a",
        detail={"reason": "refresh stale cache"},
    )
    excluded = record_operation_audit(
        "config.update",
        "investment_config",
        "router.enable_web_open_chat",
        operator="ops-b",
        detail={"reason": "config change"},
    )

    with connect() as conn:
        conn.execute(
            text(
                "update investment_operation_audits "
                "set created_at = :created_at "
                "where audit_id = :audit_id"
            ),
            {"created_at": "2026-05-29T02:00:00+00:00", "audit_id": included},
        )
        conn.execute(
            text(
                "update investment_operation_audits "
                "set created_at = :created_at "
                "where audit_id = :audit_id"
            ),
            {"created_at": "2026-05-30T02:00:00+00:00", "audit_id": excluded},
        )

    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentOperationAuditsHandler().GET,
        params={
            "action": "cache.invalidate",
            "operator": "ops-a",
            "target_type": "investment_cache_entry",
            "keyword": "stale",
            "start_date": "2026-05-29",
            "end_date": "2026-05-29",
            "limit": "20",
        },
    )

    assert payload["status"] == "success"
    assert [audit["audit_id"] for audit in payload["audits"]] == [included]
    assert excluded not in {audit["audit_id"] for audit in payload["audits"]}


def test_investment_web_api_end_to_end_smoke_without_external_services(investment_env, tmp_path, monkeypatch):
    from business.investment.auth_service import authenticate_admin, create_admin_session, create_admin_user
    from business.investment.daily_content import update_generation_success
    from business.investment.records import get_content_record
    from business.investment.router import handle_text_message
    from channel.web import web_channel
    from channel.web.web_channel import (
        InvestmentDailyContentEffectiveHandler,
        InvestmentDailyContentGenerateHandler,
        InvestmentDailyContentHandler,
        InvestmentHealthHandler,
        InvestmentRequestRecordsExportHandler,
        InvestmentRequestRecordsHandler,
        InvestmentUsersHandler,
    )

    create_admin_user("admin-e2e", "admin-pass", role="admin")
    create_admin_user("readonly-e2e", "readonly-pass", role="readonly")
    admin_token = create_admin_session(authenticate_admin("admin-e2e", "admin-pass"))
    readonly_token = create_admin_session(authenticate_admin("readonly-e2e", "readonly-pass"))
    headers = []
    current_params = {}
    current_body = {}
    current_token = admin_token

    monkeypatch.setattr(web_channel.web, "header", lambda name, value: headers.append((name, value)))
    monkeypatch.setattr(web_channel.web, "cookies", lambda: {"cow_investment_session": current_token})
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(**current_params))
    monkeypatch.setattr(web_channel.web, "data", lambda: json.dumps(current_body, ensure_ascii=False).encode("utf-8"))
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)
    monkeypatch.setattr(web_channel.web.ctx, "env", {"CONTENT_TYPE": "application/json"}, raising=False)

    def call_json(handler, *, params=None, body=None, token=None):
        nonlocal current_params, current_body, current_token
        current_params = params or {}
        current_body = body or {}
        current_token = token or admin_token
        return json.loads(handler())

    customer = {
        "openid": "openid-e2e",
        "name": "E2E Customer",
        "institution": "E2E Inst",
        "enabled": True,
        "allowed_services": ["全部"],
        "auth_start_at": "2026-01-01T00:00:00",
        "auth_end_at": "2099-12-31T23:59:59",
    }
    with pytest.raises(web_channel.web.HTTPError) as denied_error:
        call_json(InvestmentUsersHandler().POST, body=customer, token=readonly_token)
    denied = json.loads(denied_error.value.data)
    assert denied["code"] == "permission_denied"
    assert denied["permission"] == "customers.write"

    created = call_json(InvestmentUsersHandler().POST, body=customer)
    assert created["status"] == "success"
    assert created["action"] == "created"
    listed_users = call_json(InvestmentUsersHandler().GET, params={"enabled": "true", "openid": "openid-e2e"})
    assert [user["openid"] for user in listed_users["users"]] == ["openid-e2e"]

    generated_image = tmp_path / "generated-rate.png"
    generated_image.write_bytes(b"rate")
    draft = call_json(
        InvestmentDailyContentHandler().POST,
        body={"service_type": "利率", "source_text": "rate source", "operator": "admin-e2e"},
    )
    content_id = draft["content_id"]

    class ImmediateThread:
        def __init__(self, target, **_kwargs):
            self.target = target

        def start(self):
            self.target()

    def fake_generate_content(task_content_id):
        update_generation_success(task_content_id, "rate generated", str(generated_image))
        return SimpleNamespace(success=True, content_id=task_content_id, generated_text="rate generated", output_image=str(generated_image), output_files=[str(generated_image)])

    monkeypatch.setattr(web_channel.threading, "Thread", ImmediateThread)
    monkeypatch.setattr("business.investment.daily_content.generate_content", fake_generate_content)
    generated = call_json(lambda: InvestmentDailyContentGenerateHandler().POST(content_id))
    assert generated["status"] == "success"
    assert generated["generation_status"] == "started"
    assert get_content_record(content_id).output_image == str(generated_image)

    effective = call_json(
        lambda: InvestmentDailyContentEffectiveHandler().POST(content_id),
        body={"operator": "admin-e2e", "effective_date": "2026-05-28"},
    )
    assert effective["status"] == "success"
    current = call_json(InvestmentDailyContentHandler().GET, params={"service_type": "rate", "limit": "20"})
    assert current["current_effective"]["content_id"] == content_id

    cb_image = tmp_path / "cb-current.png"
    cb_image.write_bytes(b"cb")
    cb_draft = call_json(
        InvestmentDailyContentHandler().POST,
        body={
            "service_type": "转债",
            "source_text": "cb source",
            "generated_text": "cb generated",
            "output_image": str(cb_image),
            "operator": "admin-e2e",
            "effective_date": "2026-05-28",
        },
    )
    call_json(lambda: InvestmentDailyContentEffectiveHandler().POST(cb_draft["content_id"]), body={"operator": "admin-e2e"})

    rate_reply = handle_text_message("openid-e2e", "利率")
    cb_reply = handle_text_message("openid-e2e", "转债")
    assert rate_reply.success is True
    assert rate_reply.output_files == [str(generated_image)]
    assert cb_reply.success is True
    assert cb_reply.output_files == [str(cb_image)]

    records = call_json(InvestmentRequestRecordsHandler().GET, params={"limit": "20"})
    assert {record["raw_input"] for record in records["records"]} >= {"利率", "转债"}
    assert any(artifact["artifact_role"] == "output_image" for record in records["records"] for artifact in record["output_artifacts"])

    health_calls = []
    monkeypatch.setattr(
        "business.investment.health.run_health_checks",
        lambda run_smoke=False: health_calls.append(run_smoke)
        or [SimpleNamespace(name="database", ok=True, detail="ok", level="ok")],
    )
    health = call_json(InvestmentHealthHandler().GET, params={"smoke": "1"})
    assert health["status"] == "success"
    assert health["run_smoke"] is True
    assert health_calls == [True]

    headers.clear()
    export_date = datetime.now(UTC).date().isoformat()
    nonlocal_params = {"start_date": export_date, "end_date": export_date, "service_type": "", "year": "", "month": "", "quarter": ""}
    current_params = nonlocal_params
    exported = InvestmentRequestRecordsExportHandler().GET()
    rows = _xlsx_sheet_rows(exported)
    assert rows[0][:5] == ["请求时间", "OpenID", "客户姓名", "机构", "原始输入"]
    assert any(row[1:5] == ["openid-e2e", "E2E Customer", "E2E Inst", "利率"] for row in rows[1:])
    assert ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") in headers


def test_user_service_permission_edges_and_upsert(investment_env):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.user_service import (
        ImportUserRow,
        create_user,
        import_users,
        verify_permission,
    )

    assert verify_permission("missing", ServiceType.RATE).allowed is False
    assert verify_permission("missing", ServiceType.RATE).error_code == ErrorCode.UNAUTHORIZED

    create_user("disabled", enabled=False, allowed_services=[ServiceType.ALL])
    assert verify_permission("disabled", ServiceType.RATE).error_code == ErrorCode.USER_DISABLED

    create_user("limited", enabled=True, allowed_services=[ServiceType.RATE])
    assert verify_permission("limited", ServiceType.CONVERTIBLE_BOND).error_code == ErrorCode.UNAUTHORIZED

    create_user(
        "expired",
        enabled=True,
        allowed_services=[ServiceType.ALL],
        auth_end_at=datetime.now(UTC).replace(tzinfo=None) - timedelta(days=1),
    )
    assert verify_permission("expired", ServiceType.RATE).error_code == ErrorCode.AUTH_EXPIRED

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    assert verify_permission("ok", ServiceType.TECHNICAL_ANALYSIS).allowed is True

    result = import_users([
        ImportUserRow(openid="ok", name="updated", allowed_services="利率"),
        ImportUserRow(openid="new", name="new user", allowed_services="全部"),
    ])
    assert result.created == 1
    assert result.updated == 1
    assert verify_permission("ok", ServiceType.RATE).allowed is True
    assert verify_permission("ok", ServiceType.CONVERTIBLE_BOND).allowed is False


def test_user_service_crud_list_and_all_service_contract(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user, disable_user, get_user_by_openid, list_users, update_user, verify_permission

    user_id = create_user(
        "crud-openid",
        name="Alice",
        institution="Inst A",
        mobile="13800000000",
        enabled=True,
        allowed_services=["全部"],
        auth_start_at="2026-01-01T00:00:00",
        auth_end_at="2026-12-31T23:59:59",
        remark="first",
    )

    created = get_user_by_openid("crud-openid")
    assert created is not None
    assert created.id == user_id
    assert created.name == "Alice"
    assert [item for item in created.allowed_services or []] == [ServiceType.ALL]
    assert list_users(enabled=True, openid="crud") == [created]
    assert verify_permission("crud-openid", ServiceType.TECHNICAL_ANALYSIS).allowed is True
    assert verify_permission("crud-openid", ServiceType.RATE).allowed is True
    assert verify_permission("crud-openid", ServiceType.CONVERTIBLE_BOND).allowed is True

    update_user("crud-openid", name="Alice B", allowed_services=["利率"], remark="updated")
    updated = get_user_by_openid("crud-openid")
    assert updated is not None
    assert updated.name == "Alice B"
    assert updated.remark == "updated"
    assert verify_permission("crud-openid", ServiceType.RATE).allowed is True
    assert verify_permission("crud-openid", ServiceType.CONVERTIBLE_BOND).allowed is False

    disable_user("crud-openid")
    assert list_users(enabled=False, openid="crud")[0].enabled is False


def test_user_services_normalize_all_when_all_or_every_business_service_selected(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user, get_user_by_openid, update_user

    create_user(
        "all-plus-specific",
        allowed_services=["全部", "技术分析", "利率"],
    )
    all_plus_specific = get_user_by_openid("all-plus-specific")
    assert all_plus_specific is not None
    assert all_plus_specific.allowed_services == [ServiceType.ALL]

    create_user(
        "all-specific",
        allowed_services=["技术分析", "利率", "转债"],
    )
    all_specific = get_user_by_openid("all-specific")
    assert all_specific is not None
    assert all_specific.allowed_services == [ServiceType.ALL]

    update_user("all-specific", allowed_services=["技术分析", "利率"])
    partial = get_user_by_openid("all-specific")
    assert partial is not None
    assert partial.allowed_services == [ServiceType.TECHNICAL_ANALYSIS, ServiceType.RATE]


def test_user_service_excel_import_maps_fields_and_permissions_take_effect(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user, get_user_by_openid, import_users_from_excel, parse_users_excel, verify_permission

    create_user("existing-openid", name="old", allowed_services=["利率"])
    payload = _xlsx_bytes(
        [
            "openid",
            "name",
            "institution",
            "mobile",
            "enabled",
            "allowed_services",
            "auth_start_at",
            "auth_end_at",
            "remark",
        ],
        [
            [
                "existing-openid",
                "Existing New",
                "Inst B",
                "13900000000",
                "是",
                "转债",
                "2026-01-01T00:00:00",
                "2026-12-31T23:59:59",
                "renewed",
            ],
            [
                "new-openid",
                "New User",
                "Inst C",
                "13700000000",
                "true",
                "全部",
                "",
                "2026-12-31",
                "created",
            ],
        ],
    )

    rows = parse_users_excel(payload)
    assert [row.openid for row in rows] == ["existing-openid", "new-openid"]
    assert rows[0].name == "Existing New"
    assert rows[0].institution == "Inst B"
    assert rows[0].mobile == "13900000000"
    assert rows[0].enabled is True
    assert rows[0].allowed_services == "转债"
    assert rows[0].auth_start_at == datetime(2025, 12, 31, 16, 0)
    assert rows[0].auth_end_at == datetime(2026, 12, 30, 16, 0)
    assert rows[0].remark == "renewed"

    result = import_users_from_excel(payload)
    assert result.created == 1
    assert result.updated == 1
    existing = get_user_by_openid("existing-openid")
    new_user = get_user_by_openid("new-openid")
    assert existing is not None
    assert existing.name == "Existing New"
    assert new_user is not None
    assert new_user.name == "New User"
    assert verify_permission("existing-openid", ServiceType.CONVERTIBLE_BOND).allowed is True
    assert verify_permission("existing-openid", ServiceType.RATE).allowed is False
    assert verify_permission("new-openid", ServiceType.TECHNICAL_ANALYSIS).allowed is True


def test_user_service_excel_import_accepts_minimal_mobile_template_with_beijing_dates(investment_env):
    from datetime import datetime
    from business.investment.user_service import import_users_from_excel, parse_users_excel, get_user_by_openid

    payload = _xlsx_bytes(
        ["手机号", "服务权限", "授权结束日期"],
        [["13800138000", "利率", "2026-12-31"]],
    )

    rows = parse_users_excel(payload)

    assert len(rows) == 1
    assert rows[0].openid.startswith("pending-mobile-13800138000-")
    assert rows[0].mobile == "13800138000"
    assert rows[0].allowed_services == "利率"
    assert rows[0].auth_start_at is not None
    assert rows[0].auth_start_at.hour == 16
    assert rows[0].auth_start_at.minute == 0
    assert rows[0].auth_end_at == datetime(2026, 12, 30, 16, 0)

    result = import_users_from_excel(payload)
    assert result.created == 1
    assert get_user_by_openid(rows[0].openid) is not None


def test_user_service_excel_import_accepts_excel_date_cells_as_beijing_dates(investment_env):
    from datetime import datetime
    from io import BytesIO

    from openpyxl import Workbook

    from business.investment.user_service import parse_users_excel

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["手机号", "服务权限", "授权开始日期", "授权结束日期"])
    sheet.append(["13800138000", "利率", datetime(2026, 6, 1), datetime(2026, 12, 31)])
    output = BytesIO()
    workbook.save(output)

    rows = parse_users_excel(output.getvalue())

    assert rows[0].auth_start_at == datetime(2026, 5, 31, 16, 0)
    assert rows[0].auth_end_at == datetime(2026, 12, 30, 16, 0)


def test_user_service_excel_import_reports_invalid_date_with_row_and_field(investment_env):
    from business.investment.user_service import parse_users_excel

    payload = _xlsx_bytes(
        ["手机号", "服务权限", "授权结束日期"],
        [["13800138000", "利率", "2026-6-122"]],
    )

    with pytest.raises(ValueError) as error:
        parse_users_excel(payload)

    message = str(error.value)
    assert "Excel row 2 invalid date field: auth_end_at" in message
    assert "YYYY-MM-DD" in message
    assert "unconverted data remains" not in message

    invalid_month_payload = _xlsx_bytes(
        ["手机号", "服务权限", "授权结束日期"],
        [["13800138000", "利率", "2026-13-01"]],
    )
    with pytest.raises(ValueError) as month_error:
        parse_users_excel(invalid_month_payload)

    assert "Excel row 2 invalid date field: auth_end_at" in str(month_error.value)

    short_number_payload = _xlsx_bytes(
        ["手机号", "服务权限", "授权结束日期"],
        [["13800138000", "利率", "1"]],
    )
    with pytest.raises(ValueError) as number_error:
        parse_users_excel(short_number_payload)

    assert "Excel row 2 invalid date field: auth_end_at" in str(number_error.value)


def test_user_import_template_headers_are_parseable(investment_env):
    from business.investment.export_service import export_users_import_template_xlsx
    from business.investment.user_service import parse_users_excel

    rows = parse_users_excel(export_users_import_template_xlsx())

    assert len(rows) == 1
    assert rows[0].openid.startswith("pending-mobile-13800000000-")
    assert rows[0].allowed_services == "全部"


def test_web_user_import_parses_before_confirm_and_then_commits(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user, get_user_by_openid
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentUsersImportHandler

    create_user("existing-import", name="Old", allowed_services=[ServiceType.RATE])
    payload = _xlsx_bytes(
        [
            "openid",
            "name",
            "institution",
            "mobile",
            "enabled",
            "allowed_services",
            "auth_start_at",
            "auth_end_at",
            "remark",
        ],
        [
            ["existing-import", "Existing New", "Inst A", "13800000000", "启用", "全部", "", "2026-12-31", "updated"],
            ["new-import", "New User", "Inst B", "13900000000", "启用", "利率", "", "2026-12-31", "created"],
        ],
    )

    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)
    monkeypatch.setattr(web_channel.web.ctx, "headers", [], raising=False)
    monkeypatch.setattr(
        web_channel.web.webapi,
        "rawinput",
        lambda **_kwargs: {"file": SimpleNamespace(filename="customers.xlsx", value=payload), "commit": "false"},
    )

    parsed = json.loads(InvestmentUsersImportHandler().POST())

    assert parsed["status"] == "success"
    assert parsed["committed"] is False
    assert parsed["parsed"] == 2
    assert parsed["new_users"] == 1
    assert parsed["created"] == 0
    assert parsed["updated"] == 0
    assert [row["openid"] for row in parsed["preview"]] == ["existing-import", "new-import"]
    assert get_user_by_openid("existing-import").name == "Old"
    assert get_user_by_openid("new-import") is None

    monkeypatch.setattr(
        web_channel.web.webapi,
        "rawinput",
        lambda **_kwargs: {"file": SimpleNamespace(filename="customers.xlsx", value=payload), "commit": "true"},
    )
    committed = json.loads(InvestmentUsersImportHandler().POST())

    assert committed["status"] == "success"
    assert committed["committed"] is True
    assert committed["parsed"] == 2
    assert committed["new_users"] == 1
    assert committed["created"] == 1
    assert committed["updated"] == 1
    assert get_user_by_openid("existing-import").name == "Existing New"
    assert get_user_by_openid("new-import").name == "New User"


def test_records_save_failure_success_and_order(investment_env):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.records import (
        append_request_warning,
        create_request_record,
        fail_request_record,
        get_request_record,
        list_request_records,
        succeed_request_record,
    )

    first = create_request_record("openid", "利率", ServiceType.RATE)
    succeed_request_record(first, output_files=["/tmp/rate.png"], elapsed_ms=12)
    second = create_request_record("openid", "转债", ServiceType.CONVERTIBLE_BOND)
    fail_request_record(second, ErrorCode.NO_CONTENT, "今日内容尚未更新，请稍后再试。", "no active content", 5)

    assert get_request_record(first).output_files == ["/tmp/rate.png"]
    failed = get_request_record(second)
    assert failed.error_code == ErrorCode.NO_CONTENT
    assert failed.error_message == "no active content"
    assert [record.request_id for record in list_request_records(limit=2)] == [second, first]

    append_request_warning(first, "微信图片上传失败：Error code: 40164, message: invalid ip 14.153.6.203 not in whitelist")
    warned = get_request_record(first)
    assert "40164" in warned.error_message
    assert "14.153.6.203" in warned.status_warning


def test_request_record_delivery_status_is_business_facing(investment_env):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.records import (
        append_request_warning,
        create_request_record,
        fail_request_record,
        get_request_record,
        mark_request_delivered,
        succeed_request_record,
    )

    generated = create_request_record("openid", "技术分析", ServiceType.TECHNICAL_ANALYSIS)
    succeed_request_record(generated, output_files=["/tmp/report.png"], elapsed_ms=10)
    assert get_request_record(generated).delivery_status == "待客户领取"

    mark_request_delivered(generated)
    delivered = get_request_record(generated)
    assert delivered.delivery_status == "已交付"
    assert delivered.delivery_detail == "客户已收到结果"
    assert delivered.error_message == ""
    assert delivered.status_warning == ""

    append_request_warning(generated, "图片上传失败：Error code: 40164")
    failed_delivery = get_request_record(generated)
    assert failed_delivery.delivery_status == "交付异常"
    assert failed_delivery.delivery_detail == "图片上传失败：Error code: 40164"

    failed = create_request_record("openid", "利率", ServiceType.RATE)
    fail_request_record(failed, ErrorCode.NO_CONTENT, "今日内容尚未更新，请稍后再试。", "no active content", 5)
    failed_record = get_request_record(failed)
    assert failed_record.delivery_status == "未交付"
    assert failed_record.delivery_detail == "no active content"


def test_export_request_records_hides_internal_delivery_marker(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.export_service import export_request_records_xlsx
    from business.investment.records import create_request_record, mark_request_delivered, succeed_request_record

    request_id = create_request_record("openid", "利率", ServiceType.RATE)
    succeed_request_record(request_id, output_files=[], elapsed_ms=12)
    mark_request_delivered(request_id)

    rows = _xlsx_sheet_rows(export_request_records_xlsx("", "", service_type="rate"))

    assert rows[1][4] == "利率"
    assert rows[1][10] in ("", None)
    assert "[delivery:delivered]" not in str(rows[1])


def test_old_generating_request_records_are_flagged_without_mutating_status(investment_env):
    from business.investment.constants import ServiceType, Status
    from business.investment.db import connect
    from business.investment.records import create_request_record, get_request_record, list_request_records

    request_id = create_request_record("openid", "新易盛 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    old_created_at = (datetime.now(UTC) - timedelta(minutes=31)).isoformat(timespec="microseconds")
    with connect() as conn:
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :updated_at "
                "where request_id = :request_id"
            ),
            {"created_at": old_created_at, "updated_at": old_created_at, "request_id": request_id},
        )

    record = list_request_records(limit=1)[0]
    persisted = get_request_record(request_id)

    assert record.request_id == request_id
    assert record.status == Status.GENERATING
    assert record.status_warning == "未完成/可能超时"
    assert persisted.status == Status.GENERATING


def test_job_service_reuses_running_technical_analysis_record(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.job_service import find_running_job, start_job_if_absent
    from business.investment.records import succeed_request_record

    first = start_job_if_absent("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    duplicate = start_job_if_absent("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS)

    assert first.created is True
    assert duplicate.created is False
    assert duplicate.record.request_id == first.record.request_id
    assert find_running_job("openid", "300502.SZ 技术分析").request_id == first.record.request_id

    succeed_request_record(first.record.request_id, output_files=["/tmp/signal.png", "/tmp/chart.png"], elapsed_ms=12)
    next_job = start_job_if_absent("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS)

    assert next_job.created is True
    assert next_job.record.request_id != first.record.request_id


def test_job_service_reuses_running_cache_job_across_users(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.job_service import find_running_cache_job, start_cache_job_if_absent
    from business.investment.records import succeed_request_record

    cache_key = "technical_analysis:300502.SZ:2026-05-25:v1"

    first = start_cache_job_if_absent(
        "openid-a",
        "300502.SZ 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        cache_key,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
    )
    duplicate = start_cache_job_if_absent(
        "openid-b",
        "新易盛 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        cache_key,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
    )

    assert first.created is True
    assert duplicate.created is False
    assert duplicate.record.request_id == first.record.request_id
    assert find_running_cache_job(cache_key, ServiceType.TECHNICAL_ANALYSIS).request_id == first.record.request_id
    assert first.record.cache_key == cache_key
    assert first.record.normalized_target == "300502.SZ"
    assert first.record.market_date == "2026-05-25"

    succeed_request_record(first.record.request_id, output_files=["/tmp/signal.png", "/tmp/chart.png"], elapsed_ms=12)
    next_job = start_cache_job_if_absent(
        "openid-b",
        "新易盛 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        cache_key,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
    )

    assert next_job.created is True
    assert next_job.record.request_id != first.record.request_id


def test_job_service_marks_stale_running_cache_job_failed_and_allows_new_job(investment_env):
    from business.investment.constants import ServiceType, Status
    from business.investment.db import connect
    from business.investment.job_service import find_running_cache_job, start_cache_job_if_absent
    from business.investment.records import get_request_record

    cache_key = "technical_analysis:300502.SZ:2026-05-25:v1"
    first = start_cache_job_if_absent(
        "openid-a",
        "300502.SZ 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        cache_key,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
    )
    stale_time = (datetime.now(UTC) - timedelta(minutes=31)).isoformat(timespec="microseconds")
    with connect() as conn:
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at=:created_at, updated_at=:created_at where request_id=:request_id"
            ),
            {"created_at": stale_time, "request_id": first.record.request_id},
        )

    assert find_running_cache_job(cache_key, ServiceType.TECHNICAL_ANALYSIS) is None
    assert get_request_record(first.record.request_id).status == Status.FAILED

    next_job = start_cache_job_if_absent(
        "openid-b",
        "新易盛 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        cache_key,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
    )

    assert next_job.created is True
    assert next_job.record.request_id != first.record.request_id


def test_job_service_allows_different_cache_keys_to_run_together(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.job_service import start_cache_job_if_absent

    first = start_cache_job_if_absent(
        "openid-a",
        "300502.SZ 2026-05-25 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        "technical_analysis:300502.SZ:2026-05-25:v1",
        normalized_target="300502.SZ",
        market_date="2026-05-25",
    )
    second = start_cache_job_if_absent(
        "openid-b",
        "300502.SZ 2026-05-26 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        "technical_analysis:300502.SZ:2026-05-26:v1",
        normalized_target="300502.SZ",
        market_date="2026-05-26",
    )

    assert first.created is True
    assert second.created is True
    assert second.record.request_id != first.record.request_id


def test_technical_analysis_exception_marks_record_failed_and_unblocks_running_job(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.constants import ErrorCode, ServiceType, Status
    from business.investment.job_service import find_running_job
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    monkeypatch.setattr(config_service, "conf", lambda: {"custom_api_key": "sk-secret-123"})
    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])

    reply = handle_text_message(
        "ok",
        "300502.SZ 技术分析",
        technical_analysis_handler=lambda *_args: (_ for _ in ()).throw(RuntimeError("backend boom sk-secret-123")),
    )

    record = list_request_records(limit=1)[0]
    assert reply.success is False
    assert reply.error_code == ErrorCode.SYSTEM_ERROR
    assert record.status == Status.FAILED
    assert record.error_code == ErrorCode.SYSTEM_ERROR
    assert "backend boom" in record.error_message
    assert "sk-secret-123" not in record.error_message
    assert find_running_job("ok", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS) is None


def test_job_service_ignores_stale_running_technical_analysis_record(investment_env):
    from datetime import UTC, datetime, timedelta

    from sqlalchemy import text

    from business.investment.constants import ServiceType, Status
    from business.investment.db import connect
    from business.investment.job_service import find_running_job, start_job_if_absent
    from business.investment.records import create_request_record, get_request_record

    old_id = create_request_record("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    stale_time = (datetime.now(UTC) - timedelta(minutes=31)).isoformat(timespec="microseconds")
    with connect() as conn:
        conn.execute(
            text("update investment_request_records set created_at=:created_at, updated_at=:created_at where request_id=:request_id"),
            {"created_at": stale_time, "request_id": old_id},
        )

    assert find_running_job("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS) is None
    assert get_request_record(old_id).status == Status.FAILED

    next_job = start_job_if_absent("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS)

    assert next_job.created is True
    assert next_job.record.request_id != old_id


def test_job_service_concurrent_start_creates_single_running_job(investment_env):
    from concurrent.futures import ThreadPoolExecutor

    from business.investment.constants import ServiceType
    from business.investment.job_service import start_job_if_absent
    from business.investment.records import list_request_records

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda _index: start_job_if_absent("openid", "300502.SZ 技术分析", ServiceType.TECHNICAL_ANALYSIS),
                range(2),
            )
        )

    assert sum(1 for result in results if result.created) == 1
    assert len({result.record.request_id for result in results}) == 1
    records = list_request_records(limit=10)
    running = [record for record in records if record.raw_input == "300502.SZ 技术分析"]
    assert len(running) == 1


def test_router_concurrent_technical_analysis_reuses_running_job_without_duplicate_handler(
    investment_env,
    tmp_path,
):
    from concurrent.futures import ThreadPoolExecutor
    import time

    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.technical_analysis import TechnicalAnalysisResult
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    card = tmp_path / "signal.png"
    chart = tmp_path / "chart.png"
    report = tmp_path / "report.md"
    card.write_bytes(b"card")
    chart.write_bytes(b"chart")
    report.write_text("report", encoding="utf-8")
    calls = []

    def slow_handler(_openid, _raw_input, _target):
        calls.append(_target)
        time.sleep(0.2)
        return TechnicalAnalysisResult(
            True,
            signal_card_path=str(card),
            main_chart_path=str(chart),
            report_path=str(report),
            output_files=[str(card), str(chart), str(report)],
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        replies = list(
            pool.map(
                lambda _index: handle_text_message(
                    "ok",
                    "300502.SZ 技术分析",
                    technical_analysis_handler=slow_handler,
                ),
                range(2),
            )
        )

    assert len(calls) == 1
    assert sum(1 for reply in replies if reply.success) == 1
    waiting = [reply for reply in replies if not reply.success]
    assert len(waiting) == 1
    assert waiting[0].reply_text == "正在运行，请稍候。"
    records = [record for record in list_request_records(limit=10) if record.raw_input == "300502.SZ 技术分析"]
    assert len(records) == 1


def test_router_concurrent_technical_analysis_reuses_running_cache_job_across_users(
    investment_env,
    tmp_path,
    monkeypatch,
):
    from concurrent.futures import ThreadPoolExecutor
    import time

    from business.investment import technical_analysis
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.technical_analysis import TechnicalAnalysisResult
    from business.investment.user_service import create_user

    create_user("openid-a", enabled=True, allowed_services=[ServiceType.ALL])
    create_user("openid-b", enabled=True, allowed_services=[ServiceType.ALL])
    card = tmp_path / "signal.png"
    chart = tmp_path / "chart.png"
    report = tmp_path / "report.md"
    card.write_bytes(b"card")
    chart.write_bytes(b"chart")
    report.write_text("report", encoding="utf-8")
    calls = []

    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v1", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    def slow_handler(_openid, _raw_input, _target):
        calls.append(_openid)
        time.sleep(0.2)
        return TechnicalAnalysisResult(
            True,
            signal_card_path=str(card),
            main_chart_path=str(chart),
            report_path=str(report),
            output_files=[str(card), str(chart), str(report)],
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        replies = list(
            pool.map(
                lambda openid: handle_text_message(
                    openid,
                    "300502.SZ 技术分析",
                    technical_analysis_handler=slow_handler,
                ),
                ["openid-a", "openid-b"],
            )
        )

    assert len(calls) == 1
    assert sum(1 for reply in replies if reply.success) == 1
    waiting = [reply for reply in replies if not reply.success]
    assert len(waiting) == 1
    assert waiting[0].reply_text == "正在运行，请稍候。"
    records = [record for record in list_request_records(limit=10) if record.raw_input == "300502.SZ 技术分析"]
    assert len(records) == 1
    assert records[0].cache_key
    assert records[0].normalized_target == "300502.SZ"
    assert records[0].market_date == "2026-05-25"


def test_request_records_api_includes_generating_timeout_warning(investment_env, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.records import create_request_record
    from channel.web.web_channel import InvestmentRequestRecordsHandler

    request_id = create_request_record("openid", "新易盛 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    old_created_at = (datetime.now(UTC) - timedelta(minutes=31)).isoformat(timespec="microseconds")
    with connect() as conn:
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :updated_at "
                "where request_id = :request_id"
            ),
            {"created_at": old_created_at, "updated_at": old_created_at, "request_id": request_id},
        )

    payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentRequestRecordsHandler().GET,
        params={"limit": "50"},
    )
    record = payload["records"][0]

    assert record["request_id"] == request_id
    assert record["status"] == "generating"
    assert record["status_warning"] == "未完成/可能超时"


def test_batch_01_service_results_share_contract_fields(investment_env):
    from business.investment.ai_generation import AIGenerationResult
    from business.investment.daily_content import DailyContentResult
    from business.investment.render_service import RenderResult
    from business.investment.router import BusinessReply
    from business.investment.technical_analysis import TechnicalAnalysisResult

    required = {"success", "error_code", "user_prompt", "detail", "output_files"}
    results = [
        AIGenerationResult(success=False),
        RenderResult(success=False),
        TechnicalAnalysisResult(success=False),
        DailyContentResult(success=False),
        BusinessReply(handled=True, success=False, reply_text="", output_files=[], service_type="rate"),
    ]

    for result in results:
        assert required.issubset(result.__dataclass_fields__)


def test_success_request_records_output_files_table(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.records import record_success_request

    request_id = record_success_request("openid", "利率", ServiceType.RATE, ["/tmp/rate.png"], elapsed_ms=3)

    with connect() as conn:
        rows = conn.execute(
            text(
                "select owner_id, file_path, file_type, service_type "
                "from investment_output_files where owner_id = :request_id"
            ),
            {"request_id": request_id},
        ).mappings().all()

    assert [dict(row) for row in rows] == [
        {
            "owner_id": request_id,
            "file_path": "/tmp/rate.png",
            "file_type": "image",
            "service_type": ServiceType.RATE,
        }
    ]


def test_success_request_archives_generated_images_and_documents(investment_env, tmp_path):
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, get_request_record, list_output_files, succeed_request_record
    from business.investment.storage import get_storage_dirs

    image = tmp_path / "rate_card.png"
    report = tmp_path / "rate_report.md"
    image.write_bytes(b"image-v1")
    report.write_text("report-v1", encoding="utf-8")
    request_id = create_request_record("openid", "利率", ServiceType.RATE)

    succeed_request_record(
        request_id,
        output_files=[str(image), str(report)],
        elapsed_ms=3,
        artifact_roles={str(image): "output_image", str(report): "markdown_report"},
    )

    record = get_request_record(request_id)
    artifacts = list_output_files(request_id)
    archive_root = get_storage_dirs()["generated"] / "archive"

    assert len(record.output_files or []) == 2
    assert all(Path(path).is_file() for path in record.output_files or [])
    assert all(Path(path).resolve().is_relative_to(archive_root.resolve()) for path in record.output_files or [])
    assert record.output_files != [str(image), str(report)]
    assert Path(record.output_files[0]).read_bytes() == b"image-v1"
    assert Path(record.output_files[1]).read_text(encoding="utf-8") == "report-v1"
    assert [(item["file_path"], item["file_type"], item["artifact_role"]) for item in artifacts] == [
        (record.output_files[0], "image", "output_image"),
        (record.output_files[1], "markdown", "markdown_report"),
    ]


def test_request_records_save_audit_metadata(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, get_request_record, succeed_request_record

    request_id = create_request_record(
        "openid",
        "新易盛 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        stock_code="300502.SZ",
        stock_name="新易盛",
        customer_name="Alice",
        institution="Inst A",
        market_date="2026-05-25",
        cache_key="ta:300502.SZ:2026-05-25",
        cache_hit=True,
        program_version="sha256:program12345678",
        ta_version="sha256:ta123456789012",
        renderer_version="sha256:renderer123456",
        template_version="sha256:template123456",
    )
    succeed_request_record(request_id, output_files=[], elapsed_ms=12)

    record = get_request_record(request_id)

    assert record.normalized_target == "300502.SZ"
    assert record.stock_code == "300502.SZ"
    assert record.stock_name == "新易盛"
    assert record.customer_name == "Alice"
    assert record.institution == "Inst A"
    assert record.market_date == "2026-05-25"
    assert record.cache_key == "ta:300502.SZ:2026-05-25"
    assert record.cache_hit is True
    assert record.program_version == "sha256:program12345678"
    assert record.ta_version == "sha256:ta123456789012"
    assert record.renderer_version == "sha256:renderer123456"
    assert record.template_version == "sha256:template123456"


def test_unauthorized_request_service_type_is_normalized_and_labeled():
    from business.investment.constants import SERVICE_LABELS, ServiceType, normalize_service

    assert ServiceType.UNAUTHORIZED_REQUEST == "unauthorized_request"
    assert SERVICE_LABELS[ServiceType.UNAUTHORIZED_REQUEST] == "无权限请求"
    assert normalize_service("unauthorized_request") == ServiceType.UNAUTHORIZED_REQUEST
    assert normalize_service("无权限请求") == ServiceType.UNAUTHORIZED_REQUEST


def _xlsx_sheet_rows(content: bytes) -> list[list[object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_export_date_range_helpers_return_inclusive_bounds():
    from business.investment.export_service import month_range, quarter_range

    assert month_range(2026, 2) == ("2026-02-01T00:00:00", "2026-02-28T23:59:59")
    assert month_range(2024, 2) == ("2024-02-01T00:00:00", "2024-02-29T23:59:59")
    assert quarter_range(2026, 2) == ("2026-04-01T00:00:00", "2026-06-30T23:59:59")


def test_investment_date_bound_treats_plain_dates_as_beijing_days():
    from channel.web.web_channel import _investment_date_bound, _investment_month_bounds, _investment_quarter_bounds

    assert _investment_date_bound("2026-05-31") == "2026-05-30T16:00:00"
    assert _investment_date_bound("2026-05-31", end=True) == "2026-05-31T15:59:59.999999"
    assert _investment_date_bound("2026-05-31T12:30:00") == "2026-05-31T12:30:00"
    assert _investment_date_bound("2026-05-31 12:30:00", end=True) == "2026-05-31 12:30:00"
    assert _investment_month_bounds(2026, 5) == ("2026-04-30T16:00:00", "2026-05-31T15:59:59.999999")
    assert _investment_quarter_bounds(2026, 2) == ("2026-03-31T16:00:00", "2026-06-30T15:59:59.999999")


def test_export_request_records_xlsx_filters_and_includes_audit_fields(investment_env):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.db import connect
    from business.investment.export_service import export_request_records_xlsx
    from business.investment.records import create_request_record, fail_request_record, succeed_request_record

    older = create_request_record("old-openid", "利率", ServiceType.RATE)
    succeed_request_record(older, output_files=["/tmp/old.png"], elapsed_ms=1)
    included = create_request_record(
        "openid-a",
        "新易盛 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        stock_code="300502.SZ",
        stock_name="新易盛",
        customer_name="Alice",
        institution="Inst A",
        cache_hit=True,
        program_version="sha256:program",
        template_version="sha256:template",
    )
    succeed_request_record(included, output_files=["/tmp/card.png"], elapsed_ms=15)
    failed = create_request_record("openid-b", "贵州茅台 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    fail_request_record(failed, ErrorCode.STOCK_NOT_FOUND, detail="not found", elapsed_ms=5)

    with connect() as conn:
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :created_at "
                "where request_id = :request_id"
            ),
            {"created_at": "2026-04-30T23:59:59", "request_id": older},
        )
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :created_at "
                "where request_id = :request_id"
            ),
            {"created_at": "2026-05-10T08:00:00", "request_id": included},
        )
        conn.execute(
            text(
                "update investment_request_records "
                "set created_at = :created_at, updated_at = :created_at "
                "where request_id = :request_id"
            ),
            {"created_at": "2026-06-01T00:00:00", "request_id": failed},
        )

    rows = _xlsx_sheet_rows(
        export_request_records_xlsx(
            "2026-05-01T00:00:00",
            "2026-05-31T23:59:59",
            service_type=ServiceType.TECHNICAL_ANALYSIS,
            status="success",
            keyword="Alice",
        )
    )

    assert rows == [
        [
            "请求时间",
            "OpenID",
            "客户姓名",
            "机构",
            "原始输入",
            "服务类型",
            "股票代码",
            "股票名称",
            "状态",
            "错误码",
            "错误原因",
            "缓存命中",
            "输出文件",
            "耗时毫秒",
            "程序版本",
            "模板版本",
        ],
        [
            "2026-05-10T08:00:00",
            "openid-a",
            "Alice",
            "Inst A",
            "新易盛 技术分析",
            "technical_analysis",
            "300502.SZ",
            "新易盛",
            "success",
            None,
            None,
            "是",
            "/tmp/card.png",
            15,
            "sha256:program",
            "sha256:template",
        ],
    ]


def test_export_request_records_xlsx_filters_by_service_customer_and_range(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.export_service import export_request_records_xlsx
    from business.investment.records import create_request_record, succeed_request_record
    from business.investment.user_service import create_user

    create_user("openid-a", name="Alice", mobile="13800000000", enabled=True, allowed_services=[ServiceType.ALL])
    create_user("openid-b", name="Bob", mobile="13900000000", enabled=True, allowed_services=[ServiceType.ALL])
    included = create_request_record("openid-a", "新易盛 技术分析", ServiceType.TECHNICAL_ANALYSIS, stock_name="新易盛")
    wrong_service = create_request_record("openid-a", "利率", ServiceType.RATE)
    wrong_user = create_request_record("openid-b", "贵州茅台 技术分析", ServiceType.TECHNICAL_ANALYSIS)
    succeed_request_record(included, output_files=["/tmp/card.png"], elapsed_ms=1)
    succeed_request_record(wrong_service, output_files=["/tmp/rate.png"], elapsed_ms=1)
    succeed_request_record(wrong_user, output_files=["/tmp/other.png"], elapsed_ms=1)
    with connect() as conn:
        for request_id in [included, wrong_service, wrong_user]:
            conn.execute(
                text(
                    "update investment_request_records "
                    "set created_at = '2026-05-10T08:00:00' "
                    "where request_id = :request_id"
                ),
                {"request_id": request_id},
            )

    rows = _xlsx_sheet_rows(
        export_request_records_xlsx(
            "2026-05-01T00:00:00",
            "2026-05-31T23:59:59",
            service_type=ServiceType.TECHNICAL_ANALYSIS,
            customer="13800000000",
        )
    )

    assert len(rows) == 2
    assert rows[1][1] == "openid-a"
    assert rows[1][4] == "新易盛 技术分析"


def test_export_request_records_xlsx_unknown_service_returns_only_header(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.export_service import export_request_records_xlsx
    from business.investment.records import create_request_record, succeed_request_record

    request_id = create_request_record("openid", "利率", ServiceType.RATE)
    succeed_request_record(request_id, output_files=["/tmp/rate.png"], elapsed_ms=1)

    rows = _xlsx_sheet_rows(export_request_records_xlsx("", "", service_type="unknown-service"))

    assert rows == [
        [
            "请求时间",
            "OpenID",
            "客户姓名",
            "机构",
            "原始输入",
            "服务类型",
            "股票代码",
            "股票名称",
            "状态",
            "错误码",
            "错误原因",
            "缓存命中",
            "输出文件",
            "耗时毫秒",
            "程序版本",
            "模板版本",
        ]
    ]


def test_export_request_records_xlsx_empty_records_contains_only_header(investment_env):
    from business.investment.export_service import export_request_records_xlsx

    rows = _xlsx_sheet_rows(export_request_records_xlsx("2026-05-01T00:00:00", "2026-05-31T23:59:59"))

    assert len(rows) == 1
    assert rows[0][0] == "请求时间"


def test_export_users_xlsx_filters_enabled_users(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.export_service import export_users_xlsx
    from business.investment.user_service import create_user

    create_user(
        "enabled-openid",
        name="Enabled",
        institution="Inst A",
        mobile="13800000000",
        enabled=True,
        allowed_services=[ServiceType.RATE],
        auth_start_at="2026-01-01T00:00:00",
        auth_end_at="2026-12-31T23:59:59",
        remark="ok",
    )
    create_user("disabled-openid", name="Disabled", enabled=False, allowed_services=[ServiceType.ALL])

    rows = _xlsx_sheet_rows(export_users_xlsx(enabled=True))

    assert rows == [
        ["OpenID", "姓名", "机构", "手机号", "状态", "服务权限", "授权开始", "授权结束", "备注"],
        [
            "enabled-openid",
            "Enabled",
            "Inst A",
            "13800000000",
            "启用",
            "rate",
            "2026-01-01T00:00:00",
            "2026-12-31T23:59:59",
            "ok",
        ],
    ]


def test_web_export_handlers_return_xlsx_downloads(investment_env, monkeypatch):
    from channel.web import web_channel
    from channel.web.web_channel import InvestmentRequestRecordsExportHandler, InvestmentUsersExportHandler

    headers = []
    _login_default_investment_admin(monkeypatch)
    monkeypatch.setattr(web_channel.web, "header", lambda name, value: headers.append((name, value)))
    monkeypatch.setattr(
        web_channel.web,
        "input",
        lambda **_defaults: SimpleNamespace(start_date="2026-05-01", end_date="2026-05-31", service_type="", enabled=""),
    )

    request_data = InvestmentRequestRecordsExportHandler().GET()

    assert isinstance(request_data, bytes)
    assert _xlsx_sheet_rows(request_data)[0][0] == "请求时间"
    assert ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") in headers
    assert any(
        name == "Content-Disposition" and "investment-requests" in value and value.endswith(".xlsx\"")
        for name, value in headers
    )

    headers.clear()
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(enabled="true"))

    user_data = InvestmentUsersExportHandler().GET()

    assert isinstance(user_data, bytes)
    assert _xlsx_sheet_rows(user_data)[0][0] == "OpenID"
    assert ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") in headers
    assert any(
        name == "Content-Disposition" and "investment-users" in value and value.endswith(".xlsx\"")
        for name, value in headers
    )


def test_artifact_service_records_role_size_hash_and_version(investment_env, tmp_path):
    from business.investment.artifact_service import record_artifact
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.versioning import file_fingerprint

    artifact = tmp_path / "card.png"
    artifact.write_bytes(b"card-bytes")

    record_artifact(
        "request-1",
        str(artifact),
        "signal_card",
        ServiceType.TECHNICAL_ANALYSIS,
        version_tag="sha256:renderer123456",
    )

    with connect() as conn:
        row = conn.execute(
            text(
                "select owner_id, file_path, file_type, service_type, artifact_role, file_size, file_hash, version_tag "
                "from investment_output_files where owner_id = :owner_id"
            ),
            {"owner_id": "request-1"},
        ).mappings().one()

    assert row["file_path"] == str(artifact)
    assert row["file_type"] == "image"
    assert row["service_type"] == ServiceType.TECHNICAL_ANALYSIS
    assert row["artifact_role"] == "signal_card"
    assert row["file_size"] == len(b"card-bytes")
    assert row["file_hash"] == file_fingerprint(str(artifact))
    assert row["version_tag"] == "sha256:renderer123456"
    assert file_fingerprint(str(tmp_path / "missing.png")).startswith("missing:")


def test_ai_and_renderer_failures_record_sanitized_backend_detail(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.config_service import safe_log_value
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, regenerate_content
    from business.investment.records import get_content_record

    api_key = "sk-live-secret-1234567890"
    monkeypatch.setattr(config_service, "conf", lambda: {"custom_api_key": api_key})

    ai_content_id = create_content_draft(ServiceType.RATE, source_text="rate")
    ai_result = regenerate_content(
        ai_content_id,
        lambda _service_type, _source_text: SimpleNamespace(success=False, detail=f"upstream rejected {api_key}"),
        lambda _service_type, _text: SimpleNamespace(success=True, image_path="/tmp/unused.png"),
    )
    ai_record = get_content_record(ai_content_id)
    assert ai_result.success is False
    assert "upstream rejected" in ai_record.error_message
    assert api_key not in ai_record.error_message

    render_content_id = create_content_draft(ServiceType.RATE, source_text="rate")
    render_result = regenerate_content(
        render_content_id,
        lambda _service_type, _source_text: SimpleNamespace(success=True, text="standard"),
        lambda _service_type, _text: SimpleNamespace(success=False, detail=f"renderer failed with {api_key}"),
    )
    render_record = get_content_record(render_content_id)
    assert render_result.success is False
    assert "renderer failed" in render_record.error_message
    assert api_key not in render_record.error_message
    assert api_key not in safe_log_value("model.api_key", api_key)


def test_ai_generation_uses_global_model_params_and_ignores_legacy_investment_rows(investment_env, monkeypatch):
    from business.investment.ai_generation import (
        AIGenerationRequest,
        generate_convertible_bond_text,
        generate_rate_text,
        generate_technical_analysis_text,
    )
    from business.investment import config_service, db
    from business.investment.config_service import save_configs
    from business.investment.constants import ServiceType, Status

    api_key = "sk-global-contract-1234567890"
    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "global-model",
            "custom_api_base": "https://global.example/v1",
            "custom_api_key": api_key,
            "temperature": 0.25,
        },
    )
    now = datetime.now(UTC).isoformat(timespec="microseconds")
    with db.connect() as conn:
        for key, value in {
            "model.provider": "legacy-provider",
            "model.name": "legacy-model",
            "model.api_base": "https://legacy.example/v1",
            "model.api_key": "sk-legacy-secret",
            "model.temperature": 0.99,
        }.items():
            db.upsert_config(conn, key, json.dumps(value), now, "legacy")
    save_configs(
        {
            "prompt.technical_analysis": "TA prompt",
            "prompt.rate": "Rate prompt",
            "prompt.convertible_bond": "CB prompt",
        },
        operator_role="admin",
    )

    seen: list[AIGenerationRequest] = []

    class FakeAdapter:
        def generate(self, request: AIGenerationRequest) -> str:
            seen.append(request)
            return f"standard:{request.service_type}:{request.prompt}:{request.source_text}"

    ta = generate_technical_analysis_text("ta report", adapter=FakeAdapter())
    rate = generate_rate_text("rate source", adapter=FakeAdapter())
    cb = generate_convertible_bond_text("cb source", adapter=FakeAdapter())

    assert [request.service_type for request in seen] == [
        ServiceType.TECHNICAL_ANALYSIS,
        ServiceType.RATE,
        ServiceType.CONVERTIBLE_BOND,
    ]
    assert [request.prompt for request in seen] == ["TA prompt", "Rate prompt", "CB prompt"]
    assert all(request.model_provider == "custom" for request in seen)
    assert all(request.model_name == "global-model" for request in seen)
    assert all(request.api_base == "https://global.example/v1" for request in seen)
    assert all(request.api_key == api_key for request in seen)
    assert all(request.temperature == 0.25 for request in seen)

    assert ta.success is True
    assert ta.status == Status.SUCCESS
    assert ta.service_type == ServiceType.TECHNICAL_ANALYSIS
    assert ta.source_text == "ta report"
    assert ta.generated_text == "standard:technical_analysis:TA prompt:ta report"
    assert ta.text == ta.generated_text
    assert ta.failure_reason == ""
    assert api_key not in str(ta.model_params)
    assert ta.model_params == {
        "provider": "custom",
        "model": "global-model",
        "api_base": "https://global.example/v1",
        "api_key": "sk-g**********7890",
        "temperature": 0.25,
    }
    assert rate.generated_text == "standard:rate:Rate prompt:rate source"
    assert cb.generated_text == "standard:convertible_bond:CB prompt:cb source"


def test_technical_analysis_default_prompt_matches_signal_card_renderer_contract(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.ai_generation import build_generation_request
    from business.investment.constants import ServiceType

    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "global-model",
            "custom_api_base": "https://global.example/v1",
            "custom_api_key": "sk-global-contract-1234567890",
        },
    )

    request = build_generation_request(ServiceType.TECHNICAL_ANALYSIS, "ta report")

    for required in (
        "标的：",
        "信号方向：",
        "最新收盘：",
        "行情日期：",
        "趋势研判",
        "核心关键位",
        "强压力：",
        "强支撑：",
        "实操指引",
        "授权剩余时间：",
        "数据来源：",
        "业务对接：",
    ):
        assert required in request.prompt
    assert "禁止输出 Markdown 表格" in request.prompt
    assert "只输出卡片正文" in request.prompt


@pytest.mark.parametrize(
    ("global_config", "expected"),
    [
        (
            {
                "bot_type": "zhipu",
                "model": "glm-5.1",
                "zhipu_ai_api_base": "https://zhipu.example/v4",
                "zhipu_ai_api_key": "sk-zhipu",
                "temperature": 0.41,
            },
            {
                "provider": "zhipu",
                "model": "glm-5.1",
                "api_base": "https://zhipu.example/v4",
                "api_key": "sk-zhipu",
                "temperature": 0.41,
            },
        ),
        (
            {
                "bot_type": "",
                "model": "glm-5.1",
                "zhipu_ai_api_base": "https://zhipu-inferred.example/v4",
                "zhipu_ai_api_key": "sk-zhipu-inferred",
            },
            {
                "provider": "zhipu",
                "model": "glm-5.1",
                "api_base": "https://zhipu-inferred.example/v4",
                "api_key": "sk-zhipu-inferred",
                "temperature": 0.7,
            },
        ),
        (
            {
                "bot_type": "",
                "model": "qwen3-max",
                "dashscope_api_base": "https://dashscope.example/compatible-mode/v1",
                "dashscope_api_key": "sk-dashscope",
            },
            {
                "provider": "dashscope",
                "model": "qwen3-max",
                "api_base": "https://dashscope.example/compatible-mode/v1",
                "api_key": "sk-dashscope",
                "temperature": 0.7,
            },
        ),
        (
            {
                "use_linkai": True,
                "linkai_api_base": "https://link.example",
                "linkai_api_key": "sk-linkai",
                "model": "gpt-5.4-mini",
            },
            {
                "provider": "linkai",
                "model": "gpt-5.4-mini",
                "api_base": "https://link.example/v1",
                "api_key": "sk-linkai",
                "temperature": 0.7,
            },
        ),
        (
            {
                "bot_type": "deepseek",
                "model": "deepseek-v4-flash",
                "deepseek_api_base": "https://deepseek.example/v1",
                "custom_api_key": "sk-custom-should-not-be-used",
            },
            {
                "provider": "deepseek",
                "model": "deepseek-v4-flash",
                "api_base": "https://deepseek.example/v1",
                "api_key": "",
                "temperature": 0.7,
            },
        ),
    ],
)
def test_global_model_config_resolves_configured_and_inferred_providers(monkeypatch, global_config, expected):
    from business.investment import config_service
    from business.investment.ai_generation import _global_model_config

    monkeypatch.setattr(config_service, "conf", lambda: global_config)

    assert _global_model_config() == expected


def test_ai_generation_sends_image_source_files_as_multimodal_content(investment_env, tmp_path, monkeypatch):
    from business.investment import config_service
    from business.investment.ai_generation import ExistingModelAdapter, generate_rate_text

    image_bytes = b"\x89PNG\r\n\x1a\nimage"
    image_path = tmp_path / "rate-source.png"
    image_path.write_bytes(image_bytes)
    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "vision-model",
            "custom_api_base": "https://model.example/v1",
            "custom_api_key": "sk-image-source-1234567890",
        },
    )

    class FakeClient:
        def __init__(self):
            self.calls = []

        def chat_completions(self, **kwargs):
            self.calls.append(kwargs)
            return {"choices": [{"message": {"content": "standard rate text"}}]}

    client = FakeClient()

    result = generate_rate_text("", source_files=[str(image_path)], adapter=ExistingModelAdapter(client=client))

    assert result.success is True
    messages = client.calls[0]["messages"]
    assert messages[1]["role"] == "user"
    content = messages[1]["content"]
    assert content[0]["type"] == "text"
    assert str(image_path) in content[0]["text"]
    assert content[1] == {
        "type": "image_url",
        "image_url": {
            "url": f"data:image/png;base64,{base64.b64encode(image_bytes).decode('ascii')}",
        },
    }


def test_ai_generation_retries_transient_model_connection_errors(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.ai_generation import ExistingModelAdapter, generate_rate_text
    from models.openai.openai_http_client import OpenAIHTTPError

    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "mimo-v2.5",
            "custom_api_base": "https://model.example/v1",
            "custom_api_key": "sk-retry-source-1234567890",
        },
    )

    class FlakyClient:
        def __init__(self):
            self.calls = 0

        def chat_completions(self, **kwargs):
            self.calls += 1
            if self.calls == 1:
                raise OpenAIHTTPError(0, {}, "Connection error: SSL EOF")
            return {"choices": [{"message": {"content": "standard rate text"}}]}

    client = FlakyClient()

    result = generate_rate_text("rate source", adapter=ExistingModelAdapter(client=client))

    assert result.success is True
    assert result.text == "standard rate text"
    assert client.calls == 2


def test_ai_generation_extracts_image_text_before_final_card_prompt(investment_env, tmp_path, monkeypatch):
    from business.investment import config_service
    from business.investment.ai_generation import ExistingModelAdapter, generate_rate_text

    image_path = tmp_path / "rate-source.png"
    image_path.write_bytes(b"\x89PNG\r\n\x1a\nimage")
    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "vision-model",
            "custom_api_base": "https://model.example/v1",
            "custom_api_key": "sk-image-source-1234567890",
        },
    )

    class FakeClient:
        def __init__(self):
            self.calls = []

        def chat_completions(self, **kwargs):
            self.calls.append(kwargs)
            if len(self.calls) == 1:
                return {"choices": [{"message": {"content": "OCR: 2026-05-25 108.970 入场（3/8）"}}]}
            return {"choices": [{"message": {"content": "standard rate card text"}}]}

    client = FakeClient()

    result = generate_rate_text("", source_files=[str(image_path)], adapter=ExistingModelAdapter(client=client))

    assert result.success is True
    assert result.text == "standard rate card text"
    assert len(client.calls) == 2
    assert any(block.get("type") == "image_url" for block in client.calls[0]["messages"][1]["content"])
    assert client.calls[1]["messages"][1]["content"] == (
        "以下是上传图片的识别结果，请据此生成标准卡片文本。禁止要求用户再提供原文；"
        "缺失字段按系统要求填“——”。\n\nOCR: 2026-05-25 108.970 入场（3/8）"
    )


def test_ai_generation_blank_configured_prompt_falls_back_to_default(investment_env):
    from business.investment.ai_generation import AIGenerationRequest, generate_rate_text
    from business.investment.config_service import save_config

    save_config("prompt.rate", "", operator_role="admin")
    seen: list[AIGenerationRequest] = []

    class FakeAdapter:
        def generate(self, request: AIGenerationRequest) -> str:
            seen.append(request)
            return "standard rate text"

    result = generate_rate_text("rate source", adapter=FakeAdapter())

    assert result.success is True
    assert seen[0].prompt
    assert "必须只输出卡片正文" in seen[0].prompt
    assert "当日核心信号" in seen[0].prompt


def test_ai_generation_normalizes_semicolon_rate_text_for_renderer(investment_env):
    from business.investment.ai_generation import AIGenerationRequest, generate_rate_text

    class FakeAdapter:
        def generate(self, request: AIGenerationRequest) -> str:
            return (
                "标的：T主力合约；最新收盘：108.970元；行情日期：2026-05-25；分析模型：交易性择时日度信号跟踪；"
                "当日核心信号：当前维持入场信号，复合策略信号稳定在3/8；复合策略信号：入场（3/8）；"
                "多头：单均线、双均线、通道过滤；空头：阶距、隔夜共振；日度主线：空头子信号共振提示短期需保持谨慎；"
                "周度全景复盘：近一周复合策略信号由6/8逐步降至3/8；近一周整体信号：偏谨慎入场；"
                "周度主线：多头动能收敛；授权剩余时间：——；业务对接：联系人：辛仑豪/刘静烨，13681991121；数据来源：Wind"
            )

    result = generate_rate_text("rate source", adapter=FakeAdapter())

    assert result.success is True
    assert "📊 当日核心信号\n复合策略信号：入场（3/8）" in result.text
    assert "多头：单均线、双均线、通道过滤；空头：阶距、隔夜共振" in result.text
    assert "📊 周度全景复盘\n近一周整体信号：偏谨慎入场" in result.text
    assert result.generated_text == result.text


def test_ai_generation_normalizes_multiline_inline_rate_text_for_renderer(investment_env):
    from business.investment.ai_generation import AIGenerationRequest, generate_rate_text

    class FakeAdapter:
        def generate(self, request: AIGenerationRequest) -> str:
            return """标的：T主力合约
最新收盘：108.970元
行情日期：2026-05-25
分析模型：交易性择时日度信号跟踪
当日核心信号：复合策略信号入场（3/8），模型整体偏多格局未变
复合策略信号：入场（3/8）
多头：单均线、双均线、通道过滤；空头：阶距、隔夜共振
日度主线：继续跟随量化模型
周度全景复盘：近一周复合策略信号回落至3/8
近一周整体信号：入场为主
周度主线：多头主线未变
授权剩余时间：——
业务对接：幸仁豪 / 刘静怡 / 13681991121"""

    result = generate_rate_text("rate source", adapter=FakeAdapter())

    assert result.success is True
    assert "📊 当日核心信号\n复合策略信号：入场（3/8）" in result.text
    assert "▪️复合策略信号入场（3/8），模型整体偏多格局未变" in result.text
    assert "📊 周度全景复盘\n近一周整体信号：入场为主" in result.text


def test_ai_generation_failures_and_health_check_sanitize_model_config(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.ai_generation import AIGenerationRequest, generate_rate_text
    from business.investment.constants import ErrorCode, Status
    from business.investment.health import run_health_checks

    api_key = "sk-ai-failure-1234567890"
    monkeypatch.setattr(config_service, "conf", lambda: {"bot_type": "custom", "custom_api_key": api_key})

    class FailingAdapter:
        def generate(self, request: AIGenerationRequest) -> str:
            raise RuntimeError(f"upstream rejected key {request.api_key}")

    result = generate_rate_text("rate source", adapter=FailingAdapter())

    assert result.success is False
    assert result.status == Status.FAILED
    assert result.error_code == ErrorCode.SYSTEM_ERROR
    assert "upstream rejected" in result.failure_reason
    assert result.failure_reason == result.detail
    assert api_key not in result.detail
    assert api_key not in str(result.model_params)

    model_health = next(item for item in run_health_checks() if item.name == "model_config")
    assert model_health.ok is False
    assert "model.name missing" in model_health.detail
    assert "model.api_base missing" in model_health.detail
    assert "model.api_key" not in model_health.detail or api_key not in model_health.detail


def test_model_health_check_accepts_global_config_fallback(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.health import run_health_checks

    api_key = "sk-global-fallback-1234567890"
    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "fallback-model",
            "custom_api_base": "https://fallback.example/v1",
            "custom_api_key": api_key,
        },
    )

    model_health = next(item for item in run_health_checks() if item.name == "model_config")

    assert model_health.ok is True
    assert api_key not in model_health.detail


def test_ai_generation_configured_adapter_uses_existing_model_http_client(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.ai_generation import ExistingModelAdapter, generate_rate_text
    from business.investment.config_service import save_configs

    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "configured-model",
            "custom_api_base": "https://configured.example/v1",
            "custom_api_key": "sk-http-client-1234567890",
            "temperature": 0.33,
        },
    )
    save_configs({"prompt.rate": "Configured rate prompt"}, operator_role="admin")

    class FakeClient:
        def __init__(self):
            self.calls = []

        def chat_completions(self, **kwargs):
            self.calls.append(kwargs)
            return {"choices": [{"message": {"content": "standard rate text"}}]}

    client = FakeClient()

    result = generate_rate_text("rate source", adapter=ExistingModelAdapter(client=client))

    assert result.success is True
    assert result.generated_text == "standard rate text"
    assert client.calls == [
        {
            "api_key": "sk-http-client-1234567890",
            "api_base": "https://configured.example/v1",
            "model": "configured-model",
            "messages": [
                {"role": "system", "content": "Configured rate prompt"},
                {"role": "user", "content": "rate source"},
            ],
            "temperature": 0.33,
            "stream": False,
        }
    ]


def test_technical_analysis_failure_records_sanitized_backend_detail(investment_env, monkeypatch):
    from business.investment import config_service
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.technical_analysis import TechnicalAnalysisResult
    from business.investment.user_service import create_user

    api_key = "sk-ta-secret-1234567890"
    monkeypatch.setattr(config_service, "conf", lambda: {"custom_api_key": api_key})
    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])

    reply = handle_text_message(
        "ok",
        "300502.SZ 技术分析",
        technical_analysis_handler=lambda _openid, _raw_input, _target: TechnicalAnalysisResult(
            False,
            error_code=ErrorCode.TECHNICAL_ANALYSIS_FAILED,
            detail=f"skill crashed while using {api_key}",
        ),
    )

    record = list_request_records(limit=1)[0]
    assert reply.success is False
    assert "skill crashed" in reply.reply_text
    assert api_key not in reply.reply_text
    assert record.error_code == ErrorCode.TECHNICAL_ANALYSIS_FAILED
    assert "skill crashed" in record.error_message
    assert api_key not in record.error_message


def test_technical_analysis_uses_skill_cli_symbol_and_saves_all_outputs(investment_env, tmp_path, monkeypatch):
    from business.investment import technical_analysis
    from business.investment.technical_analysis import TechnicalAnalysisRequest, run_technical_analysis

    request = TechnicalAnalysisRequest(openid="ok", raw_input="300502.SZ 技术分析", target_text="300502.SZ")
    assert request.openid == "ok"
    assert request.raw_input == "300502.SZ 技术分析"
    assert request.target_text == "300502.SZ"

    calls = []
    report = tmp_path / "300502_技术分析报告_2026-05-25.md"
    chart = tmp_path / "300502_TA_2026-05-25.png"
    report.write_text("# 技术分析报告\n\n核心观点", encoding="utf-8")
    chart.write_bytes(b"chart")

    def fake_skill(symbol, output_dir):
        calls.append(("skill", symbol, output_dir.name))
        return report, chart

    def fake_ai(report_text):
        calls.append(("ai", report_text))
        return SimpleNamespace(success=True, text="signal card standard text")

    def fake_render(standard_text, output_path):
        calls.append(("render", standard_text, Path(output_path).name))
        Path(output_path).write_bytes(b"card")
        return SimpleNamespace(success=True, image_path=str(output_path), detail="")

    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(technical_analysis, "generate_technical_analysis_text", fake_ai)
    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", fake_render)

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    result = run_technical_analysis("ok", "300502.SZ 技术分析")

    assert result.success is True
    assert calls[:2] == [
        ("skill", "300502", "300502_SZ"),
        ("ai", "# 技术分析报告\n\n核心观点"),
    ]
    assert calls[2][0:2] == ("render", "signal card standard text")
    assert calls[2][2].startswith("300502_SZ_signal_card_2026-05-29_")
    assert calls[2][2].endswith(".png")
    assert result.report_path == str(report)
    assert result.main_chart_path == str(chart)
    assert Path(result.signal_card_path).name == calls[2][2]
    assert result.output_files == [result.signal_card_path, str(chart), str(report)]
    assert result.market_date == "2026-05-29"
    assert "2026-05-29" in result.cache_key


def test_technical_analysis_success_records_customer_target_versions_and_artifact_roles(investment_env, tmp_path):
    from business.investment.cache_service import find_cache_entry_by_key
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.storage import get_storage_dirs
    from business.investment.technical_analysis import TechnicalAnalysisResult
    from business.investment.user_service import create_user

    card = tmp_path / "signal.png"
    chart = tmp_path / "chart.png"
    report = tmp_path / "report.md"
    card.write_bytes(b"card")
    chart.write_bytes(b"chart")
    report.write_text("report", encoding="utf-8")
    create_user("ok", name="Alice", institution="Inst A", enabled=True, allowed_services=[ServiceType.ALL])

    reply = handle_text_message(
        "ok",
        "新易盛 技术分析",
        technical_analysis_handler=lambda _openid, _raw_input, _target: TechnicalAnalysisResult(
            True,
            signal_card_path=str(card),
            main_chart_path=str(chart),
            report_path=str(report),
            output_files=[str(card), str(chart), str(report)],
            normalized_target="300502.SZ",
            stock_code="300502.SZ",
            stock_name="新易盛",
            market_date="2026-05-25",
            program_version="sha256:program12345678",
            ta_version="sha256:ta123456789012",
            renderer_version="sha256:renderer123456",
            template_version="sha256:template123456",
            version_fingerprint="sha256:combined123456",
            cache_key="technical_analysis:300502.SZ:2026-05-25:pytest",
        ),
    )

    record = list_request_records(limit=1)[0]
    cache_entry = find_cache_entry_by_key("technical_analysis:300502.SZ:2026-05-25:pytest")
    archive_root = get_storage_dirs()["generated"] / "archive"
    assert reply.success is True
    assert record.customer_name == "Alice"
    assert record.institution == "Inst A"
    assert record.normalized_target == "300502.SZ"
    assert record.stock_code == "300502.SZ"
    assert record.stock_name == "新易盛"
    assert record.market_date == "2026-05-25"
    assert record.program_version.startswith("sha256:")
    assert record.ta_version.startswith("sha256:")
    assert record.renderer_version.startswith("sha256:")
    assert record.template_version.startswith("sha256:")

    with connect() as conn:
        rows = conn.execute(
            text(
                "select file_path, artifact_role, file_size, file_hash, version_tag "
                "from investment_output_files where owner_id = :owner_id order by id"
            ),
            {"owner_id": record.request_id},
        ).mappings().all()

    artifact_paths = [row["file_path"] for row in rows]
    assert [(Path(row["file_path"]).is_file(), row["artifact_role"]) for row in rows] == [
        (True, "signal_card"),
        (True, "main_chart"),
        (True, "markdown_report"),
    ]
    assert all(Path(path).resolve().is_relative_to(archive_root.resolve()) for path in artifact_paths)
    assert artifact_paths != [str(card), str(chart), str(report)]
    assert record.output_files == artifact_paths
    assert cache_entry is not None
    assert cache_entry.output_files == artifact_paths
    assert reply.output_files == artifact_paths[:2]
    assert all(row["file_size"] > 0 for row in rows)
    assert all(str(row["file_hash"]).startswith("sha256:") for row in rows)
    assert [row["version_tag"] for row in rows] == [
        "sha256:renderer123456",
        "sha256:ta123456789012",
        "sha256:ta123456789012",
    ]


def _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-25"):
    from business.investment import technical_analysis

    calls = []

    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v1", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    def fake_skill(symbol, output_dir):
        calls.append(("skill", symbol, len(calls)))
        report = tmp_path / f"{symbol}_技术分析报告_{generated_market_date}_{len(calls)}.md"
        chart = tmp_path / f"{symbol}_TA_{generated_market_date}_{len(calls)}.png"
        report.write_text(f"报告日期：{generated_market_date}\n核心观点", encoding="utf-8")
        chart.write_bytes(b"chart")
        return report, chart

    def fake_ai(report_text):
        calls.append(("ai", report_text))
        return SimpleNamespace(success=True, text=f"日期：{generated_market_date}\n信号卡标准文本")

    def fake_render(_standard_text, output_path):
        calls.append(("render", Path(output_path).name))
        Path(output_path).write_bytes(b"card")
        return SimpleNamespace(success=True, image_path=str(output_path), detail="")

    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(technical_analysis, "generate_technical_analysis_text", fake_ai)
    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", fake_render)
    return calls


def _write_legacy_technical_analysis_cache(
    tmp_path,
    *,
    normalized_target,
    market_date,
    version_fingerprint,
    program_version="sha256:old-program",
    ta_version="sha256:ta-v1",
    renderer_version="sha256:renderer-v1",
    template_version="sha256:template-v1",
):
    from business.investment.cache_service import build_cache_key, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, succeed_request_record

    cached_card = tmp_path / f"cached-card-{market_date}.png"
    cached_chart = tmp_path / f"cached-chart-{market_date}.png"
    cached_report = tmp_path / f"cached-report-{market_date}.md"
    cached_card.write_bytes(f"cached-card-{market_date}".encode("utf-8"))
    cached_chart.write_bytes(f"cached-chart-{market_date}".encode("utf-8"))
    cached_report.write_text(f"cached report {market_date}", encoding="utf-8")
    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, normalized_target, market_date, version_fingerprint)
    request_id = create_request_record(
        "ok",
        f"{normalized_target} 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        normalized_target=normalized_target,
    )
    succeed_request_record(
        request_id,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
        elapsed_ms=1,
        normalized_target=normalized_target,
        stock_code=normalized_target,
        stock_name=normalized_target,
        market_date=market_date,
        cache_key=cache_key,
        program_version=program_version,
        ta_version=ta_version,
        renderer_version=renderer_version,
        template_version=template_version,
    )
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target=normalized_target,
        market_date=market_date,
        version_fingerprint=version_fingerprint,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
        artifact_owner_id=request_id,
    )
    return cache_key, [str(cached_card), str(cached_chart), str(cached_report)]


def test_technical_analysis_reuses_cached_outputs_without_explicit_date_when_resolver_confirms_market_date(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]

    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False
    assert records[0].cache_key == records[1].cache_key
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert cache_entry.hit_count == 1
    assert cache_entry.market_date == "2026-05-25"
    assert cache_entry.artifact_owner_id == records[1].request_id


def test_technical_analysis_reuses_today_cache_before_close_cutoff(
    investment_env, tmp_path, monkeypatch
):
    from zoneinfo import ZoneInfo

    from business.investment import cache_policy, technical_analysis
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    monkeypatch.setattr(
        cache_policy,
        "beijing_now",
        lambda: datetime(2026, 5, 29, 15, 29, tzinfo=ZoneInfo("Asia/Shanghai")),
    )

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False


def test_technical_analysis_invalidates_today_intraday_cache_after_close_and_reruns(
    investment_env, tmp_path, monkeypatch
):
    from zoneinfo import ZoneInfo

    from business.investment import cache_policy, technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.schema import investment_cache_entries
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    monkeypatch.setattr(
        cache_policy,
        "beijing_now",
        lambda: datetime(2026, 5, 29, 15, 31, tzinfo=ZoneInfo("Asia/Shanghai")),
    )

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    assert first.success is True
    first_record = list_request_records(limit=1)[0]
    with connect() as conn:
        conn.execute(
            investment_cache_entries.update()
            .where(investment_cache_entries.c.cache_key == first_record.cache_key)
            .values(updated_at="2026-05-29T07:00:00+00:00")
        )

    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert second.success is True
    assert second.output_files != first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render", "skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].cache_hit is False
    assert records[1].cache_hit is False
    assert records[0].cache_key == records[1].cache_key
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert cache_entry.hit_count == 0
    assert cache_entry.status == "active"


def test_technical_analysis_keeps_previous_trading_day_cache_after_close(
    investment_env, tmp_path, monkeypatch
):
    from zoneinfo import ZoneInfo

    from business.investment import cache_policy, technical_analysis
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-28")
    monkeypatch.setattr(
        cache_policy,
        "beijing_now",
        lambda: datetime(2026, 5, 29, 15, 31, tzinfo=ZoneInfo("Asia/Shanghai")),
    )

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-28", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False
    assert records[0].market_date == "2026-05-28"


def test_technical_analysis_missing_cache_file_invalidates_and_reruns(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    assert first.success is True
    Path(first.output_files[0]).unlink()

    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert second.success is True
    assert second.output_files != first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render", "skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].cache_hit is False
    assert records[1].cache_hit is False
    assert records[0].cache_key == records[1].cache_key
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert cache_entry.status == "active"
    assert cache_entry.artifact_owner_id == records[0].request_id


def test_technical_analysis_cache_write_failure_does_not_leave_active_cache(investment_env, tmp_path, monkeypatch):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType, Status
    from business.investment.db import connect
    from business.investment.records import get_request_record, list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    def fail_router_cache_write(**_kwargs):
        raise RuntimeError("cache write failed after generation")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)
    monkeypatch.setattr(cache_service, "write_cache_entry", fail_router_cache_write)

    reply = handle_text_message("ok", "300502.SZ 技术分析")

    assert reply.success is False
    assert reply.error_code is not None
    assert list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS) == []
    request_record = get_request_record(list_request_records(limit=1)[0].request_id)
    assert request_record.status == Status.FAILED
    assert request_record.output_files == []
    assert request_record.cache_key == ""
    assert request_record.market_date == ""
    assert request_record.cache_hit is False
    with connect() as conn:
        artifact_count = conn.execute(
            text("select count(*) from investment_output_files where owner_id = :request_id"),
            {"request_id": request_record.request_id},
        ).scalar_one()
    assert artifact_count == 0


def test_router_default_technical_analysis_resolver_once_reuses_preview_resolution(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)
    resolver_calls = []

    class CountingResolver:
        def resolve(self, symbol, requested_market_date=""):
            resolver_calls.append((symbol, requested_market_date))
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", CountingResolver, raising=False)

    reply = handle_text_message("ok", "300502.SZ 技术分析")

    assert reply.success is True
    assert resolver_calls == [("300502.SZ", "")]
    record = list_request_records(limit=1)[0]
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert record.cache_key == cache_entry.cache_key
    assert record.market_date == "2026-05-25"


def test_technical_analysis_lock_cache_key_uses_resolver_date_over_generated_date(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-26")

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    reply = handle_text_message("ok", "300502.SZ 技术分析")

    assert reply.success is True
    record = list_request_records(limit=1)[0]
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert record.market_date == "2026-05-25"
    assert "2026-05-25" in record.cache_key
    assert record.cache_key == cache_entry.cache_key
    assert cache_entry.market_date == "2026-05-25"


def test_technical_analysis_explicit_market_date_keeps_specified_cache_date(investment_env, tmp_path, monkeypatch):
    from business.investment import technical_analysis
    from business.investment.market_date_resolver import MarketDateResolver
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.constants import ServiceType
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)
    monkeypatch.setattr(
        MarketDateResolver,
        "_latest_market_date",
        lambda _self, _symbol: SimpleNamespace(market_date="2026-05-26", known=True, source="latest"),
    )
    monkeypatch.setattr(technical_analysis, "MarketDateResolver", MarketDateResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 2026-05-25 技术分析")
    second = handle_text_message("ok", "300502.SZ 2026-05-25 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]

    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False
    assert records[0].cache_key == records[1].cache_key
    assert records[0].market_date == "2026-05-25"
    assert records[1].market_date == "2026-05-25"


def test_technical_analysis_explicit_market_date_overrides_generated_output_date(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-26")

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == "2026-05-25"
            return SimpleNamespace(market_date="2026-05-25", known=True, source="explicit")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "300502.SZ 2026-05-25 技术分析")
    second = handle_text_message("ok", "300502.SZ 2026-05-25 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]

    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False
    assert records[0].market_date == "2026-05-25"
    assert records[1].market_date == "2026-05-25"
    assert records[0].cache_key == records[1].cache_key
    assert "2026-05-25" in records[0].cache_key
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert cache_entry.market_date == "2026-05-25"
    assert "2026-05-25" in cache_entry.cache_key


def test_technical_analysis_unknown_market_date_reuses_recent_latest_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import build_cache_key, list_cache_entries, version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.stock_resolver import refresh_stock_symbols
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)
    combined_version = version_fingerprint("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")
    cached_card = tmp_path / "cached-card.png"
    cached_chart = tmp_path / "cached-chart.png"
    cached_report = tmp_path / "cached-report.md"
    cached_card.write_bytes(b"cached-card")
    cached_chart.write_bytes(b"cached-chart")
    cached_report.write_text("cached report", encoding="utf-8")
    write_cache_entry(
        cache_key=build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "002354.SZ", "2026-05-29", combined_version),
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=combined_version,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    reply = handle_text_message("ok", "天娱数科 技术分析")

    assert reply.success is True
    assert reply.output_files != [str(cached_card), str(cached_chart)]
    assert [Path(path).read_bytes() for path in reply.output_files] == [b"cached-card", b"cached-chart"]
    assert calls == []
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is True
    assert record.cache_key == build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "002354.SZ", "2026-05-29", combined_version)
    assert record.market_date == "2026-05-29"
    assert record.normalized_target == "002354.SZ"
    cache_entries = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS, limit=10)
    cached_entry = next(
        entry
        for entry in cache_entries
        if entry.cache_key == build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "002354.SZ", "2026-05-29", combined_version)
    )
    assert cached_entry.hit_count == 1


def test_technical_analysis_unknown_cache_context_reuses_latest_cache_when_only_program_wrapper_changes(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import build_cache_key, version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    old_cache_version = version_fingerprint("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")
    cached_card = tmp_path / "cached-card.png"
    cached_chart = tmp_path / "cached-chart.png"
    cached_report = tmp_path / "cached-report.md"
    cached_card.write_bytes(b"cached-card")
    cached_chart.write_bytes(b"cached-chart")
    cached_report.write_text("cached report", encoding="utf-8")
    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "002354.SZ", "2026-05-29", old_cache_version)
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=old_cache_version,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v2", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    context = technical_analysis.prepare_technical_analysis_cache_context("天娱数科 技术分析")

    assert context.program_version == "sha256:program-v2"
    assert context.version_fingerprint == old_cache_version
    assert context.cache_key == cache_key
    assert context.market_date == "2026-05-29"
    assert context.cache_lookup_version_fingerprint == old_cache_version


def test_technical_analysis_unknown_market_date_does_not_reuse_legacy_program_version_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import build_cache_key, version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, succeed_request_record
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    cached_card = tmp_path / "cached-card.png"
    cached_chart = tmp_path / "cached-chart.png"
    cached_report = tmp_path / "cached-report.md"
    cached_card.write_bytes(b"cached-card")
    cached_chart.write_bytes(b"cached-chart")
    cached_report.write_text("cached report", encoding="utf-8")
    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "002354.SZ", "2026-05-29", legacy_cache_version)
    request_id = create_request_record(
        "ok",
        "天娱数科 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
    )
    succeed_request_record(
        request_id,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
        elapsed_ms=1,
        normalized_target="002354.SZ",
        stock_code="002354.SZ",
        stock_name="天娱数科",
        market_date="2026-05-29",
        cache_key=cache_key,
        program_version="sha256:old-program",
        ta_version="sha256:ta-v1",
        renderer_version="sha256:renderer-v1",
        template_version="sha256:template-v1",
    )
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
        artifact_owner_id=request_id,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    context = technical_analysis.prepare_technical_analysis_cache_context("天娱数科 技术分析")

    assert context.program_version == "sha256:new-program"
    assert context.version_fingerprint == version_fingerprint("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")
    assert context.cache_key == ""
    assert context.market_date == ""
    assert context.cache_lookup_version_fingerprint == context.version_fingerprint


def test_direct_technical_analysis_unknown_market_date_does_not_reuse_legacy_latest_cache_without_context(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import version_fingerprint
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    cache_key, cached_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    result = technical_analysis.run_technical_analysis("ok", "天娱数科 技术分析")

    assert result.success is True
    assert result.cache_hit is False
    assert result.cache_key != cache_key
    assert result.output_files != cached_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]


def test_technical_analysis_unknown_market_date_does_not_search_legacy_latest_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import build_cache_key, version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    compatible_key, _compatible_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-27",
        version_fingerprint=legacy_cache_version,
    )
    _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-28",
        version_fingerprint=legacy_cache_version,
        ta_version="sha256:ta-v2",
    )
    missing_owner_files = []
    for name in ["missing-owner-card.png", "missing-owner-chart.png", "missing-owner-report.md"]:
        path = tmp_path / name
        path.write_text("missing owner", encoding="utf-8")
        missing_owner_files.append(str(path))
    write_cache_entry(
        cache_key=build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "002354.SZ", "2026-05-29", legacy_cache_version),
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
        output_files=missing_owner_files,
        artifact_owner_id="",
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    context = technical_analysis.prepare_technical_analysis_cache_context("天娱数科 技术分析")

    assert context.cache_key != compatible_key
    assert context.cache_key == ""
    assert context.market_date == ""
    assert context.cache_lookup_version_fingerprint == version_fingerprint(
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )


def test_router_context_uses_specific_compatible_legacy_cache_key_when_plain_lookup_would_find_incompatible_row(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.records import create_request_record, list_request_records, succeed_request_record
    from business.investment.router import handle_text_message
    from business.investment.stock_resolver import refresh_stock_symbols
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    compatible_key, compatible_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
    )
    incompatible_card = tmp_path / "incompatible-card.png"
    incompatible_chart = tmp_path / "incompatible-chart.png"
    incompatible_report = tmp_path / "incompatible-report.md"
    incompatible_card.write_bytes(b"incompatible-card")
    incompatible_chart.write_bytes(b"incompatible-chart")
    incompatible_report.write_text("incompatible report", encoding="utf-8")
    incompatible_key = "manual-incompatible-cache-key"
    incompatible_owner_id = create_request_record(
        "ok",
        "天娱数科 技术分析",
        ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
    )
    succeed_request_record(
        incompatible_owner_id,
        output_files=[str(incompatible_card), str(incompatible_chart), str(incompatible_report)],
        elapsed_ms=1,
        normalized_target="002354.SZ",
        stock_code="002354.SZ",
        stock_name="天娱数科",
        market_date="2026-05-29",
        cache_key=incompatible_key,
        program_version="sha256:old-program",
        ta_version="sha256:ta-v2",
        renderer_version="sha256:renderer-v1",
        template_version="sha256:template-v1",
    )
    write_cache_entry(
        cache_key=incompatible_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
        output_files=[str(incompatible_card), str(incompatible_chart), str(incompatible_report)],
        artifact_owner_id=incompatible_owner_id,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class KnownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="latest")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", KnownResolver, raising=False)

    reply = handle_text_message("ok", "天娱数科 技术分析")

    assert reply.success is True
    assert reply.output_files != compatible_files[:2]
    assert [Path(path).read_bytes() for path in reply.output_files] == [
        Path(compatible_files[0]).read_bytes(),
        Path(compatible_files[1]).read_bytes(),
    ]
    assert calls == []
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is True
    assert record.cache_key == compatible_key
    assert record.cache_key != incompatible_key


def test_technical_analysis_invalidates_compatible_today_intraday_cache_after_close_and_reruns(
    investment_env, tmp_path, monkeypatch
):
    from zoneinfo import ZoneInfo

    from business.investment import cache_policy, cache_service, technical_analysis
    from business.investment.cache_service import list_cache_entries, version_fingerprint
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.schema import investment_cache_entries
    from business.investment.stock_resolver import refresh_stock_symbols
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    compatible_key, compatible_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        cache_policy,
        "beijing_now",
        lambda: datetime(2026, 5, 29, 15, 31, tzinfo=ZoneInfo("Asia/Shanghai")),
    )
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )
    with connect() as conn:
        conn.execute(
            investment_cache_entries.update()
            .where(investment_cache_entries.c.cache_key == compatible_key)
            .values(updated_at="2026-05-29T07:00:00+00:00")
        )

    class KnownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="latest")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", KnownResolver, raising=False)

    reply = handle_text_message("ok", "天娱数科 技术分析")

    assert reply.success is True
    assert reply.output_files != compatible_files[:2]
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is False
    assert record.cache_key != compatible_key
    entries = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS, include_invalidated=True)
    assert {entry.status for entry in entries if entry.cache_key == compatible_key} == {"invalidated"}


def test_technical_analysis_context_cache_key_misses_when_owner_becomes_incompatible(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import list_cache_entries, version_fingerprint
    from business.investment.constants import ServiceType
    from business.investment.db import connect
    from business.investment.schema import investment_request_records
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    compatible_key, compatible_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class KnownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="latest")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", KnownResolver, raising=False)
    context = technical_analysis.prepare_technical_analysis_cache_context("天娱数科 技术分析")
    owner_id = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0].artifact_owner_id
    with connect() as conn:
        conn.execute(
            investment_request_records.update()
            .where(investment_request_records.c.request_id == owner_id)
            .values(ta_version="sha256:ta-v2")
        )

    result = technical_analysis.run_technical_analysis("ok", "天娱数科 技术分析", cache_context=context)

    assert context.cache_key == compatible_key
    assert result.success is True
    assert result.cache_hit is False
    assert result.cache_key == compatible_key
    assert result.output_files != compatible_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]


def test_technical_analysis_legacy_latest_misses_when_no_compatible_owner_row_in_window(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import version_fingerprint
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-20",
        version_fingerprint=legacy_cache_version,
    )
    incompatible_key, _incompatible_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
        ta_version="sha256:ta-v2",
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    context = technical_analysis.prepare_technical_analysis_cache_context("天娱数科 技术分析")

    assert context.cache_key != incompatible_key
    assert context.market_date == ""
    assert context.cache_key == ""
    assert context.cache_lookup_version_fingerprint == version_fingerprint(
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )


def test_technical_analysis_known_market_date_reuses_legacy_program_version_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import version_fingerprint
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.stock_resolver import refresh_stock_symbols
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-29")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    cache_key, cached_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="002354.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class KnownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-29", known=True, source="latest")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", KnownResolver, raising=False)

    reply = handle_text_message("ok", "天娱数科 技术分析")

    assert reply.success is True
    assert reply.output_files != cached_files[:2]
    assert [Path(path).read_bytes() for path in reply.output_files] == [
        Path(cached_files[0]).read_bytes(),
        Path(cached_files[1]).read_bytes(),
    ]
    assert calls == []
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is True
    assert record.cache_key == cache_key
    assert record.market_date == "2026-05-29"


def test_technical_analysis_explicit_market_date_reuses_legacy_program_version_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import version_fingerprint
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-28")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    cache_key, cached_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="300502.SZ",
        market_date="2026-05-28",
        version_fingerprint=legacy_cache_version,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class ExplicitResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == "2026-05-28"
            return SimpleNamespace(market_date="2026-05-28", known=True, source="explicit")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", ExplicitResolver, raising=False)

    reply = handle_text_message("ok", "300502.SZ 2026-05-28 技术分析")

    assert reply.success is True
    assert reply.output_files != cached_files[:2]
    assert [Path(path).read_bytes() for path in reply.output_files] == [
        Path(cached_files[0]).read_bytes(),
        Path(cached_files[1]).read_bytes(),
    ]
    assert calls == []
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is True
    assert record.cache_key == cache_key
    assert record.market_date == "2026-05-28"


def test_technical_analysis_explicit_market_date_does_not_fallback_to_other_legacy_date(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import version_fingerprint
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-28")
    legacy_cache_version = version_fingerprint(
        "sha256:old-program",
        "sha256:ta-v1",
        "sha256:renderer-v1",
        "sha256:template-v1",
    )
    _cache_key, cached_files = _write_legacy_technical_analysis_cache(
        tmp_path,
        normalized_target="300502.SZ",
        market_date="2026-05-29",
        version_fingerprint=legacy_cache_version,
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)
    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:new-program", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class ExplicitResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == "2026-05-28"
            return SimpleNamespace(market_date="2026-05-28", known=True, source="explicit")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", ExplicitResolver, raising=False)

    reply = handle_text_message("ok", "300502.SZ 2026-05-28 技术分析")

    assert reply.success is True
    assert reply.output_files != cached_files[:2]
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is False
    assert record.market_date == "2026-05-28"


def test_technical_analysis_cache_key_version_changes_when_output_versions_change(
    investment_env, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import version_fingerprint
    from business.investment.stock_resolver import refresh_stock_symbols

    refresh_stock_symbols([{"code": "002354.SZ", "name": "天娱数科", "market": "SZ", "source": "pytest"}], source="pytest")

    class UnknownResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "002354.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    def context_version(ta_version, renderer_version, template_version):
        monkeypatch.setattr(
            technical_analysis,
            "_versions",
            lambda: ("sha256:program-v1", ta_version, renderer_version, template_version),
        )
        return technical_analysis.prepare_technical_analysis_cache_context("天娱数科 技术分析").version_fingerprint

    base = context_version("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")

    assert base == version_fingerprint("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")
    assert context_version("sha256:ta-v2", "sha256:renderer-v1", "sha256:template-v1") != base
    assert context_version("sha256:ta-v1", "sha256:renderer-v2", "sha256:template-v1") != base
    assert context_version("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v2") != base


def test_technical_analysis_unknown_market_date_does_not_reuse_cache_outside_fallback_window(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import build_cache_key, version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)
    combined_version = version_fingerprint("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")
    old_card = tmp_path / "old-card.png"
    old_chart = tmp_path / "old-chart.png"
    old_report = tmp_path / "old-report.md"
    old_card.write_bytes(b"old-card")
    old_chart.write_bytes(b"old-chart")
    old_report.write_text("old report", encoding="utf-8")
    write_cache_entry(
        cache_key=build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-20", combined_version),
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-20",
        version_fingerprint=combined_version,
        output_files=[str(old_card), str(old_chart), str(old_report)],
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)

    class UnknownResolver:
        def resolve(self, _symbol, requested_market_date=""):
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)

    reply = handle_text_message("ok", "300502.SZ 技术分析")

    assert reply.success is True
    assert reply.output_files != [str(old_card), str(old_chart)]
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is False
    assert record.market_date == "2026-05-25"


def test_technical_analysis_explicit_market_date_does_not_fallback_to_latest_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service, technical_analysis
    from business.investment.cache_service import build_cache_key, version_fingerprint, write_cache_entry
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path, generated_market_date="2026-05-28")
    combined_version = version_fingerprint("sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1")
    cached_card = tmp_path / "cached-card.png"
    cached_chart = tmp_path / "cached-chart.png"
    cached_report = tmp_path / "cached-report.md"
    cached_card.write_bytes(b"cached-card")
    cached_chart.write_bytes(b"cached-chart")
    cached_report.write_text("cached report", encoding="utf-8")
    write_cache_entry(
        cache_key=build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-29", combined_version),
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-29",
        version_fingerprint=combined_version,
        output_files=[str(cached_card), str(cached_chart), str(cached_report)],
    )
    monkeypatch.setattr(cache_service, "_today", lambda: datetime(2026, 5, 29).date(), raising=False)

    class ExplicitResolver:
        def resolve(self, _symbol, requested_market_date=""):
            assert requested_market_date == "2026-05-28"
            return SimpleNamespace(market_date="2026-05-28", known=True, source="explicit")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", ExplicitResolver, raising=False)

    reply = handle_text_message("ok", "300502.SZ 2026-05-28 技术分析")

    assert reply.success is True
    assert reply.output_files != [str(cached_card), str(cached_chart)]
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    record = list_request_records(limit=1)[0]
    assert record.cache_hit is False
    assert record.market_date == "2026-05-28"


def test_technical_analysis_stock_name_reuses_same_standard_code_cache(investment_env, tmp_path, monkeypatch):
    from business.investment import technical_analysis
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.stock_resolver import refresh_stock_symbols
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    refresh_stock_symbols([{"code": "300502.SZ", "name": "新易盛", "market": "SZ", "source": "pytest"}], source="pytest")
    calls = _patch_fake_technical_analysis_pipeline(monkeypatch, tmp_path)

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)

    first = handle_text_message("ok", "新易盛 技术分析")
    second = handle_text_message("ok", "新易盛 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False
    assert records[0].normalized_target == "300502.SZ"
    assert records[1].normalized_target == "300502.SZ"


def test_technical_analysis_different_market_date_misses_cache(investment_env, tmp_path, monkeypatch):
    from business.investment import technical_analysis
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    calls = []
    current_market_date = ""

    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v1", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    def fake_skill(symbol, output_dir):
        calls.append(("skill", symbol, current_market_date))
        report = tmp_path / f"{symbol}_技术分析报告_{current_market_date}.md"
        chart = tmp_path / f"{symbol}_TA_{current_market_date}.png"
        report.write_text(f"报告日期：{current_market_date}\n核心观点", encoding="utf-8")
        chart.write_bytes(f"chart-{current_market_date}".encode("utf-8"))
        return report, chart

    def fake_ai(report_text):
        calls.append(("ai", current_market_date))
        return SimpleNamespace(success=True, text=f"日期：{current_market_date}\n信号卡标准文本")

    def fake_render(_standard_text, output_path):
        calls.append(("render", current_market_date))
        Path(output_path).write_bytes(f"card-{current_market_date}".encode("utf-8"))
        return SimpleNamespace(success=True, image_path=str(output_path), detail="")

    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(technical_analysis, "generate_technical_analysis_text", fake_ai)
    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", fake_render)

    current_market_date = "2026-05-25"
    first = handle_text_message("ok", "300502.SZ 2026-05-25 技术分析")
    current_market_date = "2026-05-26"
    second = handle_text_message("ok", "300502.SZ 2026-05-26 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files != first.output_files
    assert [call[0] for call in calls].count("skill") == 2
    records = list_request_records(limit=2)
    assert records[0].market_date == "2026-05-26"
    assert records[1].market_date == "2026-05-25"
    assert records[0].cache_hit is False
    assert records[1].cache_hit is False

    current_market_date = "2026-05-25"
    third = handle_text_message("ok", "300502.SZ 2026-05-25 技术分析")

    assert third.success is True
    assert third.output_files == first.output_files
    assert [call[0] for call in calls].count("skill") == 2
    assert Path(third.output_files[0]).read_bytes() == b"card-2026-05-25"
    assert list_request_records(limit=1)[0].cache_hit is True


def test_technical_analysis_unknown_generated_market_date_does_not_cache_current_date(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    report = tmp_path / "report_without_date.md"
    chart = tmp_path / "chart_without_date.png"
    report.write_text("核心观点 without date", encoding="utf-8")
    chart.write_bytes(b"chart")

    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v1", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )
    calls = []

    class UnknownResolver:
        def resolve(self, _symbol, requested_market_date=""):
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    def fake_skill(symbol, _output_dir):
        calls.append(("skill", symbol))
        return report, chart

    def fake_ai(_report_text):
        calls.append(("ai", len(calls)))
        return SimpleNamespace(success=True, text="信号卡标准文本 without date")

    def fake_render(_standard_text, output_path):
        calls.append(("render", Path(output_path).name))
        Path(output_path).write_bytes(b"card")
        return SimpleNamespace(success=True, image_path=str(output_path), detail="")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)
    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(technical_analysis, "generate_technical_analysis_text", fake_ai)
    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", fake_render)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert first.success is True
    assert second.success is True
    assert [call[0] for call in calls] == ["skill", "ai", "render", "skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].market_date == ""
    assert records[1].market_date == ""
    assert records[0].cache_key == ""
    assert records[1].cache_key == ""
    assert records[0].cache_hit is False
    assert records[1].cache_hit is False
    assert "market_date unknown" in records[0].status_warning
    assert list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS) == []


def test_technical_analysis_invalid_generated_market_date_does_not_cache(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    invalid_market_date = "2026-02-31"
    report = tmp_path / f"300502_技术分析报告_{invalid_market_date}.md"
    chart = tmp_path / f"300502_TA_{invalid_market_date}.png"
    report.write_text(f"报告日期：{invalid_market_date}\n核心观点", encoding="utf-8")
    chart.write_bytes(b"chart")

    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v1", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )

    class UnknownResolver:
        def resolve(self, _symbol, requested_market_date=""):
            assert requested_market_date == ""
            return SimpleNamespace(market_date="", known=False, source="unknown")

    def fake_skill(_symbol, _output_dir):
        return report, chart

    def fake_ai(_report_text):
        return SimpleNamespace(success=True, text=f"日期：{invalid_market_date}\n信号卡标准文本")

    def fake_render(_standard_text, output_path):
        Path(output_path).write_bytes(b"card")
        return SimpleNamespace(success=True, image_path=str(output_path), detail="")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", UnknownResolver, raising=False)
    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(technical_analysis, "generate_technical_analysis_text", fake_ai)
    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", fake_render)

    reply = handle_text_message("ok", "300502.SZ 技术分析")

    assert reply.success is True
    record = list_request_records(limit=1)[0]
    assert record.market_date == ""
    assert record.cache_key == ""
    assert list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS) == []


def test_technical_analysis_resolver_market_date_is_used_when_generated_outputs_have_no_date(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import technical_analysis
    from business.investment.cache_service import list_cache_entries
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    report = tmp_path / "report_without_date.md"
    chart = tmp_path / "chart_without_date.png"
    report.write_text("核心观点 without date", encoding="utf-8")
    chart.write_bytes(b"chart")

    monkeypatch.setattr(
        technical_analysis,
        "_versions",
        lambda: ("sha256:program-v1", "sha256:ta-v1", "sha256:renderer-v1", "sha256:template-v1"),
    )
    calls = []

    class FakeResolver:
        def resolve(self, symbol, requested_market_date=""):
            assert symbol == "300502.SZ"
            assert requested_market_date == ""
            return SimpleNamespace(market_date="2026-05-25", known=True, source="fake")

    def fake_skill(symbol, _output_dir):
        calls.append(("skill", symbol))
        return report, chart

    def fake_ai(_report_text):
        calls.append(("ai", len(calls)))
        return SimpleNamespace(success=True, text="信号卡标准文本 without date")

    def fake_render(_standard_text, output_path):
        calls.append(("render", Path(output_path).name))
        Path(output_path).write_bytes(b"card")
        return SimpleNamespace(success=True, image_path=str(output_path), detail="")

    monkeypatch.setattr(technical_analysis, "MarketDateResolver", FakeResolver, raising=False)
    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(technical_analysis, "generate_technical_analysis_text", fake_ai)
    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", fake_render)

    first = handle_text_message("ok", "300502.SZ 技术分析")
    second = handle_text_message("ok", "300502.SZ 技术分析")

    assert first.success is True
    assert second.success is True
    assert second.output_files == first.output_files
    assert [call[0] for call in calls] == ["skill", "ai", "render"]
    records = list_request_records(limit=2)
    assert records[0].cache_hit is True
    assert records[1].cache_hit is False
    assert records[0].market_date == "2026-05-25"
    assert records[1].market_date == "2026-05-25"
    assert records[0].cache_key
    assert records[1].cache_key
    assert records[0].cache_key == records[1].cache_key
    cache_entry = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)[0]
    assert cache_entry.market_date == "2026-05-25"


def test_market_date_resolver_rejects_invalid_explicit_dates(monkeypatch):
    from business.investment.market_date_resolver import MarketDateResolver

    calls = []

    monkeypatch.setattr(
        MarketDateResolver,
        "_latest_market_date",
        lambda _self, _symbol: calls.append(_symbol)
        or SimpleNamespace(market_date="2026-05-26", known=True, source="latest"),
    )

    resolver = MarketDateResolver()

    result = resolver.resolve("300502.SZ", "2026-02-31")

    assert result.market_date == ""
    assert result.known is False
    assert result.source == "unknown"
    assert calls == []


def test_write_cache_entry_rewrites_payload_without_resetting_hit_count(investment_env):
    from business.investment.cache_service import build_cache_key, increment_cache_hit, list_cache_entries, write_cache_entry
    from business.investment.constants import ServiceType

    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-25", "v1")
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=["/tmp/old-card.png"],
        artifact_owner_id="request-old",
    )
    increment_cache_hit(cache_key)

    rewritten = write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=["/tmp/new-card.png", "/tmp/new-chart.png"],
        artifact_owner_id="request-new",
    )

    assert rewritten.hit_count == 1
    assert rewritten.output_files == ["/tmp/new-card.png", "/tmp/new-chart.png"]
    assert rewritten.artifact_owner_id == "request-new"
    assert rewritten.status == "active"
    entries = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)
    assert len(entries) == 1
    assert entries[0].hit_count == 1
    assert entries[0].output_files == ["/tmp/new-card.png", "/tmp/new-chart.png"]
    assert entries[0].artifact_owner_id == "request-new"
    assert entries[0].status == "active"


def test_beijing_now_returns_beijing_timezone_datetime():
    from business.investment.cache_policy import beijing_now

    now = beijing_now()

    assert now.tzinfo is not None
    assert now.tzname() == "CST"
    assert now.utcoffset() == timedelta(hours=8)


def test_technical_analysis_cache_policy_keeps_cache_before_close_cutoff():
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 7, 0, tzinfo=UTC),
            now=datetime(2026, 5, 31, 15, 29),
        )
        is False
    )


def test_technical_analysis_cache_policy_expires_today_cache_written_before_close_cutoff():
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 15, 29, 59),
            now=datetime(2026, 5, 31, 15, 30),
        )
        is True
    )


def test_technical_analysis_cache_policy_keeps_non_today_market_date_after_close_cutoff():
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-30",
            datetime(2026, 5, 30, 15, 0),
            now=datetime(2026, 5, 31, 15, 30),
        )
        is False
    )


def test_technical_analysis_cache_policy_keeps_today_cache_written_at_or_after_close_cutoff():
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 15, 30),
            now=datetime(2026, 5, 31, 15, 31),
        )
        is False
    )


def test_technical_analysis_cache_policy_compares_utc_and_local_times_as_beijing_time():
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            "2026-05-31T07:29:59+00:00",
            now="2026-05-31T07:30:00+00:00",
        )
        is True
    )
    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            "2026-05-31T15:30:00+08:00",
            now="2026-05-31T07:30:00+00:00",
        )
        is False
    )


def test_technical_analysis_cache_policy_uses_default_cutoff_when_config_missing(investment_env):
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 15, 29, 59),
            now=datetime(2026, 5, 31, 15, 30),
        )
        is True
    )


def test_technical_analysis_cache_policy_uses_configured_cutoff_time(investment_env):
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close
    from business.investment.config_service import save_config

    save_config("investment.technical_analysis.cache_close_invalidate_time", "14:45", operator_role="admin")

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 14, 44, 59),
            now=datetime(2026, 5, 31, 14, 45),
        )
        is True
    )


def test_technical_analysis_cache_policy_falls_back_to_default_for_invalid_cutoff(investment_env):
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close
    from business.investment.config_service import save_config

    save_config("investment.technical_analysis.cache_close_invalidate_time", "14:45:00", operator_role="admin")

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 14, 44, 59),
            now=datetime(2026, 5, 31, 14, 45),
        )
        is False
    )


def test_technical_analysis_cache_policy_requires_strict_hh_mm_cutoff(investment_env):
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close
    from business.investment.config_service import save_config

    save_config("investment.technical_analysis.cache_close_invalidate_time", "1:02", operator_role="admin")

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 1, 1, 59),
            now=datetime(2026, 5, 31, 1, 2),
        )
        is False
    )


def test_technical_analysis_cache_policy_rejects_padded_hh_mm_cutoff(investment_env):
    from business.investment.cache_policy import technical_analysis_cache_expired_after_close
    from business.investment.config_service import save_config

    save_config("investment.technical_analysis.cache_close_invalidate_time", " 14:45", operator_role="admin")

    assert (
        technical_analysis_cache_expired_after_close(
            "2026-05-31",
            datetime(2026, 5, 31, 14, 44, 59),
            now=datetime(2026, 5, 31, 14, 45),
        )
        is False
    )


def test_find_cache_entry_missing_cache_file_invalidates_active_entry(investment_env, tmp_path):
    from business.investment.cache_service import build_cache_key, find_cache_entry, list_cache_entries, write_cache_entry
    from business.investment.constants import ServiceType

    card = tmp_path / "card.png"
    chart = tmp_path / "chart.png"
    report = tmp_path / "report.md"
    card.write_bytes(b"card")
    chart.write_bytes(b"chart")
    report.write_text("report", encoding="utf-8")
    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-25", "v1")
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=[str(card), str(chart), str(report)],
        artifact_owner_id="request-1",
    )
    chart.unlink()

    entry = find_cache_entry(
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        require_files=True,
    )

    assert entry is None
    assert list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS) == []
    entries = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS, include_invalidated=True)
    assert len(entries) == 1
    assert entries[0].cache_key == cache_key
    assert entries[0].status == "invalidated"


def test_find_cache_entry_missing_cache_file_does_not_invalidate_rewritten_active_entry(
    investment_env, tmp_path, monkeypatch
):
    from business.investment import cache_service
    from business.investment.cache_service import build_cache_key, list_cache_entries, write_cache_entry
    from business.investment.constants import ServiceType

    old_card = tmp_path / "old-card.png"
    new_card = tmp_path / "new-card.png"
    new_card.write_bytes(b"new-card")
    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-25", "v1")
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=[str(old_card)],
        artifact_owner_id="request-old",
    )

    def rewrite_entry_before_missing_file_result(paths):
        assert paths == [str(old_card)]
        write_cache_entry(
            cache_key=cache_key,
            service_type=ServiceType.TECHNICAL_ANALYSIS,
            normalized_target="300502.SZ",
            market_date="2026-05-25",
            version_fingerprint="v1",
            output_files=[str(new_card)],
            artifact_owner_id="request-new",
        )
        return False

    monkeypatch.setattr(cache_service, "_files_available", rewrite_entry_before_missing_file_result)

    entry = cache_service.find_cache_entry(
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        require_files=True,
    )

    assert entry is None
    entries = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)
    assert len(entries) == 1
    assert entries[0].cache_key == cache_key
    assert entries[0].status == "active"
    assert entries[0].output_files == [str(new_card)]
    assert entries[0].artifact_owner_id == "request-new"


def test_write_cache_entry_concurrent_same_key_uses_single_active_entry(investment_env):
    from concurrent.futures import ThreadPoolExecutor

    from business.investment.cache_service import build_cache_key, list_cache_entries, write_cache_entry
    from business.investment.constants import ServiceType

    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-25", "v1")

    def write(output_file: str):
        return write_cache_entry(
            cache_key=cache_key,
            service_type=ServiceType.TECHNICAL_ANALYSIS,
            normalized_target="300502.SZ",
            market_date="2026-05-25",
            version_fingerprint="v1",
            output_files=[output_file],
            artifact_owner_id=output_file,
        )

    output_files = [f"/tmp/card-{index}.png" for index in range(12)]
    with ThreadPoolExecutor(max_workers=6) as executor:
        results = list(executor.map(write, output_files))

    assert {result.cache_key for result in results} == {cache_key}
    entries = list_cache_entries(service_type=ServiceType.TECHNICAL_ANALYSIS)
    assert len(entries) == 1
    assert entries[0].cache_key == cache_key
    assert entries[0].status == "active"
    assert entries[0].hit_count >= 0
    assert entries[0].output_files in [[output_file] for output_file in output_files]
    assert entries[0].artifact_owner_id in set(output_files)


def test_web_investment_cache_handlers_list_and_clear_entries(investment_env, monkeypatch):
    from business.investment.cache_service import build_cache_key, write_cache_entry
    from business.investment.constants import ServiceType
    from channel.web.web_channel import InvestmentCacheClearHandler, InvestmentCacheEntryInvalidateHandler, InvestmentCacheHandler

    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-25", "v1")
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=["/tmp/card.png", "/tmp/chart.png", "/tmp/report.md"],
        artifact_owner_id="request-1",
    )

    list_payload = _call_investment_json_handler(monkeypatch, InvestmentCacheHandler().GET, params={"limit": "20"})

    assert list_payload["status"] == "success", list_payload
    assert list_payload["entries"][0]["cache_key"] == cache_key
    assert list_payload["entries"][0]["output_files"] == ["/tmp/card.png", "/tmp/chart.png", "/tmp/report.md"]

    invalidate_payload = _call_investment_json_handler(
        monkeypatch,
        lambda: InvestmentCacheEntryInvalidateHandler().POST(cache_key),
        body={"operator": "tester"},
    )

    assert invalidate_payload["status"] == "success"
    assert invalidate_payload["invalidated"] is True

    default_list_payload = _call_investment_json_handler(monkeypatch, InvestmentCacheHandler().GET, params={"limit": "20"})
    assert default_list_payload["status"] == "success"
    assert default_list_payload["entries"] == []

    invalidated_list_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentCacheHandler().GET,
        params={"limit": "20", "include_invalidated": "1"},
    )
    assert invalidated_list_payload["status"] == "success"
    assert invalidated_list_payload["entries"][0]["cache_key"] == cache_key
    assert invalidated_list_payload["entries"][0]["status"] == "invalidated"

    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=["/tmp/card.png"],
        artifact_owner_id="request-2",
    )
    clear_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentCacheClearHandler().POST,
        body={"service_type": "technical_analysis", "market_date": "2026-05-25", "operator": "tester"},
    )

    assert clear_payload["status"] == "success"
    assert clear_payload["removed"] == 1


def test_web_investment_cache_handler_sanitizes_limit_and_rejects_unmatched_service_type(investment_env, monkeypatch):
    from business.investment.cache_service import build_cache_key, write_cache_entry
    from business.investment.constants import ServiceType
    from channel.web.web_channel import InvestmentCacheHandler

    cache_key = build_cache_key(ServiceType.TECHNICAL_ANALYSIS, "300502.SZ", "2026-05-25", "v1")
    write_cache_entry(
        cache_key=cache_key,
        service_type=ServiceType.TECHNICAL_ANALYSIS,
        normalized_target="300502.SZ",
        market_date="2026-05-25",
        version_fingerprint="v1",
        output_files=["/tmp/card.png"],
        artifact_owner_id="request-1",
    )

    invalid_limit_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentCacheHandler().GET,
        params={"service_type": "technical_analysis", "limit": "bad"},
    )
    unmatched_service_payload = _call_investment_json_handler(
        monkeypatch,
        InvestmentCacheHandler().GET,
        params={"service_type": "unknown-service", "limit": "20"},
    )

    assert invalid_limit_payload["status"] == "success"
    assert [entry["cache_key"] for entry in invalid_limit_payload["entries"]] == [cache_key]
    assert invalid_limit_payload["market_dates"] == ["2026-05-25"]
    assert unmatched_service_payload["status"] == "success"
    assert unmatched_service_payload["entries"] == []


def test_file_serve_handler_rejects_paths_outside_allowed_storage(investment_env, tmp_path, monkeypatch):
    from channel.web import web_channel
    from channel.web.web_channel import FileServeHandler

    _login_default_investment_admin(monkeypatch)
    secret = tmp_path / "outside-secret.txt"
    secret.write_text("secret", encoding="utf-8")
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(path=str(secret)))
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)

    with pytest.raises(web_channel.web.HTTPError):
        FileServeHandler().GET()


def test_file_serve_handler_allows_investment_storage_file(investment_env, monkeypatch):
    from business.investment.storage import get_storage_dirs
    from channel.web import web_channel
    from channel.web.web_channel import FileServeHandler

    _login_default_investment_admin(monkeypatch)
    image = get_storage_dirs()["generated"] / "allowed.png"
    image.write_bytes(b"png")
    monkeypatch.setattr(web_channel.web, "input", lambda **_defaults: SimpleNamespace(path=str(image)))
    monkeypatch.setattr(web_channel.web, "header", lambda *args, **kwargs: None)

    assert FileServeHandler().GET() == b"png"


def test_stock_resolver_resolves_codes_names_and_business_prompts(investment_env, monkeypatch):
    from business.investment import stock_resolver
    from business.investment.constants import ErrorCode, user_message
    from business.investment.stock_resolver import (
        list_stock_symbols,
        refresh_stock_symbols,
        resolve_stock as core_resolve_stock,
        stock_dictionary_stats,
    )
    from business.investment.technical_analysis import resolve_stock as technical_resolve_stock
    from business.investment.technical_analysis import run_technical_analysis

    monkeypatch.setattr(stock_resolver, "refresh_from_auto", lambda: {})
    inserted = refresh_stock_symbols(
        [
            {"code": "300502.SZ", "name": "新易盛", "market": "SZ", "ts_code": "300502.SZ"},
            {"code": "000001.SZ", "name": "重名", "market": "SZ"},
            {"code": "600001.SH", "name": "重名", "market": "SH", "source": "row-source"},
        ],
        source="test-source",
    )

    assert inserted == 3
    assert core_resolve_stock("300502") == ("300502.SZ", None)
    assert core_resolve_stock("300502.SZ") == ("300502.SZ", None)
    assert core_resolve_stock("600519.SH") == ("600519.SH", None)
    assert core_resolve_stock("新易盛") == ("300502.SZ", None)
    assert core_resolve_stock("不存在的股票") == (None, ErrorCode.STOCK_NOT_FOUND)
    assert core_resolve_stock("重名") == (None, ErrorCode.STOCK_AMBIGUOUS)
    assert technical_resolve_stock("新易盛") == ("300502.SZ", None)
    assert [row["code"] for row in list_stock_symbols("新", limit=5)] == ["300502.SZ"]
    assert stock_dictionary_stats()["total"] == 3

    not_found = run_technical_analysis("ok", "不存在的股票 技术分析")
    ambiguous_symbol, ambiguous_error = technical_resolve_stock("重名")

    assert not_found.success is False
    assert not_found.error_code == ErrorCode.STOCK_NOT_FOUND
    assert not_found.user_prompt == user_message(ErrorCode.STOCK_NOT_FOUND)
    assert ambiguous_symbol is None
    assert ambiguous_error == ErrorCode.STOCK_AMBIGUOUS
    assert user_message(ambiguous_error) == "股票名称匹配到多个标的，请改用股票代码。"


def test_stock_resolver_auto_refreshes_once_for_name_miss_then_resolves(investment_env, monkeypatch):
    from business.investment import stock_resolver

    calls = []

    def fake_refresh():
        calls.append("refresh")
        stock_resolver.refresh_stock_symbols(
            [{"code": "300502.SZ", "name": "新易盛", "market": "SZ", "source": "fake"}],
            source="fake",
        )
        return {"fake": {"count": 1}}

    monkeypatch.setattr(stock_resolver, "refresh_from_auto", fake_refresh)

    assert stock_resolver.resolve_stock("新易盛") == ("300502.SZ", None)
    assert calls == ["refresh"]


def test_stock_resolver_auto_refresh_miss_respects_flag_errors_and_codes(investment_env, monkeypatch):
    from business.investment import stock_resolver
    from business.investment.constants import ErrorCode

    calls = []

    def fake_refresh():
        calls.append("refresh")
        raise RuntimeError("refresh failed")

    monkeypatch.setattr(stock_resolver, "refresh_from_auto", fake_refresh)

    assert stock_resolver.resolve_stock("新易盛", auto_refresh_on_miss=False) == (None, ErrorCode.STOCK_NOT_FOUND)
    assert stock_resolver.resolve_stock("新易盛") == (None, ErrorCode.STOCK_NOT_FOUND)
    assert stock_resolver.resolve_stock("300502") == ("300502.SZ", None)
    assert stock_resolver.resolve_stock("300502.SZ") == ("300502.SZ", None)
    assert calls == ["refresh"]


def test_stock_resolver_auto_refresh_reports_ambiguous_after_refresh(investment_env, monkeypatch):
    from business.investment import stock_resolver
    from business.investment.constants import ErrorCode

    def fake_refresh():
        stock_resolver.refresh_stock_symbols(
            [
                {"code": "000001.SZ", "name": "重名", "market": "SZ"},
                {"code": "600001.SH", "name": "重名", "market": "SH"},
            ],
            source="fake",
        )
        return {"fake": {"count": 2}}

    monkeypatch.setattr(stock_resolver, "refresh_from_auto", fake_refresh)

    assert stock_resolver.resolve_stock("重名") == (None, ErrorCode.STOCK_AMBIGUOUS)


class _FakeDataFrame:
    def __init__(self, rows):
        self._rows = rows

    def to_dict(self, orient):
        assert orient == "records"
        return self._rows


def test_stock_resolver_refreshes_from_akshare_fake_dataframe(investment_env, monkeypatch):
    from business.investment import stock_resolver

    fake_akshare = SimpleNamespace(
        stock_info_a_code_name=lambda: _FakeDataFrame(
            [
                {"code": "600519", "name": "贵州茅台"},
                {"代码": "300502", "名称": "新易盛"},
            ]
        )
    )
    monkeypatch.setitem(sys.modules, "akshare", fake_akshare)

    assert stock_resolver.refresh_from_akshare() == 2

    rows = {row["code"]: row for row in stock_resolver.list_stock_symbols(limit=10)}
    assert rows["600519.SH"]["name"] == "贵州茅台"
    assert rows["600519.SH"]["market"] == "SH"
    assert rows["600519.SH"]["source"] == "akshare"
    assert rows["300502.SZ"]["name"] == "新易盛"
    assert rows["300502.SZ"]["market"] == "SZ"
    assert rows["300502.SZ"]["source"] == "akshare"


def test_stock_resolver_refreshes_large_symbol_batch(investment_env):
    from business.investment import stock_resolver

    rows = [
        {"code": f"{index:06d}.SZ", "name": f"Test Stock {index}", "market": "SZ"}
        for index in range(6000)
    ]

    assert stock_resolver.refresh_stock_symbols(rows, source="large-batch") == 6000
    assert stock_resolver.stock_dictionary_stats()["total"] == 6000


def test_stock_resolver_tushare_token_priority_and_masking(investment_env, tmp_path, monkeypatch):
    from business.investment import stock_resolver
    from business.investment.config_service import get_config, save_config

    monkeypatch.setattr(stock_resolver.Path, "home", lambda: tmp_path)
    (tmp_path / ".tushare_token").write_text("file-token-1234567890", encoding="utf-8")
    monkeypatch.setenv("TUSHARE_TOKEN", "env-token-1234567890")

    assert stock_resolver.get_tushare_token() == "env-token-1234567890"

    save_config("tushare.token", "config-token-1234567890", operator_role="admin")

    assert stock_resolver.get_tushare_token() == "config-token-1234567890"
    assert stock_resolver.get_tushare_token(masked=True) == "conf**********7890"
    assert get_config("tushare.token", masked=True) == "conf**********7890"


def test_stock_resolver_tushare_token_falls_back_to_file(investment_env, tmp_path, monkeypatch):
    from business.investment import stock_resolver

    monkeypatch.delenv("TUSHARE_TOKEN", raising=False)
    monkeypatch.setattr(stock_resolver.Path, "home", lambda: tmp_path)
    (tmp_path / ".tushare_token").write_text("file-token-1234567890\n", encoding="utf-8")

    assert stock_resolver.get_tushare_token() == "file-token-1234567890"


def test_stock_resolver_refresh_from_tushare_requires_token(investment_env, tmp_path, monkeypatch):
    from business.investment import stock_resolver

    monkeypatch.delenv("TUSHARE_TOKEN", raising=False)
    monkeypatch.setattr(stock_resolver.Path, "home", lambda: tmp_path)

    with pytest.raises(RuntimeError, match="tushare token not configured"):
        stock_resolver.refresh_from_tushare()


def test_stock_resolver_refreshes_from_tushare_fake_dataframe(investment_env, monkeypatch):
    from business.investment import stock_resolver
    from business.investment.config_service import save_config

    calls = []

    class FakePro:
        def stock_basic(self, **kwargs):
            calls.append(kwargs)
            return _FakeDataFrame(
                [
                    {"ts_code": "600519.SH", "symbol": "600519", "name": "贵州茅台", "exchange": "SSE"},
                    {"ts_code": "300502.SZ", "symbol": "300502", "name": "新易盛", "exchange": "SZSE"},
                ]
            )

    fake_tushare = SimpleNamespace(pro_api=lambda token: calls.append({"token": token}) or FakePro())
    monkeypatch.setitem(sys.modules, "tushare", fake_tushare)
    save_config("tushare.token", "config-token-1234567890", operator_role="admin")

    assert stock_resolver.refresh_from_tushare() == 2

    assert calls == [
        {"token": "config-token-1234567890"},
        {"exchange": "", "list_status": "L", "fields": "ts_code,symbol,name,exchange"},
    ]
    rows = {row["code"]: row for row in stock_resolver.list_stock_symbols(limit=10)}
    assert rows["600519.SH"]["ts_code"] == "600519.SH"
    assert rows["600519.SH"]["market"] == "SH"
    assert rows["600519.SH"]["source"] == "tushare"
    assert rows["300502.SZ"]["ts_code"] == "300502.SZ"
    assert rows["300502.SZ"]["market"] == "SZ"
    assert rows["300502.SZ"]["source"] == "tushare"


def test_stock_resolver_auto_reports_source_counts_and_errors(investment_env, monkeypatch):
    from business.investment import stock_resolver

    monkeypatch.setattr(stock_resolver, "refresh_from_akshare", lambda: 2)
    monkeypatch.setattr(stock_resolver, "get_tushare_token", lambda masked=False: "token")
    monkeypatch.setattr(stock_resolver, "refresh_from_tushare", lambda: 3)

    assert stock_resolver.refresh_from_auto() == {"akshare": {"count": 2}, "tushare": {"count": 3}}

    monkeypatch.setattr(stock_resolver, "refresh_from_akshare", lambda: (_ for _ in ()).throw(RuntimeError("ak failed")))
    monkeypatch.setattr(stock_resolver, "refresh_from_tushare", lambda: (_ for _ in ()).throw(RuntimeError("ts failed")))

    assert stock_resolver.refresh_from_auto() == {
        "akshare": {"error": "ak failed"},
        "tushare": {"error": "ts failed"},
    }


def test_refresh_investment_stocks_script_dispatches_sources(investment_env, monkeypatch, capsys):
    from scripts import refresh_investment_stocks

    calls = []
    monkeypatch.setattr(refresh_investment_stocks.storage, "initialize_storage", lambda: calls.append("init"))
    monkeypatch.setattr(refresh_investment_stocks.stock_resolver, "refresh_from_auto", lambda: calls.append("auto") or {"akshare": {"count": 2}})
    monkeypatch.setattr(refresh_investment_stocks.stock_resolver, "refresh_from_akshare", lambda: calls.append("akshare") or 3)
    monkeypatch.setattr(refresh_investment_stocks.stock_resolver, "refresh_from_tushare", lambda: calls.append("tushare") or 4)

    assert refresh_investment_stocks.main(["--source", "auto", "--json"]) == 0
    assert refresh_investment_stocks.main(["--source", "akshare", "--json"]) == 0
    assert refresh_investment_stocks.main(["--source", "tushare", "--json"]) == 0

    assert calls == ["init", "auto", "init", "akshare", "init", "tushare"]
    payloads = [json.loads(line) for line in capsys.readouterr().out.strip().splitlines()]
    assert [payload["source"] for payload in payloads] == ["auto", "akshare", "tushare"]
    assert [payload["count"] for payload in payloads] == [2, 3, 4]
    assert all(payload["success"] is True for payload in payloads)
    assert all("db_path" in payload for payload in payloads)


def test_refresh_investment_stocks_script_exits_one_when_all_sources_fail(investment_env, monkeypatch, capsys):
    from scripts import refresh_investment_stocks

    monkeypatch.setattr(refresh_investment_stocks.storage, "initialize_storage", lambda: None)
    monkeypatch.setattr(
        refresh_investment_stocks.stock_resolver,
        "refresh_from_auto",
        lambda: {"akshare": {"error": "ak failed"}, "tushare": {"error": "ts failed"}},
    )

    assert refresh_investment_stocks.main(["--source", "auto", "--json"]) == 1

    payload = json.loads(capsys.readouterr().out)
    assert payload["source"] == "auto"
    assert payload["success"] is False
    assert payload["count"] == 0
    assert "akshare: ak failed" in payload["error"]
    assert "tushare: ts failed" in payload["error"]


def test_refresh_investment_stocks_script_reports_single_source_exceptions(investment_env, monkeypatch, capsys):
    from scripts import refresh_investment_stocks

    monkeypatch.setattr(refresh_investment_stocks.storage, "initialize_storage", lambda: None)
    monkeypatch.setattr(
        refresh_investment_stocks.stock_resolver,
        "refresh_from_akshare",
        lambda: (_ for _ in ()).throw(RuntimeError("provider unavailable")),
    )

    assert refresh_investment_stocks.main(["--source", "akshare", "--json"]) == 1

    payload = json.loads(capsys.readouterr().out)
    assert payload["source"] == "akshare"
    assert payload["success"] is False
    assert payload["count"] == 0
    assert payload["error"] == "provider unavailable"


def test_refresh_investment_stocks_script_runs_by_file_path():
    repo_root = Path(__file__).resolve().parents[1]

    result = subprocess.run(
        [sys.executable, "scripts/refresh_investment_stocks.py", "--help"],
        cwd=repo_root,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0
    assert "--source" in result.stdout


def test_sqlite_to_pg_migration_rejects_missing_sqlite_file(tmp_path):
    from scripts import migrate_investment_sqlite_to_pg

    missing = tmp_path / "missing.db"

    with pytest.raises(FileNotFoundError, match="SQLite source does not exist"):
        migrate_investment_sqlite_to_pg.migrate_sqlite_to_postgres(
            missing,
            "postgresql+psycopg://user:secret@localhost:5432/cowagent",
        )


def test_sqlite_to_pg_migration_rejects_non_postgresql_url(tmp_path):
    from scripts import migrate_investment_sqlite_to_pg

    sqlite_path = tmp_path / "investment.db"
    sqlite_path.write_bytes(b"")

    with pytest.raises(ValueError, match="target URL must be PostgreSQL"):
        migrate_investment_sqlite_to_pg.migrate_sqlite_to_postgres(
            sqlite_path,
            f"sqlite:///{tmp_path / 'target.db'}",
        )


def test_sqlite_to_pg_migration_reads_all_tables_and_upserts(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from business.investment import schema
    from scripts import migrate_investment_sqlite_to_pg

    sqlite_path = tmp_path / "investment.db"
    engine = create_engine(f"sqlite:///{sqlite_path}", future=True)
    schema.metadata.create_all(engine)
    now = "2026-05-25T10:00:00"
    with engine.begin() as conn:
        conn.execute(
            schema.investment_users.insert(),
            {
                "id": 1,
                "openid": "openid-1",
                "name": "Alice",
                "enabled": 1,
                "allowed_services": "[\"rate\"]",
                "created_at": now,
                "updated_at": now,
            },
        )
        conn.execute(
            schema.investment_request_records.insert(),
            {
                "request_id": "request-1",
                "openid": "openid-1",
                "raw_input": "利率",
                "service_type": "rate",
                "status": "success",
                "created_at": now,
                "updated_at": now,
            },
        )
        conn.execute(
            schema.investment_daily_contents.insert(),
            {
                "content_id": "content-1",
                "service_type": "rate",
                "status": "draft",
                "created_at": now,
                "updated_at": now,
            },
        )
        conn.execute(
            schema.investment_output_files.insert(),
            {
                "id": 1,
                "owner_id": "content-1",
                "file_path": "/tmp/rate.png",
                "file_type": "image",
                "service_type": "rate",
                "created_at": now,
            },
        )
        conn.execute(
            schema.investment_configs.insert(),
            {
                "config_key": "model.api_key",
                "config_value": "sk-test-secret",
                "updated_at": now,
                "updated_by": "admin",
            },
        )
        conn.execute(
            schema.investment_configs.insert(),
            {
                "config_key": "wechatmp.token",
                "config_value": "wx-test-token",
                "updated_at": now,
                "updated_by": "admin",
            },
        )
        conn.execute(
            schema.investment_configs.insert(),
            {
                "config_key": "tushare.token",
                "config_value": "ts-test-token",
                "updated_at": now,
                "updated_by": "admin",
            },
        )
        conn.execute(
            schema.investment_stock_symbols.insert(),
            {
                "code": "300502.SZ",
                "name": "新易盛",
                "market": "SZ",
                "ts_code": "300502.SZ",
                "source": "seed",
                "updated_at": now,
            },
        )

    tables = migrate_investment_sqlite_to_pg.read_sqlite_tables(sqlite_path)
    assert list(tables) == [
        "investment_users",
        "investment_request_records",
        "investment_daily_contents",
        "investment_output_files",
        "investment_configs",
        "investment_stock_symbols",
    ]
    assert {name: len(rows) for name, rows in tables.items()} == {
        **{name: 1 for name in tables},
        "investment_configs": 3,
    }

    calls = []

    class FakeConnection:
        dialect = SimpleNamespace(name="postgresql")

        def execute(self, statement):
            calls.append(("execute", statement))

    class FakeBegin:
        def __enter__(self):
            return FakeConnection()

        def __exit__(self, exc_type, exc, tb):
            return False

    class FakeEngine:
        def begin(self):
            return FakeBegin()

    monkeypatch.setattr(schema.metadata, "create_all", lambda conn: pytest.fail("migration target schema must be built by Alembic"))
    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg,
        "upgrade_investment_schema",
        lambda pg_url: calls.append(("alembic_upgrade", pg_url)),
        raising=False,
    )
    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg.investment_db,
        "upsert_config",
        lambda conn, key, value, updated_at, updated_by: calls.append(("config", key, value, updated_at, updated_by)),
    )
    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg.investment_db,
        "upsert_stock_symbols",
        lambda conn, rows: calls.append(("stocks", rows)),
    )

    summary = migrate_investment_sqlite_to_pg.copy_tables_to_postgres(
        "postgresql+psycopg://user:secret@localhost:5432/cowagent",
        tables,
        engine_factory=lambda url, future: calls.append(("engine", url, future)) or FakeEngine(),
    )

    assert summary == {
        **{name: 1 for name in tables},
        "investment_configs": 1,
    }
    assert ("engine", "postgresql+psycopg://user:secret@localhost:5432/cowagent", True) in calls
    assert ("alembic_upgrade", "postgresql+psycopg://user:secret@localhost:5432/cowagent") in calls
    assert calls.index(("alembic_upgrade", "postgresql+psycopg://user:secret@localhost:5432/cowagent")) < next(
        index for index, call in enumerate(calls) if call[0] in {"execute", "config", "stocks"}
    )
    assert any(call[:3] == ("config", "tushare.token", "ts-test-token") for call in calls)
    assert not any(call[:2] == ("config", "model.api_key") for call in calls)
    assert not any(call[:2] == ("config", "wechatmp.token") for call in calls)
    assert any(call[0] == "stocks" and call[1][0]["code"] == "300502.SZ" for call in calls)
    assert len([call for call in calls if call[0] == "execute"]) == 6


def test_sqlite_to_pg_migration_resets_postgres_sequences_after_copy(monkeypatch):
    from scripts import migrate_investment_sqlite_to_pg

    calls = []

    class FakeConnection:
        dialect = SimpleNamespace(name="postgresql")

        def execute(self, statement):
            calls.append(("execute", statement))

    class FakeBegin:
        def __enter__(self):
            return FakeConnection()

        def __exit__(self, exc_type, exc, tb):
            return False

    class FakeEngine:
        def begin(self):
            return FakeBegin()

    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg,
        "upgrade_investment_schema",
        lambda pg_url: calls.append(("alembic_upgrade", pg_url)),
        raising=False,
    )
    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg.schema.metadata,
        "create_all",
        lambda conn: pytest.fail("migration target schema must be built by Alembic"),
    )
    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg.investment_db,
        "upsert_stock_symbols",
        lambda conn, rows: calls.append(("stocks", rows)),
    )

    tables = OrderedDict((name, []) for name in migrate_investment_sqlite_to_pg.TABLE_NAMES)
    tables["investment_users"] = [
        {
            "id": 7,
            "openid": "openid-7",
            "name": "Alice",
            "enabled": 1,
            "allowed_services": "[\"rate\"]",
            "created_at": "2026-05-25T10:00:00",
            "updated_at": "2026-05-25T10:00:00",
        }
    ]
    tables["investment_output_files"] = [
        {
            "id": 11,
            "owner_id": "content-1",
            "file_path": "/tmp/rate.png",
            "file_type": "image",
            "service_type": "rate",
            "created_at": "2026-05-25T10:00:00",
        }
    ]

    migrate_investment_sqlite_to_pg.copy_tables_to_postgres(
        "postgresql+psycopg://user:secret@localhost:5432/cowagent",
        tables,
        engine_factory=lambda url, future: FakeEngine(),
    )

    executed_sql = [str(call[1]) for call in calls if call[0] == "execute"]
    sequence_sql = [sql for sql in executed_sql if "setval" in sql]

    assert len(sequence_sql) == 2
    assert "pg_get_serial_sequence('investment_users', 'id')" in sequence_sql[0]
    assert "pg_get_serial_sequence('investment_output_files', 'id')" in sequence_sql[1]
    assert all("(select count(*)" in sql and "> 0" in sql for sql in sequence_sql)
    assert executed_sql.index(sequence_sql[0]) > 1
    assert executed_sql.index(sequence_sql[1]) > 1


def test_sqlite_to_pg_migration_output_does_not_include_pg_url(tmp_path, monkeypatch, capsys):
    from scripts import migrate_investment_sqlite_to_pg

    sqlite_path = tmp_path / "investment.db"
    sqlite_path.write_bytes(b"")
    pg_url = "postgresql+psycopg://user:super-secret-password@localhost:5432/cowagent"

    monkeypatch.setattr(
        migrate_investment_sqlite_to_pg,
        "migrate_sqlite_to_postgres",
        lambda sqlite_arg, pg_arg: {
            "status": "success",
            "tables": {"investment_users": 1},
        },
    )

    assert migrate_investment_sqlite_to_pg.main(["--sqlite", str(sqlite_path), "--pg", pg_url]) == 0
    output = capsys.readouterr().out
    payload = json.loads(output)
    assert payload == {"status": "success", "tables": {"investment_users": 1}}
    assert pg_url not in output
    assert "super-secret-password" not in output


def test_tushare_token_config_permission_is_sensitive(investment_env):
    from business.investment.config_service import can_modify_config, save_config

    assert can_modify_config("tushare.token", "uploader") is False
    assert can_modify_config("tushare.token", "operator") is False
    assert can_modify_config("tushare.token", "technical_admin") is False
    assert can_modify_config("tushare.token", "technical_operator") is True

    with pytest.raises(PermissionError):
        save_config("tushare.token", "blocked-token", operator_role="operator")


def test_technical_analysis_sh_suffix_enters_skill_and_failures_return_business_prompts(investment_env, tmp_path, monkeypatch):
    from business.investment import technical_analysis
    from business.investment.constants import ErrorCode, user_message
    from business.investment.technical_analysis import run_technical_analysis

    report = tmp_path / "600519_技术分析报告_2026-05-25.md"
    chart = tmp_path / "600519_TA_2026-05-25.png"
    report.write_text("ta report", encoding="utf-8")
    chart.write_bytes(b"chart")
    calls = []

    def fake_skill(symbol, _output_dir):
        calls.append(symbol)
        return report, chart

    monkeypatch.setattr(technical_analysis, "_run_skill", fake_skill)
    monkeypatch.setattr(
        technical_analysis,
        "generate_technical_analysis_text",
        lambda _report_text: SimpleNamespace(success=False, detail="model failed"),
    )

    ai_failed = run_technical_analysis("ok", "600519.SH 技术分析")

    assert calls == ["600519"]
    assert ai_failed.success is False
    assert ai_failed.error_code == ErrorCode.TECHNICAL_ANALYSIS_FAILED
    assert ai_failed.user_prompt == user_message(ErrorCode.TECHNICAL_ANALYSIS_FAILED)

    monkeypatch.setattr(
        technical_analysis,
        "generate_technical_analysis_text",
        lambda _report_text: SimpleNamespace(success=True, text="standard"),
    )
    monkeypatch.setattr(
        technical_analysis,
        "render_technical_analysis_card",
        lambda _standard_text, _output_path: SimpleNamespace(success=False, detail="renderer failed"),
    )

    render_failed = run_technical_analysis("ok", "600519.SH 技术分析")

    assert render_failed.success is False
    assert render_failed.error_code == ErrorCode.IMAGE_GENERATION_FAILED
    assert render_failed.user_prompt == user_message(ErrorCode.IMAGE_GENERATION_FAILED)

    calls.clear()
    def successful_render(_standard_text, output_path):
        Path(output_path).write_bytes(b"card")
        return SimpleNamespace(success=True, image_path=output_path, detail="")

    monkeypatch.setattr(technical_analysis, "render_technical_analysis_card", successful_render)

    standard_code = run_technical_analysis("ok", "600519 技术分析")

    assert standard_code.success is True
    assert calls == ["600519"]


def test_router_records_technical_analysis_report_chart_and_card_paths(investment_env, tmp_path):
    from business.investment.constants import ServiceType
    from business.investment.records import list_request_records
    from business.investment.router import handle_text_message
    from business.investment.technical_analysis import TechnicalAnalysisResult
    from business.investment.user_service import create_user

    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    card = str(tmp_path / "card.png")
    chart = str(tmp_path / "chart.png")
    report = str(tmp_path / "report.md")

    reply = handle_text_message(
        "ok",
        "300502.SZ 技术分析",
        technical_analysis_handler=lambda _openid, _raw_input, _target: TechnicalAnalysisResult(
            True,
            signal_card_path=card,
            main_chart_path=chart,
            report_path=report,
            output_files=[card, chart, report],
        ),
    )

    record = list_request_records(limit=1)[0]
    assert reply.success is True
    assert reply.output_files == [card, chart]
    assert record.output_files == [card, chart, report]


def test_daily_content_activation_and_query(investment_env, tmp_path):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.daily_content import (
        create_content_draft,
        get_latest_effective_content,
        set_content_effective,
    )

    empty = get_latest_effective_content(ServiceType.RATE)
    assert empty.success is False
    assert empty.error_code == ErrorCode.NO_CONTENT

    first_img = tmp_path / "first.png"
    second_img = tmp_path / "second.png"
    first_img.write_bytes(b"png")
    second_img.write_bytes(b"png2")

    first = create_content_draft(ServiceType.RATE, source_text="first")
    second = create_content_draft(ServiceType.RATE, source_text="second")
    set_content_effective(first, str(first_img), operator="admin")
    set_content_effective(second, str(second_img), operator="admin")

    latest = get_latest_effective_content(ServiceType.RATE)
    assert latest.success is True
    assert latest.output_image == str(second_img)


def test_daily_content_versions_are_effective_per_service_and_date(investment_env, tmp_path):
    from datetime import date, timedelta

    from business.investment.constants import ServiceType, Status
    from business.investment.daily_content import (
        create_content_draft,
        get_latest_effective_content,
        set_content_effective,
    )
    from business.investment.records import get_content_record

    today = date.today()
    yesterday = (today - timedelta(days=1)).isoformat()
    tomorrow = (today + timedelta(days=1)).isoformat()

    yesterday_first_img = tmp_path / "yesterday-first.png"
    yesterday_second_img = tmp_path / "yesterday-second.png"
    tomorrow_img = tmp_path / "tomorrow.png"
    for path in (yesterday_first_img, yesterday_second_img, tomorrow_img):
        path.write_bytes(b"png")

    yesterday_first = create_content_draft(ServiceType.RATE, source_text="old")
    tomorrow_content = create_content_draft(ServiceType.RATE, source_text="future")
    yesterday_second = create_content_draft(ServiceType.RATE, source_text="replacement")

    set_content_effective(yesterday_first, str(yesterday_first_img), effective_date=yesterday, operator="operator-a")
    set_content_effective(tomorrow_content, str(tomorrow_img), effective_date=tomorrow, operator="operator-b")
    set_content_effective(yesterday_second, str(yesterday_second_img), effective_date=yesterday, operator="operator-c")

    assert get_content_record(yesterday_first).status == Status.ARCHIVED
    assert get_content_record(yesterday_first).archived_at
    assert get_content_record(yesterday_second).status == Status.EFFECTIVE
    assert get_content_record(yesterday_second).effective_date == yesterday
    assert get_content_record(tomorrow_content).status == Status.EFFECTIVE
    assert get_latest_effective_content(ServiceType.RATE).content_id == yesterday_second


def test_rate_direct_output_mode_uses_uploaded_png_without_ai_or_renderer(investment_env, tmp_path):
    from business.investment.constants import ServiceType, Status
    from business.investment.daily_content import create_rate_content_draft, generate_content, save_source_file
    from business.investment.records import get_content_record

    uploaded = save_source_file(ServiceType.RATE, "final-rate.png", b"png")
    content_id = create_rate_content_draft(
        source_files=[uploaded],
        source_text="",
        operator="operator-a",
        direct_output_mode=True,
    )

    result = generate_content(
        content_id,
        ai_generator=lambda *_args, **_kwargs: pytest.fail("direct output mode must not call AI"),
        renderer=lambda *_args, **_kwargs: pytest.fail("direct output mode must not call renderer"),
    )

    record = get_content_record(content_id)
    assert result.success is True
    assert result.output_image != uploaded
    assert Path(result.output_image).is_file()
    assert Path(result.output_image).read_bytes() == b"png"
    assert result.generated_text == ""
    assert record.output_image == result.output_image
    assert record.direct_output_mode is True
    assert record.status == Status.GENERATED


@pytest.mark.parametrize(
    ("case_name", "source_path_factory", "expected_detail"),
    [
        (
            "missing png",
            lambda uploads_dir, _tmp_path: str(uploads_dir / "rate" / "missing.png"),
            "direct output mode requires an existing PNG upload",
        ),
        (
            "directory png",
            lambda uploads_dir, _tmp_path: _mkdir_and_return(uploads_dir / "rate" / "directory.png"),
            "direct output mode requires an existing PNG upload",
        ),
        (
            "relative png",
            lambda _uploads_dir, _tmp_path: "relative.png",
            "direct output mode requires an absolute uploaded PNG path",
        ),
        (
            "non-png",
            lambda uploads_dir, _tmp_path: _write_and_return(uploads_dir / "rate" / "rate.xlsx", b"excel"),
            "direct output mode only supports PNG files",
        ),
    ],
)
def test_rate_direct_output_mode_rejects_invalid_uploaded_png_paths(
    investment_env,
    tmp_path,
    case_name,
    source_path_factory,
    expected_detail,
):
    from business.investment.constants import ErrorCode, ServiceType, Status
    from business.investment.daily_content import create_rate_content_draft, generate_content, save_source_file
    from business.investment.records import get_content_record, list_output_files
    from business.investment.storage import get_storage_dirs

    source_path = source_path_factory(get_storage_dirs()["uploads"], tmp_path)
    content_id = create_rate_content_draft(
        source_files=[source_path],
        source_text="",
        operator=f"operator-{case_name}",
        direct_output_mode=True,
    )

    result = generate_content(
        content_id,
        ai_generator=lambda *_args, **_kwargs: pytest.fail("invalid direct output mode must not call AI"),
        renderer=lambda *_args, **_kwargs: pytest.fail("invalid direct output mode must not call renderer"),
    )

    record = get_content_record(content_id)
    assert result.success is False
    assert result.error_code == ErrorCode.INPUT_ERROR
    assert expected_detail in result.detail
    assert record.status == Status.GENERATE_FAILED
    assert expected_detail in record.error_message
    assert record.generated_text == ""
    assert record.output_image == ""
    assert list_output_files(content_id) == []


def _mkdir_and_return(path: Path) -> str:
    path.mkdir(parents=True, exist_ok=True)
    return str(path)


def _write_and_return(path: Path, content: bytes) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return str(path)


def test_daily_content_operation_audits_track_create_generate_effective_and_archive(investment_env, tmp_path):
    from business.investment.audit_service import list_operation_audits
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, generate_content, set_content_effective

    first = create_content_draft(ServiceType.RATE, source_text="first", operator="operator-a")
    second = create_content_draft(ServiceType.RATE, source_text="second", operator="operator-b")
    generate_content(
        first,
        ai_generator=lambda _service_type, _source_text: SimpleNamespace(success=True, text="first text"),
        renderer=lambda _service_type, _text: SimpleNamespace(success=True, image_path=str(tmp_path / "first.png")),
    )
    set_content_effective(first, operator="operator-a")
    set_content_effective(second, str(tmp_path / "second.png"), operator="operator-b")

    audits = list_operation_audits(limit=20)
    actions = [audit.action for audit in audits]

    assert "content.create" in actions
    assert "content.generate" in actions
    assert "content.publish" in actions
    assert "content.archive" in actions
    effective = next(audit for audit in audits if audit.action == "content.publish" and audit.target_id == second)
    assert effective.operator == "operator-b"
    assert effective.target_type == "daily_content"
    assert effective.operation_category == "content"
    assert effective.detail["service_type"] == ServiceType.RATE


def test_daily_content_create_upload_generate_and_failure_records(investment_env):
    from business.investment.constants import ErrorCode, ServiceType, Status
    from business.investment.daily_content import (
        create_convertible_bond_content_draft,
        create_rate_content_draft,
        generate_content,
        update_content_source,
    )
    from business.investment.records import get_content_record

    content_id = create_rate_content_draft(source_files=["/upload/original.xlsx"], source_text="old", operator="operator-a")
    update_content_source(content_id, source_files=["/upload/latest.xlsx"], source_text="updated rate data")

    calls = []

    def fake_ai(service_type, source_text):
        calls.append(("ai", service_type, source_text))
        return SimpleNamespace(success=True, text="standard rate text")

    def fake_renderer(service_type, standard_text):
        calls.append(("renderer", service_type, standard_text))
        return SimpleNamespace(success=True, image_path="/generated/rate.png")

    result = generate_content(content_id, ai_generator=fake_ai, renderer=fake_renderer)

    assert result.success is True
    assert result.generated_text == "standard rate text"
    assert result.output_image == "/generated/rate.png"
    assert calls == [
        ("ai", ServiceType.RATE, "updated rate data"),
        ("renderer", ServiceType.RATE, "standard rate text"),
    ]

    record = get_content_record(content_id)
    assert record.service_type == ServiceType.RATE
    assert record.source_files == ["/upload/latest.xlsx"]
    assert record.source_text == "updated rate data"
    assert record.generated_text == "standard rate text"
    assert record.output_image == "/generated/rate.png"
    assert record.status == Status.GENERATED
    assert record.error_message == ""

    failed_id = create_convertible_bond_content_draft(source_text="cb data\n\njoke", operator="operator-b")
    failed = generate_content(
        failed_id,
        ai_generator=lambda _service_type, _source_text: SimpleNamespace(success=False, detail="model rejected payload"),
        renderer=fake_renderer,
    )
    failed_record = get_content_record(failed_id)
    assert failed.success is False
    assert failed.error_code == ErrorCode.SYSTEM_ERROR
    assert failed_record.service_type == ServiceType.CONVERTIBLE_BOND
    assert failed_record.status == Status.GENERATE_FAILED
    assert failed_record.error_message == "model rejected payload"


def test_daily_content_default_generation_passes_uploaded_source_files(investment_env, tmp_path, monkeypatch):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_rate_content_draft, generate_content

    uploaded = tmp_path / "uploaded-rate.png"
    uploaded.write_bytes(b"png")
    content_id = create_rate_content_draft(source_files=[str(uploaded)], source_text="", operator="operator-a")
    captured = {}

    def fake_generate_standard_text(service_type, source_text, *, source_files=None):
        captured["service_type"] = service_type
        captured["source_text"] = source_text
        captured["source_files"] = source_files
        return SimpleNamespace(success=True, text="standard rate text")

    monkeypatch.setattr(
        "business.investment.ai_generation.generate_standard_text",
        fake_generate_standard_text,
    )

    result = generate_content(
        content_id,
        renderer=lambda _service_type, _text: SimpleNamespace(success=True, image_path="/generated/rate.png"),
    )

    assert result.success is True
    assert captured == {
        "service_type": ServiceType.RATE,
        "source_text": "",
        "source_files": [str(uploaded)],
    }


def test_daily_content_default_image_generation_renders_png_with_fake_model(investment_env, tmp_path, monkeypatch):
    from business.investment import config_service
    from business.investment.config_service import save_configs
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_rate_content_draft, generate_content
    from business.investment.records import get_content_record
    from business.investment.storage import get_storage_dirs

    source_image = tmp_path / "uploaded-rate.png"
    source_image.write_bytes(b"\x89PNG\r\n\x1a\nimage")
    output_dir = tmp_path / "generated"
    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "vision-model",
            "custom_api_base": "https://model.example/v1",
            "custom_api_key": "sk-e2e-image-1234567890",
        },
    )
    save_configs(
        {"render.output_dir": str(output_dir)},
        operator_role="admin",
    )
    standard_text = Path("skills/signal-card-renderer/examples/bond_sample.txt").read_text(encoding="utf-8")

    class FakeClient:
        def __init__(self):
            self.calls = []

        def chat_completions(self, **kwargs):
            self.calls.append(kwargs)
            user_content = kwargs["messages"][1]["content"]
            if len(self.calls) == 1:
                assert any(block.get("type") == "image_url" for block in user_content)
                return {"choices": [{"message": {"content": "OCR: 2026-05-25 108.970 入场（3/8）"}}]}
            assert "必须只输出卡片正文" in kwargs["messages"][0]["content"]
            assert "OCR: 2026-05-25 108.970 入场（3/8）" in user_content
            return {"choices": [{"message": {"content": standard_text}}]}

    fake_client = FakeClient()
    monkeypatch.setattr(
        "models.openai.openai_http_client.get_default_client",
        lambda: fake_client,
    )

    content_id = create_rate_content_draft(
        source_files=[str(source_image)],
        source_text="",
        operator="operator-a",
        effective_date="2026-05-25",
    )
    result = generate_content(content_id)

    assert result.success is True
    assert Path(result.output_image).is_file()
    assert Path(result.output_image).stat().st_size > 0
    assert Path(result.output_image).resolve().is_relative_to((get_storage_dirs()["generated"] / "archive").resolve())
    assert Path(result.output_image).name.startswith("output_image_rate_2026-05-25_v1_")
    assert Path(result.output_image).name.endswith(".png")
    assert result.output_image != str(output_dir / f"{ServiceType.RATE}_card.png")
    assert get_content_record(content_id).output_image == result.output_image
    assert fake_client.calls


def test_daily_content_records_filter_by_service_type_for_console_pages(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_convertible_bond_content_draft, create_rate_content_draft
    from business.investment.records import list_content_records

    rate_id = create_rate_content_draft(source_text="rate")
    cb_id = create_convertible_bond_content_draft(source_text="cb")

    rate_records = list_content_records(service_type=ServiceType.RATE)
    cb_records = list_content_records(service_type=ServiceType.CONVERTIBLE_BOND)

    assert [record.content_id for record in rate_records] == [rate_id]
    assert [record.content_id for record in cb_records] == [cb_id]


def test_daily_content_upload_saves_files_under_investment_upload_dir(investment_env):
    from business.investment.constants import ServiceType
    from business.investment.daily_content import save_source_file

    saved = save_source_file(ServiceType.RATE, "rates.xlsx", b"rate-data")

    assert Path(saved).read_bytes() == b"rate-data"
    assert investment_env / "storage" / "uploads" in Path(saved).parents

    with pytest.raises(ValueError):
        save_source_file(ServiceType.RATE, "../escape.txt", b"bad")


def test_daily_content_regenerate_updates_output_and_only_latest_is_effective(investment_env, tmp_path):
    from business.investment.constants import ServiceType, Status
    from business.investment.daily_content import (
        create_rate_content_draft,
        generate_content,
        get_latest_effective_content,
        regenerate_content,
        set_content_effective,
    )
    from business.investment.records import get_content_record

    first_id = create_rate_content_draft(source_text="first")
    second_id = create_rate_content_draft(source_text="second")
    first_image = tmp_path / "first.png"
    second_image = tmp_path / "second.png"
    regenerated_image = tmp_path / "second-v2.png"
    first_image.write_bytes(b"first")
    second_image.write_bytes(b"second")
    regenerated_image.write_bytes(b"second-v2")

    generate_content(
        first_id,
        ai_generator=lambda _service_type, _source_text: SimpleNamespace(success=True, text="first text"),
        renderer=lambda _service_type, _text: SimpleNamespace(success=True, image_path=str(first_image)),
    )
    generate_content(
        second_id,
        ai_generator=lambda _service_type, _source_text: SimpleNamespace(success=True, text="second text"),
        renderer=lambda _service_type, _text: SimpleNamespace(success=True, image_path=str(second_image)),
    )
    set_content_effective(first_id, operator="operator-a")
    set_content_effective(second_id, operator="operator-b")

    assert get_content_record(first_id).status == Status.ARCHIVED
    second_record = get_content_record(second_id)
    assert second_record.status == Status.EFFECTIVE
    assert second_record.output_image != str(second_image)
    assert Path(second_record.output_image).read_bytes() == b"second"
    assert get_latest_effective_content(ServiceType.RATE).output_image == second_record.output_image

    regenerated = regenerate_content(
        second_id,
        ai_generator=lambda _service_type, _source_text: SimpleNamespace(success=True, text="second text regenerated"),
        renderer=lambda _service_type, _text: SimpleNamespace(success=True, image_path=str(regenerated_image)),
    )

    assert regenerated.success is True
    assert regenerated.generated_text == "second text regenerated"
    assert regenerated.output_image != str(regenerated_image)
    assert Path(regenerated.output_image).read_bytes() == b"second-v2"
    updated_record = get_content_record(second_id)
    assert updated_record.generated_text == "second text regenerated"
    assert updated_record.output_image == regenerated.output_image
    assert updated_record.status == Status.GENERATED


def test_daily_content_convertible_bond_no_effective_content_prompt(investment_env):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.daily_content import get_latest_effective_content

    empty = get_latest_effective_content(ServiceType.CONVERTIBLE_BOND)
    assert empty.success is False
    assert empty.error_code == ErrorCode.NO_CONTENT
    assert empty.user_prompt == "今日内容尚未更新，请稍后再试。"


@pytest.mark.parametrize("service_type", ["rate", "convertible_bond"])
def test_daily_content_effective_content_image_missing_returns_no_content(investment_env, tmp_path, service_type):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.daily_content import create_content_draft, get_latest_effective_content, set_content_effective

    service = ServiceType(service_type)
    missing_image = tmp_path / f"missing-{service_type}.png"
    content_id = create_content_draft(service, source_text="content")
    set_content_effective(content_id, str(missing_image), operator="admin")

    result = get_latest_effective_content(service)

    assert result.success is False
    assert result.error_code == ErrorCode.NO_CONTENT
    assert result.user_prompt == "今日内容尚未更新，请稍后再试。"
    assert "effective content image missing" in result.detail
    assert str(missing_image) in result.detail


def test_render_service_validates_output_files(investment_env, tmp_path):
    from business.investment.constants import ErrorCode, ServiceType
    from business.investment.render_service import RenderRequest, render_card

    output = tmp_path / "card.png"

    def fake_renderer(_request, target):
        Path(target).write_bytes(b"png")

    result = render_card(
        RenderRequest(service_type=ServiceType.RATE, standard_text="text", output_path=str(output)),
        renderer=fake_renderer,
    )
    assert result.success is True
    assert result.image_path == str(output)

    def empty_renderer(_request, target):
        Path(target).write_bytes(b"")

    failed = render_card(
        RenderRequest(service_type=ServiceType.RATE, standard_text="text", output_path=str(output)),
        renderer=empty_renderer,
    )
    assert failed.success is False
    assert failed.error_code == ErrorCode.IMAGE_GENERATION_FAILED


def test_render_service_contract_uses_configured_templates_and_output_dir(investment_env, tmp_path):
    from business.investment.config_service import save_configs
    from business.investment.constants import ServiceType, Status
    from business.investment.render_service import RenderRequest, render_card

    template_ta = tmp_path / "template_ta.html"
    template_bond = tmp_path / "template_bond.html"
    template_cb = tmp_path / "template_cb.html"
    for template in (template_ta, template_bond, template_cb):
        template.write_text("<html></html>", encoding="utf-8")
    output_dir = tmp_path / "rendered"
    save_configs(
        {
            "render.template_ta_path": str(template_ta),
            "render.template_rate_path": str(template_bond),
            "render.template_cb_path": str(template_cb),
            "render.output_dir": str(output_dir),
        },
        operator_role="admin",
    )
    calls = []

    def fake_renderer(request, target):
        calls.append((request.service_type, request.standard_text, request.output_dir, request.output_path, request.template_path, target))
        Path(target).write_bytes(b"png")

    result = render_card(
        RenderRequest(service_type=ServiceType.CONVERTIBLE_BOND, standard_text="cb standard text"),
        renderer=fake_renderer,
    )

    assert result.success is True
    assert result.status == Status.SUCCESS
    assert result.service_type == ServiceType.CONVERTIBLE_BOND
    assert result.standard_text == "cb standard text"
    assert result.output_dir == str(output_dir)
    assert result.output_path == result.image_path
    assert result.image_path.startswith(str(output_dir))
    assert result.failure_reason == ""
    assert calls == [
        (
            ServiceType.CONVERTIBLE_BOND,
            "cb standard text",
            str(output_dir),
            result.output_path,
            str(template_cb),
            result.output_path,
        )
    ]


def test_render_health_check_reports_renderer_template_and_chromium_details(investment_env, tmp_path, monkeypatch):
    from business.investment import health
    from business.investment.config_service import save_configs

    missing_renderer = tmp_path / "missing-render-card.py"
    missing_ta = tmp_path / "missing-template-ta.html"
    missing_bond = tmp_path / "missing-template-bond.html"
    missing_cb = tmp_path / "missing-template-cb.html"
    save_configs(
        {
            "render.renderer_path": str(missing_renderer),
            "render.template_ta_path": str(missing_ta),
            "render.template_rate_path": str(missing_bond),
            "render.template_cb_path": str(missing_cb),
        },
        operator_role="admin",
    )
    monkeypatch.setattr(
        health,
        "_check_playwright_chromium",
        lambda: health.HealthItem("playwright_chromium", False, "chromium executable missing: C:/missing/chrome.exe"),
        raising=False,
    )

    items = {item.name: item for item in health.run_health_checks()}

    assert items["signal_card_renderer"].ok is False
    assert str(missing_renderer) in items["signal_card_renderer"].detail
    assert items["template_ta"].ok is False
    assert str(missing_ta) in items["template_ta"].detail
    assert items["template_bond"].ok is False
    assert str(missing_bond) in items["template_bond"].detail
    assert items["template_cb"].ok is False
    assert str(missing_cb) in items["template_cb"].detail
    assert items["playwright_chromium"].ok is False
    assert "chromium executable missing" in items["playwright_chromium"].detail


def test_health_check_levels_dependencies_and_wechatmp_config(investment_env, monkeypatch):
    from business.investment import config_service, health

    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "channel_type": "wechatmp_service",
            "wechatmp_app_id": "",
            "wechatmp_app_secret": "",
            "wechatmp_token": "token-ok",
            "wechatmp_aes_key": "",
        },
    )
    monkeypatch.setattr(
        health,
        "_dependency_available",
        lambda module_name: (module_name in {"pandas", "scipy"}, "" if module_name in {"pandas", "scipy"} else "not installed"),
        raising=False,
    )
    monkeypatch.setattr(health, "_has_chinese_font", lambda: False, raising=False)
    monkeypatch.setattr(
        health,
        "_check_playwright_chromium",
        lambda: health.HealthItem("playwright_chromium", True, "chromium ok", level="ok"),
        raising=False,
    )

    items = {item.name: item for item in health.run_health_checks()}

    assert {item.level for item in items.values()}.issubset({"ok", "warning", "error"})
    assert items["dependency_pandas"].level == "ok"
    assert items["dependency_scipy"].level == "ok"
    assert items["dependency_talib"].level == "error"
    assert items["dependency_akshare"].level == "warning"
    assert items["dependency_tushare"].level == "warning"
    assert items["dependency_baostock"].level == "warning"
    assert items["chinese_font"].level == "warning"
    assert items["wechatmp_channel_enabled"].level == "ok"
    assert items["wechatmp_app_id"].level == "error"
    assert items["wechatmp_app_secret"].level == "error"
    assert items["wechatmp_token"].level == "ok"
    assert items["wechatmp_aes_key"].level == "warning"


def test_health_dependency_import_failure_reports_error_detail(investment_env, monkeypatch):
    from business.investment import health

    def fake_import(module_name):
        if module_name == "talib":
            raise OSError("DLL load failed while importing _ta_lib")
        return object()

    monkeypatch.setattr(health.importlib, "import_module", fake_import)

    item = health._check_dependency("dependency_talib", "talib", required=True)

    assert item.level == "error"
    assert item.ok is False
    assert "talib package import failed" in item.detail
    assert "DLL load failed" in item.detail


def test_run_health_checks_skips_smoke_by_default_and_runs_when_requested(investment_env, monkeypatch):
    from business.investment import health

    calls = []
    monkeypatch.setattr(
        health,
        "_run_technical_analysis_smoke",
        lambda: calls.append("ta") or health.HealthItem("smoke_technical_analysis", True, "ok", level="ok"),
        raising=False,
    )
    monkeypatch.setattr(
        health,
        "_run_renderer_smoke_checks",
        lambda: calls.append("renderer") or [
            health.HealthItem("smoke_renderer_ta", True, "ok", level="ok"),
            health.HealthItem("smoke_renderer_rate", True, "ok", level="ok"),
            health.HealthItem("smoke_renderer_cb", True, "ok", level="ok"),
        ],
        raising=False,
    )

    light_items = {item.name: item for item in health.run_health_checks()}

    assert calls == []
    assert "smoke_technical_analysis" not in light_items
    assert "smoke_renderer_ta" not in light_items

    full_items = {item.name: item for item in health.run_health_checks(run_smoke=True)}

    assert calls == ["ta", "renderer"]
    assert full_items["smoke_technical_analysis"].level == "ok"
    assert full_items["smoke_renderer_ta"].level == "ok"
    assert full_items["smoke_renderer_rate"].level == "ok"
    assert full_items["smoke_renderer_cb"].level == "ok"


def test_health_check_reports_stock_dictionary_and_tushare_token_without_leaking_secret(investment_env):
    from business.investment.config_service import save_config
    from business.investment.health import run_health_checks
    from business.investment.stock_resolver import refresh_stock_symbols

    save_config("tushare.token", "ts-health-secret-1234567890", operator_role="admin")
    refresh_stock_symbols(
        [
            {"code": "300502.SZ", "name": "新易盛", "market": "SZ", "source": "seed"},
            {"code": "600519.SH", "name": "贵州茅台", "market": "SH", "source": "seed"},
        ],
        source="seed",
    )

    items = {item.name: item for item in run_health_checks()}
    details = "\n".join(item.detail for item in items.values())

    assert items["stock_dictionary_table"].ok is True
    assert items["stock_dictionary_table"].detail == "exists"
    assert items["stock_dictionary_count"].ok is True
    assert items["stock_dictionary_count"].detail == "2"
    assert items["stock_dictionary_latest"].ok is True
    assert items["stock_dictionary_latest"].detail
    assert items["stock_dictionary_latest_source"].ok is True
    assert items["stock_dictionary_latest_source"].detail == "seed"
    assert items["tushare_token"].ok is True
    assert items["tushare_token"].detail == "configured"
    assert "ts-health-secret-1234567890" not in details


def test_health_check_reports_missing_and_unwritable_directories(investment_env, tmp_path, monkeypatch):
    from business.investment import health
    from business.investment.config_service import save_configs

    missing_upload = tmp_path / "missing-upload"
    generated_dir = tmp_path / "generated"
    generated_dir.mkdir()
    save_configs(
        {
            "storage.upload_dir": str(missing_upload),
            "storage.generated_dir": str(generated_dir),
        },
        operator_role="admin",
    )

    items = {item.name: item for item in health.run_health_checks()}

    assert items["upload_dir"].ok is False
    assert f"directory does not exist: {missing_upload}" == items["upload_dir"].detail

    def fail_mkstemp(*_args, **_kwargs):
        raise PermissionError("readonly")

    monkeypatch.setattr(health.tempfile, "mkstemp", fail_mkstemp)

    items = {item.name: item for item in health.run_health_checks()}

    assert items["generated_dir"].ok is False
    assert "directory not writable" in items["generated_dir"].detail
    assert "readonly" in items["generated_dir"].detail


def test_health_check_reports_all_dependencies_available(investment_env, tmp_path, monkeypatch):
    from business.investment import health
    from business.investment import config_service
    from business.investment.config_service import save_configs
    from business.investment.stock_resolver import refresh_stock_symbols

    files = {
        "technical_analysis.skill_path": tmp_path / "analyze_universal.py",
        "render.renderer_path": tmp_path / "render_card.py",
        "render.template_ta_path": tmp_path / "template_ta.html",
        "render.template_rate_path": tmp_path / "template_bond.html",
        "render.template_cb_path": tmp_path / "template_cb.html",
    }
    for path in files.values():
        path.write_text("ok", encoding="utf-8")
    upload_dir = tmp_path / "uploads"
    generated_dir = tmp_path / "generated"
    upload_dir.mkdir()
    generated_dir.mkdir()
    monkeypatch.setattr(
        config_service,
        "conf",
        lambda: {
            "bot_type": "custom",
            "model": "model",
            "custom_api_base": "https://model.example/v1",
            "custom_api_key": "sk-health-ok-1234567890",
        },
    )
    save_configs(
        {
            **{key: str(path) for key, path in files.items()},
            "storage.upload_dir": str(upload_dir),
            "storage.generated_dir": str(generated_dir),
            "tushare.token": "ts-health-ok-1234567890",
        },
        operator_role="admin",
    )
    monkeypatch.setattr(
        health,
        "_check_playwright_chromium",
        lambda: health.HealthItem("playwright_chromium", True, "chromium ok"),
        raising=False,
    )
    refresh_stock_symbols(
        [{"code": "300502.SZ", "name": "新易盛", "market": "SZ", "source": "seed"}],
        source="seed",
    )

    items = {item.name: item for item in health.run_health_checks()}

    assert all(item.ok for item in items.values())


def test_router_handles_rate_success_unauthorized_and_miss(investment_env, tmp_path):
    from business.investment.constants import ErrorCode
    from business.investment.constants import ServiceType
    from business.investment.config_service import save_config
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.records import list_request_records
    from business.investment.router import DEFAULT_UNMATCHED_PROMPT, handle_text_message, parse_route
    from business.investment.user_service import create_user

    secret = "sk-secret"
    save_config("tushare.token", secret, operator_role="admin")
    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])
    image = tmp_path / "rate.png"
    image.write_bytes(b"png")
    content_id = create_content_draft(ServiceType.RATE, source_text="rate")
    set_content_effective(content_id, str(image), operator="admin")

    assert parse_route("300502.SZ 技术分析").service_type == ServiceType.TECHNICAL_ANALYSIS
    assert parse_route("利率").service_type == ServiceType.RATE
    assert parse_route("转债").service_type == ServiceType.CONVERTIBLE_BOND
    assert parse_route("hello").matched is False

    success = handle_text_message("ok", "利率")
    assert success.success is True
    assert success.output_files == [str(image)]
    assert "[图片:" in success.reply_text

    denied = handle_text_message("missing", "利率")
    assert denied.success is False
    assert denied.reply_text == "您暂未开通该服务，如需开通请联系服务人员。"

    bypassed = handle_text_message("missing", "利率", skip_permission=True)
    assert bypassed.success is True
    assert bypassed.output_files == [str(image)]

    missing_image = tmp_path / f"missing-rate-{secret}.png"
    missing_content_id = create_content_draft(ServiceType.RATE, source_text="missing rate")
    set_content_effective(missing_content_id, str(missing_image), operator="admin")

    missing_image_reply = handle_text_message("ok", "利率")
    assert missing_image_reply.success is False
    assert missing_image_reply.reply_text == "今日内容尚未更新，请稍后再试。"
    assert missing_image_reply.output_files == []
    assert str(missing_image) not in missing_image_reply.reply_text
    assert "effective content image missing" in missing_image_reply.detail
    assert secret not in missing_image_reply.detail
    missing_image_record = list_request_records(limit=1)[0]
    assert missing_image_record.service_type == ServiceType.RATE
    assert missing_image_record.raw_input == "利率"
    assert missing_image_record.error_code == ErrorCode.NO_CONTENT
    assert missing_image_record.user_prompt == "今日内容尚未更新，请稍后再试。"
    assert "effective content image missing" in missing_image_record.error_message
    assert secret not in missing_image_record.error_message
    assert missing_image_record.output_files == []

    miss = handle_text_message("ok", "hello")
    assert miss.success is False
    assert miss.handled is True
    assert miss.reply_text == DEFAULT_UNMATCHED_PROMPT
    miss_record = list_request_records(limit=1)[0]
    assert miss_record.service_type == ServiceType.UNMATCHED
    assert miss_record.error_code == ErrorCode.INPUT_ERROR
    assert miss_record.user_prompt == DEFAULT_UNMATCHED_PROMPT


def test_parse_route_uses_configured_investment_skill_triggers(investment_env):
    from business.investment.config_service import save_config
    from business.investment.constants import ServiceType
    from business.investment.router import parse_route

    save_config("skill.rate.triggers", ["今日利率"], operator_role="admin", operator="pytest")
    save_config("skill.technical-analysis.triggers", ["走势分析"], operator_role="admin", operator="pytest")

    assert parse_route("利率").matched is False
    assert parse_route("今日利率").service_type == ServiceType.RATE

    route = parse_route("300502.SZ 走势分析")
    assert route.matched is True
    assert route.service_type == ServiceType.TECHNICAL_ANALYSIS
    assert route.target_text == "300502.SZ"


def test_parse_route_ignores_disabled_investment_skill(investment_env):
    from business.investment.config_service import save_config
    from business.investment.router import parse_route

    save_config("skill.rate.enabled", False, operator_role="admin", operator="pytest")

    assert parse_route("利率").matched is False


def test_router_can_explicitly_fallback_to_general_agent_for_unmatched_text(investment_env):
    from business.investment.config_service import save_config
    from business.investment.constants import ServiceType
    from business.investment.router import DEFAULT_UNMATCHED_PROMPT, handle_text_message
    from business.investment.user_service import create_user

    save_config("router.enable_agent_fallback", True, operator_role="admin")
    create_user("ok", enabled=True, allowed_services=[ServiceType.ALL])

    miss = handle_text_message("ok", "hello")

    assert miss.success is False
    assert miss.handled is False
    assert miss.reply_text == DEFAULT_UNMATCHED_PROMPT


def test_web_channel_routes_investment_commands_from_admin_chat(investment_env, tmp_path):
    from bridge.reply import ReplyType
    from business.investment.constants import ServiceType
    from business.investment.router import DEFAULT_UNMATCHED_PROMPT
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.user_service import create_user
    from channel.web.web_channel import WebChannel, _build_investment_web_reply

    create_user("web-session-user", enabled=True, allowed_services=[ServiceType.ALL])
    image = tmp_path / "rate card.png"
    image.write_bytes(b"png")
    content_id = create_content_draft(ServiceType.RATE, source_text="rate")
    set_content_effective(content_id, str(image), operator="admin")

    reply = _build_investment_web_reply("web-session-user", "利率")

    assert reply is not None
    assert reply.type == ReplyType.TEXT
    assert "![rate card.png](/api/file?path=" in reply.content
    assert "%20" in reply.content
    assert "[图片:" not in reply.content
    plain = _build_investment_web_reply("web-session", "普通聊天")
    assert plain is not None
    assert plain.type == ReplyType.TEXT
    assert plain.content == DEFAULT_UNMATCHED_PROMPT
    assert WebChannel().channel_type == "web"


def test_web_channel_uses_configured_investment_skill_triggers(investment_env, tmp_path):
    from bridge.reply import ReplyType
    from business.investment.config_service import save_config
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.user_service import create_user
    from channel.web.web_channel import _build_investment_web_reply

    create_user("web-session", enabled=True, allowed_services=[ServiceType.ALL])
    save_config("skill.rate.triggers", ["今日利率"], operator_role="admin", operator="pytest")
    image = tmp_path / "rate override.png"
    image.write_bytes(b"png")
    content_id = create_content_draft(ServiceType.RATE, source_text="rate")
    set_content_effective(content_id, str(image), operator="admin")

    reply = _build_investment_web_reply("web-session", "今日利率")

    assert reply is not None
    assert reply.type == ReplyType.TEXT
    assert "![rate override.png](/api/file?path=" in reply.content


def test_web_channel_uses_admin_session_instead_of_customer_permission(investment_env, tmp_path):
    from bridge.reply import ReplyType
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.user_service import create_user
    from channel.web.web_channel import _build_investment_web_reply

    create_user("disabled-web", enabled=False, allowed_services=[ServiceType.ALL])
    image = tmp_path / "admin-rate.png"
    image.write_bytes(b"png")
    content_id = create_content_draft(ServiceType.RATE, source_text="rate")
    set_content_effective(content_id, str(image), operator="admin")

    reply = _build_investment_web_reply("disabled-web", "利率")

    assert reply is not None
    assert reply.type == ReplyType.TEXT
    assert "![admin-rate.png](/api/file?path=" in reply.content


def test_web_open_chat_uses_plain_model_without_agent_bridge(investment_env, monkeypatch):
    from bridge.reply import Reply, ReplyType
    from business.investment.config_service import save_config
    from bridge.context import Context, ContextType
    from channel.chat_channel import ChatChannel
    from channel.web.web_channel import WebChannel, WebMessage

    save_config("router.enable_web_open_chat", True, operator_role="admin")
    calls = []

    def fake_fetch_reply_content(query, context):
        calls.append((query, context.get("session_id")))
        return Reply(ReplyType.TEXT, "plain model reply")

    import bridge.bridge as bridge_module

    monkeypatch.setattr(bridge_module.Bridge(), "fetch_reply_content", fake_fetch_reply_content)
    monkeypatch.setattr(
        ChatChannel,
        "_generate_reply",
        lambda *_args, **_kwargs: pytest.fail("web open chat must not enter AgentBridge path"),
    )

    context = Context(ContextType.TEXT, "普通聊天", {"msg": WebMessage("msg-1", "普通聊天")})
    context["session_id"] = "web-session"

    reply = WebChannel()._generate_reply(context)

    assert reply.type == ReplyType.TEXT
    assert reply.content == "plain model reply"
    assert calls == [("普通聊天", "web-session")]


def test_wechatmp_channel_uses_effective_content_and_permission_prompts(investment_env, tmp_path, monkeypatch):
    from bridge.context import Context, ContextType
    from bridge.reply import ReplyType
    from business.investment.constants import ServiceType
    from business.investment.daily_content import create_content_draft, set_content_effective
    from business.investment.user_service import create_user
    import channel.wechatmp.wechatmp_channel as wechatmp_channel

    instances = wechatmp_channel.WechatMPChannel.__closure__[1].cell_contents
    instances.clear()
    monkeypatch.setattr(
        wechatmp_channel,
        "conf",
        lambda: {
            "wechatmp_app_id": "wx-test",
            "wechatmp_app_secret": "secret",
            "wechatmp_token": "token",
            "wechatmp_aes_key": "",
            "single_chat_prefix": [""],
            "concurrency_in_session": 1,
        },
    )
    channel = wechatmp_channel.WechatMPChannel()

    rate_image = tmp_path / "wechat-rate.png"
    cb_image = tmp_path / "wechat-cb.png"
    rate_image.write_bytes(b"rate")
    cb_image.write_bytes(b"cb")
    rate_id = create_content_draft(ServiceType.RATE, source_text="rate", effective_date="2026-05-28")
    cb_id = create_content_draft(ServiceType.CONVERTIBLE_BOND, source_text="cb", effective_date="2026-05-28")
    set_content_effective(rate_id, str(rate_image), operator="tester")
    set_content_effective(cb_id, str(cb_image), operator="tester")

    create_user("openid-ok", enabled=True, allowed_services=[ServiceType.ALL], auth_end_at="2099-12-31T23:59:59")
    create_user("openid-disabled", enabled=False, allowed_services=[ServiceType.ALL])
    create_user("openid-expired", enabled=True, allowed_services=[ServiceType.ALL], auth_end_at="2020-01-01T00:00:00")

    def context(openid, content):
        msg = SimpleNamespace(from_user_id=openid, other_user_id=openid, msg_id=f"{openid}-{content}")
        return Context(
            ContextType.TEXT,
            content,
            {
                "msg": msg,
                "session_id": openid,
                "receiver": openid,
                "channel_type": "wechatmp",
                "isgroup": False,
            },
        )

    rate_reply = channel._generate_reply(context("openid-ok", "利率"))
    cb_reply = channel._generate_reply(context("openid-ok", "转债"))
    missing_reply = channel._generate_reply(context("openid-missing", "利率"))
    disabled_reply = channel._generate_reply(context("openid-disabled", "利率"))
    expired_reply = channel._generate_reply(context("openid-expired", "利率"))

    assert rate_reply.type == ReplyType.IMAGE_URL
    assert rate_reply.content == [str(rate_image)]
    assert cb_reply.type == ReplyType.IMAGE_URL
    assert cb_reply.content == [str(cb_image)]
    assert missing_reply.type == ReplyType.TEXT
    assert missing_reply.content == "您暂未开通该服务，如需开通请联系服务人员。"
    assert disabled_reply.type == ReplyType.TEXT
    assert disabled_reply.content == "您的服务已停用，如需恢复请联系服务人员。"
    assert expired_reply.type == ReplyType.TEXT
    assert expired_reply.content == "您的授权已过期，如需续期请联系服务人员。"
